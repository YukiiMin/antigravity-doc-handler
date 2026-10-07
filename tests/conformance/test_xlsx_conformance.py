"""
tests.conformance.test_xlsx_conformance — End-to-End Conformance test suite for Phase 2.1 XLSX MVP.
Xâu chuỗi trọn vẹn vòng đời các công cụ MCP:
Preflight -> Register -> Lint -> Inspect -> Mutate -> Recalc -> Diff -> Universal Gates Validate.
"""

import os
import shutil
import tempfile
import unittest
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from doctools.infra.file_store import FileStore
from doctools.operations.xlsx import (
    register_xlsx_inspect_tools,
    register_xlsx_mutate_tools,
    register_xlsx_preflight_tools,
    register_xlsx_recalc_tools,
    register_xlsx_template_tools,
    register_xlsx_validate_tools,
)
from doctools.registry import ToolRegistry


class TestXlsxConformance(unittest.TestCase):
    """Kiểm thử chuỗi chuyển đổi toàn diện cho Phase 2.1 XLSX MVP."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.mkdtemp()
        self.file_store = FileStore(base_dir=os.path.join(self.temp_dir, "filestore"))
        self.registry = ToolRegistry(file_store=self.file_store)

        # Đăng ký trọn bộ công cụ MCP Phase 2.1
        register_xlsx_preflight_tools(self.registry)
        register_xlsx_template_tools(self.registry)
        register_xlsx_inspect_tools(self.registry)
        register_xlsx_mutate_tools(self.registry)
        register_xlsx_recalc_tools(self.registry)
        register_xlsx_validate_tools(self.registry)

        # Khởi tạo template bảng tính chuẩn
        self.template_path = os.path.join(self.temp_dir, "invoice_template.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Invoice"

        ws["A1"] = "INVOICE REPORT"
        ws["A2"] = "Code"
        ws["B2"] = "Description"
        ws["C2"] = "Qty"
        ws["D2"] = "Price"
        ws["E2"] = "Amount"

        # Prototype row ở dòng 3
        thin = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        ws["A3"] = "001"
        ws["A3"].number_format = "@"
        ws["A3"].border = thin
        ws["B3"] = "Default Item"
        ws["B3"].border = thin
        ws["C3"] = 1
        ws["C3"].border = thin
        ws["D3"] = 100
        ws["D3"].border = thin
        ws["E3"] = "=C3*D3"
        ws["E3"].border = thin

        # Dòng 5: Tổng cộng
        ws["D5"] = "Total"
        ws["E5"] = "=SUM(E3:E4)"

        wb.save(self.template_path)

    def tearDown(self) -> None:
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_complete_xlsx_lifecycle_pipeline(self) -> None:
        """Thực thi tuần tự toàn bộ chuỗi công cụ MCP xlsx.* từ Preflight đến Gates."""
        # 1. xlsx.preflight
        res_preflight = self.registry.execute("xlsx.preflight", {"file_ref": self.template_path})
        self.assertTrue(res_preflight.success)
        inv = res_preflight.stats.extra["inventory"]
        self.assertEqual(len(inv["sheets"]), 1)
        self.assertEqual(inv["sheets"][0]["name"], "Invoice")

        # 2. xlsx.register_template
        manifest = {
            "template_id": "tpl_invoice_master",
            "version": 1,
            "reference_sheet": "Invoice",
            "prototype_rows": {"Invoice": 3},
            "locked_zones": ["Invoice!A1:E2", "Invoice!D5:E5"],
        }
        res_reg = self.registry.execute(
            "xlsx.register_template",
            {"template_ref": self.template_path, "manifest": manifest},
        )
        self.assertTrue(res_reg.success)

        # 3. xlsx.lint_template
        res_lint = self.registry.execute(
            "xlsx.lint_template",
            {"template_ref": self.template_path, "manifest": manifest},
        )
        self.assertTrue(res_lint.success)

        # 4. xlsx.inspect
        res_inspect = self.registry.execute("xlsx.inspect", {"file_ref": self.template_path})
        self.assertTrue(res_inspect.success)
        self.assertEqual(res_inspect.stats.extra["total_formulas"], 2)

        # 5. xlsx.mutate (nhân bản dữ liệu theo MutationSpec)
        mutated_path = os.path.join(self.temp_dir, "invoice_mutated.xlsx")
        mutate_spec = {
            "template_ref_or_id": self.template_path,
            "expansions": [
                {
                    "sheet": "Invoice",
                    "prototype_row": 3,
                    "consume_prototype": True,
                    "rows_data": [
                        {"A": "010", "B": "Server Node Alpha", "C": 2, "D": 500},
                        {"A": "011", "B": "Database Storage", "C": 5, "D": 200},
                        {"A": "012", "B": "Load Balancer", "C": 1, "D": 300},
                    ],
                    "id_columns": ["A"],
                }
            ],
        }
        res_mutate = self.registry.execute(
            "xlsx.mutate",
            {"spec": mutate_spec, "output_path": mutated_path},
        )
        self.assertTrue(res_mutate.success)
        self.assertTrue(os.path.exists(mutated_path))

        # 6. xlsx.recalc (tiêm cache giá trị vào công thức)
        recalc_path = os.path.join(self.temp_dir, "invoice_recalc.xlsx")
        cached_data = {
            "Invoice": {
                "E3": 1000,
                "E4": 1000,
                "E5": 300,
                "E7": 2300,
            }
        }
        res_recalc = self.registry.execute(
            "xlsx.recalc",
            {
                "file_ref": mutated_path,
                "method": "cache_writer",
                "cached_values": cached_data,
                "output_path": recalc_path,
            },
        )
        self.assertTrue(res_recalc.success)
        self.assertTrue(os.path.exists(recalc_path))

        # 7. xlsx.diff (so khớp template gốc vs file kết quả)
        res_diff = self.registry.execute(
            "xlsx.diff",
            {"original_ref": self.template_path, "modified_ref": recalc_path},
        )
        self.assertTrue(res_diff.success)
        diff_rep = res_diff.stats.extra["diff_report"]
        self.assertEqual(diff_rep["total_errors_mod"], 0)
        self.assertIn("Invoice", diff_rep["sheet_diffs"])
        self.assertTrue(diff_rep["sheet_diffs"]["Invoice"]["row_delta"] > 0)

        # 8. xlsx.validate (kiểm định 13 Universal Gates UG-01..13)
        res_val = self.registry.execute("xlsx.validate", {"file_ref": recalc_path})
        self.assertTrue(res_val.success)
        self.assertEqual(len(res_val.diagnostics.errors), 0)
        self.assertIn("13_UNIVERSAL_GATES_VERIFIED", res_val.guarantees_applied)


if __name__ == "__main__":
    unittest.main()
