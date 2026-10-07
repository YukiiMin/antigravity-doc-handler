"""
tests/test_xlsx/test_template_registry.py — Bộ kiểm thử cho Template Registry, Manifest và Linter XLSX.
Kiểm tra các yêu cầu FR-03, FR-04, TC-18, TC-19:
- Phân tích và kiểm định schema manifest
- Phát hiện anchor thiếu hoặc mơ hồ (E-TPL-ANCHOR-MISSING/AMBIGUOUS)
- Cảnh báo dòng mẫu có nguy cơ bị tính nhầm (W-TPL-SAMPLE-ROW)
- Đăng ký template vào kho, kiểm tra tính bất biến và sha256
- Tích hợp điều phối qua ToolRegistry trung tâm.
"""

from pathlib import Path
import tempfile
import unittest
import openpyxl

from doctools.contract.issues import Severity
from doctools.contract.xlsx.manifest import XlsxTemplateManifest
from doctools.core.xlsx.template.manifest_parser import (
    ManifestParseError,
    parse_locked_zone,
    parse_manifest,
    validate_manifest_against_template,
)
from doctools.core.xlsx.template.template_linter import XlsxTemplateLinter
from doctools.core.xlsx.template.template_registry import XlsxTemplateRegistry
from doctools.infra.file_store import FileStore
from doctools.operations.xlsx.template_ops import (
    register_xlsx_template_tools,
    xlsx_lint_template,
    xlsx_list_templates,
    xlsx_register_template,
)
from doctools.registry import ToolRegistry


class TestXlsxTemplateRegistry(unittest.TestCase):
    """Test suite cho manifest, linter và registry của XLSX template."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_path = Path(self.temp_dir.name)
        self.file_store = FileStore(base_dir=self.base_path / "filestore")
        self.registry = XlsxTemplateRegistry(file_store=self.file_store)

    def tearDown(self):
        self.temp_dir.cleanup()

    def _create_sample_workbook(self, path: Path) -> None:
        """Tạo workbook mẫu có sheet Example và Function 1."""
        wb = openpyxl.Workbook()
        ws_ex = wb.active
        ws_ex.title = "Example"
        ws_ex["A1"] = "BẢNG KẾT QUẢ KIỂM THỬ"
        ws_ex["A3"] = "Test ID"
        ws_ex["B3"] = "Mô tả test case"
        ws_ex["C3"] = "Trạng thái"
        ws_ex["A4"] = "TC001"
        ws_ex["B4"] = "Sample item"
        ws_ex["C4"] = "Pass"

        ws_fn = wb.create_sheet(title="Function 1")
        ws_fn["A1"] = "Function 1"

        wb.save(path)
        wb.close()

    def test_manifest_parse_valid(self):
        """Manifest hợp lệ được parse thành công."""
        raw_manifest = {
            "template_id": "tpl_test_matrix",
            "version": 1,
            "reference_sheet": "Example",
            "sibling_sheets": ["Function 1"],
            "anchors": {
                "header": {"kind": "header", "keyword": "Test ID", "scope": "first_table"},
            },
            "prototype_rows": {"Example": 4},
            "locked_zones": ["Example!A1:C2", "*!A1:B2"],
            "id_columns": ["Test ID"],
        }
        manifest = parse_manifest(raw_manifest)
        self.assertEqual(manifest.template_id, "tpl_test_matrix")
        self.assertEqual(manifest.reference_sheet, "Example")
        self.assertEqual(len(manifest.sibling_sheets), 1)

    def test_manifest_locked_zone_syntax(self):
        """Cú pháp locked_zones được parse chính xác."""
        sheet_pat, bounds = parse_locked_zone("Statistics!C12:I19")
        self.assertEqual(sheet_pat, "Statistics")
        self.assertEqual(bounds, (3, 12, 9, 19))

        with self.assertRaises(ValueError):
            parse_locked_zone("InvalidZoneWithoutBang")

    def test_validate_manifest_detects_missing_sheet(self):
        """Phát hiện lỗi khi reference_sheet không tồn tại trong workbook."""
        wb_path = self.base_path / "test.xlsx"
        self._create_sample_workbook(wb_path)

        manifest = XlsxTemplateManifest(
            template_id="tpl_missing",
            reference_sheet="NonExistentSheet",
        )
        issues = validate_manifest_against_template(manifest, wb_path)
        err_codes = [i.code for i in issues]
        self.assertIn("E-TPL-REF-SHEET-NOT-FOUND", err_codes)

    def test_validate_manifest_detects_missing_anchor(self):
        """Phát hiện lỗi khi anchor keyword không có trong workbook (TC-18)."""
        wb_path = self.base_path / "test_anchor.xlsx"
        self._create_sample_workbook(wb_path)

        manifest = XlsxTemplateManifest(
            template_id="tpl_missing_anchor",
            reference_sheet="Example",
            anchors={"header": {"keyword": "KeywordDoesNotExist"}},
        )
        issues = validate_manifest_against_template(manifest, wb_path)
        err_codes = [i.code for i in issues]
        self.assertIn("E-TPL-ANCHOR-MISSING", err_codes)

    def test_validate_manifest_detects_ambiguous_anchor(self):
        """Phát hiện lỗi khi anchor keyword trùng lặp trong first_table scope (TC-18)."""
        wb_path = self.base_path / "test_ambig.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Example"
        ws["A3"] = "Test ID"
        ws["D3"] = "Test ID"  # Trùng lặp trong cùng bảng
        wb.save(wb_path)
        wb.close()

        manifest = XlsxTemplateManifest(
            template_id="tpl_ambig",
            reference_sheet="Example",
            anchors={"header": {"keyword": "Test ID", "scope": "first_table"}},
        )
        issues = validate_manifest_against_template(manifest, wb_path)
        err_codes = [i.code for i in issues]
        self.assertIn("E-TPL-ANCHOR-AMBIGUOUS", err_codes)

    def test_linter_detects_placeholders_and_sample_rows(self):
        """Linter phát hiện placeholder và cảnh báo W-TPL-SAMPLE-ROW (TC-18)."""
        wb_path = self.base_path / "test_lint.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Example"
        ws["A1"] = "Báo cáo {{ project_name }}"
        ws["A4"] = "Sample item"
        wb.save(wb_path)
        wb.close()

        manifest = XlsxTemplateManifest(
            template_id="tpl_lint",
            reference_sheet="Example",
            prototype_rows={"Example": 4},
        )
        linter = XlsxTemplateLinter()
        report = linter.lint(wb_path, manifest=manifest)
        self.assertTrue(any("{{ project_name }}" in p for p in report.placeholders_found))
        codes = [i.code for i in report.issues]
        self.assertIn("W-TPL-SAMPLE-ROW", codes)

    def test_register_template_success_and_immutability(self):
        """Đăng ký template thành công, kiểm tra sha256 và chống ghi đè phiên bản."""
        wb_path = self.base_path / "master.xlsx"
        self._create_sample_workbook(wb_path)

        manifest = XlsxTemplateManifest(
            template_id="tpl_master",
            version=1,
            reference_sheet="Example",
            sibling_sheets=["Function 1"],
            anchors={"header": {"keyword": "Test ID"}},
            prototype_rows={"Example": 4},
        )

        record, issues = self.registry.register(wb_path, manifest)
        self.assertIsNotNone(record)
        self.assertEqual(len([i for i in issues if i.severity == Severity.ERROR]), 0)
        self.assertEqual(record.template_id, "tpl_master")
        self.assertEqual(record.version, 1)

        # Truy xuất lại
        fetched = self.registry.get("tpl_master", 1)
        self.assertIsNotNone(fetched)
        self.assertEqual(fetched.sha256, record.sha256)

        # Thử ghi đè cùng version với nội dung khác -> Version collision
        diff_path = self.base_path / "diff.xlsx"
        self._create_sample_workbook(diff_path)
        wb2 = openpyxl.load_workbook(diff_path)
        ws2 = wb2["Example"]
        ws2["Z100"] = "Different Content"
        wb2.save(diff_path)
        wb2.close()

        rec2, issues2 = self.registry.register(diff_path, manifest)
        self.assertIsNone(rec2)
        self.assertTrue(any(i.code == "E-TPL-VERSION-COLLISION" for i in issues2))

    def test_mcp_template_tools_dispatch(self):
        """Kiểm tra tích hợp các MCP tool trong ToolRegistry."""
        wb_path = self.base_path / "mcp_tpl.xlsx"
        self._create_sample_workbook(wb_path)

        mcp_registry = ToolRegistry(file_store=self.file_store)
        register_xlsx_template_tools(mcp_registry, template_registry=self.registry)

        manifest_data = {
            "template_id": "tpl_mcp_test",
            "version": 1,
            "reference_sheet": "Example",
            "anchors": {"header": {"keyword": "Test ID"}},
        }

        # 1. xlsx.register_template
        reg_env = mcp_registry.execute(
            "xlsx.register_template",
            {"template_ref": str(wb_path), "manifest": manifest_data},
        )
        self.assertTrue(reg_env.success)
        self.assertIn("template_immutable_stored", reg_env.guarantees_applied)

        # 2. xlsx.lint_template
        lint_env = mcp_registry.execute(
            "xlsx.lint_template",
            {"template_ref": str(wb_path), "manifest": manifest_data},
        )
        self.assertTrue(lint_env.success)

        # 3. xlsx.list_templates
        list_env = mcp_registry.execute("xlsx.list_templates", {})
        self.assertTrue(list_env.success)
        self.assertGreaterEqual(list_env.stats.elements_processed, 1)


if __name__ == "__main__":
    unittest.main()
