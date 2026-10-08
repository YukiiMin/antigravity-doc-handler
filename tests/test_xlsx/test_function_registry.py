"""
tests.test_xlsx.test_function_registry
Test suite for Phase 2.3d: Function Registry & Fx Differential Matrix (TC-50, TC-51).
Verifies:
- TC-50: Differential matrix summary, prefix tracking (_xlfn.), and EV-14 modern functions
- TC-51: Function auditing (W-FX-UNKNOWN, W-FX-COMPAT-PREFIX, W-FX-RISKY)
"""

from pathlib import Path
import tempfile
import unittest

import openpyxl

from doctools.contract.envelope import ResultEnvelope
from doctools.core.xlsx.template.function_registry import (
    FunctionEntry,
    FunctionRegistry,
)
from doctools.operations.xlsx.inspect_ops import (
    register_xlsx_inspect_tools,
    xlsx_function_catalog,
)
from doctools.registry import ToolRegistry


class TestFunctionRegistry(unittest.TestCase):
    """Unit tests for Phase 2.3d: Function Registry."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_file = Path(self.temp_dir.name) / "test_fx.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "FxSheet"

        # Formula 1: Modern function requiring prefix (XLOOKUP)
        ws["A1"] = "='_xlfn.XLOOKUP'(1, A2:A10, B2:B10)"
        # Formula 2: Risky function (WEBSERVICE)
        ws["A2"] = '=WEBSERVICE("http://example.com/api")'
        # Formula 3: Unknown / custom user function (MY_CUSTOM_FUNC)
        ws["A3"] = "=MY_CUSTOM_FUNC(10, 20)"
        # Formula 4: Classical standard function (SUM)
        ws["A4"] = "=SUM(B1:B10)"

        wb.save(self.test_file)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_tc50_modern_functions_registered(self) -> None:
        """TC-50: Verify modern Excel functions independent of legacy openpyxl FORMULAE."""
        reg = FunctionRegistry()
        for fn_name in ["XLOOKUP", "FILTER", "TEXTJOIN", "IFS", "SWITCH", "LET", "LAMBDA"]:
            entry = reg.get(fn_name)
            self.assertIsNotNone(entry, f"Expected {fn_name} in FunctionRegistry")
            self.assertEqual(entry.name, fn_name)

        # Verify prefix tracking
        xlookup = reg.get("XLOOKUP")
        self.assertEqual(xlookup.prefix, "_xlfn.")
        filter_fn = reg.get("FILTER")
        self.assertEqual(filter_fn.prefix, "_xlfn._xlws.")

    def test_tc50_differential_matrix_summary(self) -> None:
        """TC-50: Verify differential matrix summary statistics."""
        res: ResultEnvelope = xlsx_function_catalog()
        self.assertTrue(res.success)
        extra = res.stats.extra

        diff = extra["differential_matrix"]
        self.assertGreater(diff["total_registered_functions"], 30)
        self.assertIn("status_distribution", diff)
        self.assertGreaterEqual(diff["status_distribution"]["MATCH"], 20)
        self.assertGreater(diff["match_rate"], 50.0)

    def test_tc51_function_audit_warnings(self) -> None:
        """TC-51: Verify W-FX-UNKNOWN, W-FX-COMPAT-PREFIX, and W-FX-RISKY diagnostics."""
        res: ResultEnvelope = xlsx_function_catalog(file_ref_or_path=self.test_file)
        self.assertTrue(res.success)
        all_issues = res.diagnostics.errors + res.diagnostics.warnings + res.diagnostics.info
        issue_codes = [iss.code for iss in all_issues]

        # 1. Unknown function warning
        self.assertIn("W-FX-UNKNOWN", issue_codes)
        # 2. Compatibility prefix warning
        self.assertIn("W-FX-COMPAT-PREFIX", issue_codes)
        # 3. Risky function warning
        self.assertIn("W-FX-RISKY", issue_codes)

        # Check audit breakdown
        audit = res.stats.extra["workbook_audit"]
        self.assertIn("MY_CUSTOM_FUNC", audit["unknown_functions"])
        risky_names = [r["name"] for r in audit["risky_functions"]]
        self.assertIn("WEBSERVICE", risky_names)

    def test_tool_registry_registration(self) -> None:
        """Verify xlsx.function_catalog is registered in ToolRegistry."""
        registry = ToolRegistry()
        register_xlsx_inspect_tools(registry)
        tool_names = [t["name"] for t in registry.list_tools()]
        self.assertIn("xlsx.function_catalog", tool_names)


if __name__ == "__main__":
    unittest.main()
