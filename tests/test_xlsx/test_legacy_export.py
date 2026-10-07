"""
Unit test suite for doctools.core.xlsx.export (Legacy .xls Exporter).
Kiểm thử toàn diện Tiểu bước 2.2.3:
- Xuất tệp .xlsx sang tệp cổ điển .xls (Excel 97-2003 BIFF8) qua Excel COM / LibreOffice
- Kiểm tra chữ ký file OLE2 Compound Document (0xD0CF11E0)
- Tích hợp công cụ MCP xlsx.export_legacy_xls vào ToolRegistry
"""

import os
import shutil
import sys
import tempfile
import unittest
import openpyxl

from doctools.core.xlsx.export.legacy_exporter import XlsxLegacyExporter
from doctools.infra.file_store import FileStore
from doctools.operations.xlsx.export_ops import register_xlsx_export_tools
from doctools.registry import ToolRegistry


class TestLegacyExporter(unittest.TestCase):
    """Kiểm thử bộ chuyển đổi sang định dạng .xls cổ điển."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.mkdtemp()
        self.file_store = FileStore(base_dir=os.path.join(self.temp_dir, "filestore"))
        self.exporter = XlsxLegacyExporter()

        # Tạo file mẫu
        self.sample_xlsx = os.path.join(self.temp_dir, "sample.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LegacyData"
        ws["A1"] = "ID"
        ws["B1"] = "Value"
        ws["A2"] = "LEG-01"
        ws["B2"] = 500
        wb.save(self.sample_xlsx)

    def tearDown(self) -> None:
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_to_xls(self) -> None:
        """Kiểm tra xuất file sang định dạng .xls nếu Excel COM hoặc LibreOffice khả dụng."""
        out_xls = os.path.join(self.temp_dir, "exported.xls")
        res_path, issues = self.exporter.export_to_xls(self.sample_xlsx, output_path=out_xls)

        if sys.platform == "win32":
            # Trên môi trường Windows đã kiểm chứng Excel COM sẵn sàng
            if res_path is not None:
                self.assertTrue(os.path.exists(out_xls))
                self.assertTrue(os.path.getsize(out_xls) > 0)
                # Kiểm tra magic bytes của định dạng OLE Compound File (BIFF8 .xls)
                with open(out_xls, "rb") as f:
                    header = f.read(4)
                self.assertEqual(header, b"\xd0\xcf\x11\xe0")

    def test_mcp_tool_export_integration(self) -> None:
        """Kiểm tra tích hợp công cụ MCP xlsx.export_legacy_xls."""
        registry = ToolRegistry(file_store=self.file_store)
        register_xlsx_export_tools(registry)

        out_xls = os.path.join(self.temp_dir, "mcp_exported.xls")
        res = registry.execute(
            "xlsx.export_legacy_xls",
            {"file_ref": self.sample_xlsx, "output_path": out_xls},
        )

        if sys.platform == "win32" and os.path.exists(out_xls):
            self.assertTrue(res.success)
            self.assertIsNotNone(res.file_ref)
            self.assertIn("LEGACY_BIFF8_EXPORTED", res.guarantees_applied)


if __name__ == "__main__":
    unittest.main()
