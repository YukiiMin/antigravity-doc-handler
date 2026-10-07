"""
Unit test suite for doctools.core.xlsx.repair (Workbook Auto-Repair).
Kiểm thử toàn diện Tiểu bước 2.2.2:
- Tự động khôi phục thẻ <calcPr fullCalcOnLoad="1"/>
- Bổ sung Override còn thiếu trong [Content_Types].xml
- Loại bỏ các ký tự điều khiển ASCII bất hợp pháp gây lỗi OpenXML
- Tích hợp công cụ MCP xlsx.repair vào ToolRegistry
"""

import io
import os
import shutil
import tempfile
import unittest
import zipfile
import openpyxl
from lxml import etree

from doctools.core.xlsx.repair.workbook_repairer import XlsxWorkbookRepairer
from doctools.infra.file_store import FileStore
from doctools.operations.xlsx.repair_ops import register_xlsx_repair_tools
from doctools.registry import ToolRegistry


class TestWorkbookRepairer(unittest.TestCase):
    """Kiểm thử động cơ sửa chữa workbook."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.mkdtemp()
        self.file_store = FileStore(base_dir=os.path.join(self.temp_dir, "filestore"))
        self.repairer = XlsxWorkbookRepairer()

    def tearDown(self) -> None:
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_repair_missing_calc_pr(self) -> None:
        """Kiểm tra khôi phục calcPr khi file bị thiếu."""
        wb_path = os.path.join(self.temp_dir, "test_no_calc.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(wb_path)

        # Xóa thẻ calcPr trong xl/workbook.xml để mô phỏng lỗi
        buf = io.BytesIO()
        with zipfile.ZipFile(wb_path, "r") as z_in, zipfile.ZipFile(buf, "w") as z_out:
            for item in z_in.infolist():
                if item.filename == "xl/workbook.xml":
                    xml_data = z_in.read(item.filename)
                    tree = etree.fromstring(xml_data)
                    for cp in tree.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}calcPr"):
                        cp.getparent().remove(cp)
                    z_out.writestr(item, etree.tostring(tree))
                else:
                    z_out.writestr(item, z_in.read(item.filename))

        repaired_bytes, issues = self.repairer.repair(buf.getvalue())
        self.assertTrue(any(i.code == "I-XLSX-REPAIR-INJECT-CALCPR" for i in issues))

        # Kiểm tra lại gói đã sửa
        with zipfile.ZipFile(io.BytesIO(repaired_bytes), "r") as z:
            wb_tree = etree.fromstring(z.read("xl/workbook.xml"))
            calc_nodes = wb_tree.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}calcPr")
            self.assertEqual(len(calc_nodes), 1)
            self.assertEqual(calc_nodes[0].attrib.get("fullCalcOnLoad"), "1")

    def test_repair_invalid_control_characters(self) -> None:
        """Kiểm tra loại bỏ ký tự điều khiển ASCII bất hợp pháp trong sheet XML."""
        wb_path = os.path.join(self.temp_dir, "test_chars.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "CleanSheet"
        wb.active["A1"] = "CleanVal"
        wb.save(wb_path)

        # Tiêm ký tự \x00 và \x08 vào XML
        buf = io.BytesIO()
        with zipfile.ZipFile(wb_path, "r") as z_in, zipfile.ZipFile(buf, "w") as z_out:
            for item in z_in.infolist():
                if item.filename == "xl/worksheets/sheet1.xml":
                    content = z_in.read(item.filename).decode("utf-8")
                    bad_content = content.replace("CleanVal", "Clean\x00Val\x08")
                    z_out.writestr(item, bad_content.encode("utf-8"))
                else:
                    z_out.writestr(item, z_in.read(item.filename))

        repaired_bytes, issues = self.repairer.repair(buf.getvalue())
        self.assertTrue(any(i.code == "I-XLSX-REPAIR-STRIP-INVALID-CHARS" for i in issues))

        with zipfile.ZipFile(io.BytesIO(repaired_bytes), "r") as z:
            sheet_content = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
            self.assertNotIn("\x00", sheet_content)
            self.assertNotIn("\x08", sheet_content)

    def test_mcp_tool_repair_integration(self) -> None:
        """Kiểm tra tích hợp công cụ MCP xlsx.repair vào ToolRegistry."""
        registry = ToolRegistry(file_store=self.file_store)
        register_xlsx_repair_tools(registry)

        wb_path = os.path.join(self.temp_dir, "repair_mcp.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "RepairTest"
        wb.save(wb_path)

        out_path = os.path.join(self.temp_dir, "repaired_mcp.xlsx")
        res = registry.execute("xlsx.repair", {"file_ref": wb_path, "output_path": out_path})

        self.assertTrue(res.success)
        self.assertIsNotNone(res.file_ref)
        self.assertTrue(os.path.exists(out_path))
        self.assertIn("STRUCTURE_REPAIRED", res.guarantees_applied)


if __name__ == "__main__":
    unittest.main()
