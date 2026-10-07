"""
Unit test suite for doctools.core.xlsx.recalc (Recalc Engine & Cache Writer).
Kiểm thử toàn diện Tiểu bước 2.1.6:
- Tiêm giá trị cache <v> kiểu số (n), chuỗi (str), logic (b) bằng lxml
- Bảo toàn 100% cú pháp công thức <f>
- Kích hoạt cấu hình fullCalcOnLoad="1" trong calcPr
- Điều phối chiến lược RecalcEngine (cache_writer, auto)
- Tích hợp công cụ MCP xlsx.recalc vào ToolRegistry
"""

import io
import os
import shutil
import tempfile
import unittest
import zipfile
import openpyxl
from lxml import etree

from doctools.core.xlsx.recalc.cache_writer import XlsxCacheWriter
from doctools.core.xlsx.recalc.recalc_engine import RecalcEngine
from doctools.infra.file_store import FileStore
from doctools.operations.xlsx.recalc_ops import register_xlsx_recalc_tools
from doctools.registry import ToolRegistry


class TestRecalcCacheWriter(unittest.TestCase):
    """Kiểm định bộ ghi cache lxml và động cơ tính toán lại."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.mkdtemp()
        self.file_store = FileStore(base_dir=os.path.join(self.temp_dir, "filestore"))
        self.sample_xlsx = os.path.join(self.temp_dir, "calc_sample.xlsx")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CalcSheet"
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = "=SUM(A1:A2)"
        ws["B3"] = '=IF(A3>25,"Pass","Fail")'
        ws["C3"] = "=A3>25"
        wb.save(self.sample_xlsx)

    def tearDown(self) -> None:
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_inject_cached_values_lxml(self) -> None:
        """Kiểm tra tiêm cache <v> trực tiếp bằng lxml cho các kiểu dữ liệu khác nhau."""
        cached_data = {
            "CalcSheet": {
                "A3": 30,
                "B3": "Pass",
                "C3": True,
            }
        }

        updated_bytes = XlsxCacheWriter.inject_cached_values(
            xlsx_input=self.sample_xlsx,
            cached_values=cached_data,
            ensure_full_calc=True,
        )
        self.assertIsNotNone(updated_bytes)

        # Kiểm tra nội dung XML bên trong ZIP
        with zipfile.ZipFile(io.BytesIO(updated_bytes), "r") as z:
            # 1. Kiểm tra xl/workbook.xml có calcPr fullCalcOnLoad="1"
            wb_xml = z.read("xl/workbook.xml")
            wb_tree = etree.fromstring(wb_xml)
            calc_nodes = wb_tree.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}calcPr")
            self.assertTrue(len(calc_nodes) > 0)
            self.assertEqual(calc_nodes[0].attrib.get("fullCalcOnLoad"), "1")

            # 2. Kiểm tra sheet1.xml
            sheet_xml = z.read("xl/worksheets/sheet1.xml")
            s_tree = etree.fromstring(sheet_xml)
            ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

            # Ô A3: số 30, giữ nguyên formula
            cell_a3 = s_tree.xpath('.//s:c[@r="A3"]', namespaces=ns)[0]
            f_a3 = cell_a3.find("s:f", ns)
            v_a3 = cell_a3.find("s:v", ns)
            self.assertEqual(f_a3.text, "SUM(A1:A2)")
            self.assertEqual(v_a3.text, "30")

            # Ô B3: chuỗi Pass, t="str"
            cell_b3 = s_tree.xpath('.//s:c[@r="B3"]', namespaces=ns)[0]
            self.assertEqual(cell_b3.attrib.get("t"), "str")
            self.assertEqual(cell_b3.find("s:v", ns).text, "Pass")

            # Ô C3: boolean True, t="b", v="1"
            cell_c3 = s_tree.xpath('.//s:c[@r="C3"]', namespaces=ns)[0]
            self.assertEqual(cell_c3.attrib.get("t"), "b")
            self.assertEqual(cell_c3.find("s:v", ns).text, "1")

    def test_recalc_engine_cache_writer_strategy(self) -> None:
        """Kiểm tra điều phối phương thức cache_writer trong RecalcEngine."""
        engine = RecalcEngine()
        res = engine.recalculate(
            xlsx_path=self.sample_xlsx,
            method="cache_writer",
            cached_values={"CalcSheet": {"A3": 30}},
        )
        self.assertTrue(res.success)
        self.assertEqual(res.method_used, "cache_writer")

    def test_mcp_tool_recalc_integration(self) -> None:
        """Kiểm tra tích hợp công cụ MCP xlsx.recalc trong ToolRegistry."""
        registry = ToolRegistry(file_store=self.file_store)
        register_xlsx_recalc_tools(registry)

        out_path = os.path.join(self.temp_dir, "recalculated.xlsx")
        res = registry.execute(
            "xlsx.recalc",
            {
                "file_ref": self.sample_xlsx,
                "method": "cache_writer",
                "cached_values": {"CalcSheet": {"A3": 30}},
                "output_path": out_path,
            },
        )
        self.assertTrue(res.success)
        self.assertIsNotNone(res.file_ref)
        self.assertTrue(os.path.exists(out_path))
        self.assertIn("RECALCULATED", res.guarantees_applied)


if __name__ == "__main__":
    unittest.main()
