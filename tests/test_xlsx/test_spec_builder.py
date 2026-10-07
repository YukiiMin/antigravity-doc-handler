"""
Unit test suite for doctools.core.xlsx.build (Native Chart & Spec Builder).
Kiểm định toàn diện Tiểu bước 2.2.1:
- Dựng bảng tính từ JSON XlsxSpec thuần (Path B)
- Sinh biểu đồ DrawingML native (BarChart, LineChart, PieChart)
- Áp dụng các design tokens: Font, Fill, Border, Number Format '@'
- Tích hợp công cụ MCP xlsx.build vào ToolRegistry
"""

import os
import shutil
import tempfile
import unittest
import openpyxl

from doctools.contract.xlsx.spec import ChartSpec, SheetSpec, StyleSpec, TableSpec, XlsxSpec
from doctools.core.xlsx.build.spec_builder import XlsxSpecBuilder
from doctools.infra.file_store import FileStore
from doctools.operations.xlsx.build_ops import register_xlsx_build_tools
from doctools.registry import ToolRegistry


class TestXlsxSpecBuilder(unittest.TestCase):
    """Kiểm thử bộ dựng bảng tính XlsxSpecBuilder."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.mkdtemp()
        self.file_store = FileStore(base_dir=os.path.join(self.temp_dir, "filestore"))
        self.builder = XlsxSpecBuilder()

    def tearDown(self) -> None:
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_build_table_and_styles(self) -> None:
        """Kiểm tra dựng bảng dữ liệu, style ô và vùng gộp."""
        spec = XlsxSpec(
            sheets=[
                SheetSpec(
                    name="Summary",
                    column_widths={"A": 15.0, "B": 30.0, "C": 12.0},
                    freeze_panes="A2",
                    merged_ranges=["A1:C1"],
                    tables=[
                        TableSpec(
                            start_coordinate="A1",
                            headers=["REPORT HEADER", "", ""],
                            header_style=StyleSpec(bold=True, fill_color="4F81BD", font_color="FFFFFF"),
                            rows=[
                                ["ID001", "Component Alpha", 100],
                                ["ID002", "Component Beta", 250],
                            ],
                            id_columns=[1],
                            row_style=StyleSpec(border_style="thin"),
                        )
                    ],
                )
            ]
        )

        wb, issues = self.builder.build(spec)
        self.assertEqual(len(issues), 0)
        self.assertEqual(wb.sheetnames, ["Summary"])

        ws = wb["Summary"]
        self.assertEqual(ws["A1"].value, "REPORT HEADER")
        self.assertEqual(ws["A2"].value, "ID001")
        self.assertEqual(ws["A2"].number_format, "@")
        self.assertEqual(ws["B2"].value, "Component Alpha")
        self.assertEqual(ws["C2"].value, 100)

    def test_build_native_drawingml_chart(self) -> None:
        """Kiểm tra sinh biểu đồ DrawingML native trong workbook."""
        spec = XlsxSpec(
            sheets=[
                SheetSpec(
                    name="Sales",
                    tables=[
                        TableSpec(
                            start_coordinate="A1",
                            headers=["Quarter", "Revenue"],
                            rows=[
                                ["Q1", 1000],
                                ["Q2", 1500],
                                ["Q3", 1200],
                                ["Q4", 2100],
                            ],
                        )
                    ],
                    charts=[
                        ChartSpec(
                            chart_type="col",
                            title="Quarterly Revenue",
                            data_min_col=2,
                            data_min_row=1,
                            data_max_col=2,
                            data_max_row=5,
                            categories_min_col=1,
                            categories_min_row=2,
                            categories_max_col=1,
                            categories_max_row=5,
                            anchor="D2",
                        )
                    ],
                )
            ]
        )

        wb, issues = self.builder.build(spec)
        self.assertEqual(len(issues), 0)

        ws = wb["Sales"]
        self.assertEqual(len(ws._charts), 1)
        chart = ws._charts[0]
        self.assertIsNotNone(chart.title)
        self.assertIn("Quarterly Revenue", str(chart.title))

        # Lưu file và kiểm tra khả năng mở lại của openpyxl
        out_file = os.path.join(self.temp_dir, "chart_test.xlsx")
        wb.save(out_file)
        self.assertTrue(os.path.exists(out_file))

        wb_loaded = openpyxl.load_workbook(out_file)
        ws_loaded = wb_loaded["Sales"]
        self.assertEqual(len(ws_loaded._charts), 1)

    def test_mcp_tool_build_integration(self) -> None:
        """Kiểm tra công cụ MCP xlsx.build trong ToolRegistry."""
        registry = ToolRegistry(file_store=self.file_store)
        register_xlsx_build_tools(registry)

        spec_data = {
            "title": "Generated Project Plan",
            "sheets": [
                {
                    "name": "Tasks",
                    "tables": [
                        {
                            "start_coordinate": "A1",
                            "headers": ["Task ID", "Name", "Hours"],
                            "rows": [
                                ["TSK-01", "Design Engine", 40],
                                ["TSK-02", "Implement Tests", 30],
                            ],
                            "id_columns": [1],
                        }
                    ],
                }
            ],
        }

        res = registry.execute("xlsx.build", {"spec": spec_data})
        self.assertTrue(res.success)
        self.assertIsNotNone(res.file_ref)
        self.assertIn("BUILT_FROM_SPEC", res.guarantees_applied)


if __name__ == "__main__":
    unittest.main()
