"""
tests.test_xlsx.test_write_catalog — Unit tests for Phase 2.4.5: Extended Write Catalog (FR-35, TC-55).
Verifies 100% round-trip fidelity for all extended write operations:
- set_format round-trip
- set_validation round-trip
- set_conditional_format round-trip
- copy_sheet round-trip
- ToolRegistry registration
"""

from pathlib import Path
import tempfile
import unittest

import openpyxl

from doctools.operations.xlsx.write_catalog import (
    register_xlsx_write_catalog_tools,
    xlsx_copy_sheet,
    xlsx_set_conditional_format,
    xlsx_set_format,
    xlsx_set_validation,
)
from doctools.registry import ToolRegistry


class TestWriteCatalog(unittest.TestCase):
    """Round-trip verification for Extended Write Catalog operations."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_file = Path(self.temp_dir.name) / "test_catalog.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "MainSheet"
        ws["A1"] = "ID"
        ws["B1"] = "Status"
        ws["A2"] = "001"
        ws["B2"] = "PASS"
        ws["C2"] = "=MainSheet!A2 + 10"  # Self-referencing formula

        wb.save(self.test_file)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_tc55_set_format_roundtrip(self) -> None:
        """TC-55: Verify set_format applies styles and persists 100% roundtrip."""
        res = xlsx_set_format(
            file_path=self.test_file,
            sheet="MainSheet",
            range="A2:A2",
            number_format="@",
            bold=True,
            fill_color="FFFFCC00",
            border_style="thin",
            alignment="center",
            wrap_text=True,
        )
        self.assertTrue(res.success)

        # Inspect back from saved file
        wb = openpyxl.load_workbook(self.test_file, data_only=False)
        cell = wb["MainSheet"]["A2"]
        self.assertEqual(cell.number_format, "@")
        self.assertTrue(cell.font.bold)
        self.assertIsNotNone(cell.fill.start_color.rgb)
        self.assertEqual(cell.border.bottom.style, "thin")
        self.assertEqual(cell.alignment.horizontal, "center")
        self.assertTrue(cell.alignment.wrap_text)

    def test_tc55_set_validation_roundtrip(self) -> None:
        """TC-55: Verify set_validation adds and persists Data Validation."""
        res = xlsx_set_validation(
            file_path=self.test_file,
            sheet="MainSheet",
            range="B2:B10",
            type="list",
            formula1='"PASS,FAIL,SKIP"',
        )
        self.assertTrue(res.success)

        # Inspect back
        wb = openpyxl.load_workbook(self.test_file, data_only=False)
        ws = wb["MainSheet"]
        self.assertEqual(len(ws.data_validations.dataValidation), 1)
        dv = ws.data_validations.dataValidation[0]
        self.assertEqual(dv.type, "list")
        self.assertEqual(dv.formula1, '"PASS,FAIL,SKIP"')
        self.assertIn("B2:B10", str(dv.sqref))

    def test_tc55_set_conditional_format_roundtrip(self) -> None:
        """TC-55: Verify set_conditional_format adds CF rule."""
        res = xlsx_set_conditional_format(
            file_path=self.test_file,
            sheet="MainSheet",
            range="B2:B10",
            operator="equal",
            formula=['"FAIL"'],
            fill_color="FFEE1111",
        )
        self.assertTrue(res.success)

        # Inspect back
        wb = openpyxl.load_workbook(self.test_file, data_only=False)
        ws = wb["MainSheet"]
        self.assertEqual(len(ws.conditional_formatting), 1)
        cf = list(ws.conditional_formatting)[0]
        self.assertIn("B2:B10", str(cf.sqref))

    def test_tc55_copy_sheet_roundtrip(self) -> None:
        """TC-55: Verify copy_sheet clones sheet with validation & formula rewiring."""
        # 1. Add validation to MainSheet first
        xlsx_set_validation(
            file_path=self.test_file,
            sheet="MainSheet",
            range="B2:B10",
            formula1='"PASS,FAIL"',
        )

        # 2. Clone sheet via copy_sheet
        res = xlsx_copy_sheet(
            file_path=self.test_file,
            source_sheet="MainSheet",
            target_title="ClonedSheet",
            rewire_self_refs=True,
        )
        self.assertTrue(res.success)

        # 3. Inspect back
        wb = openpyxl.load_workbook(self.test_file, data_only=False)
        self.assertIn("ClonedSheet", wb.sheetnames)
        cloned_ws = wb["ClonedSheet"]

        # Data Validation preserved
        self.assertEqual(len(cloned_ws.data_validations.dataValidation), 1)
        # Self-referencing formula rewired
        self.assertEqual(cloned_ws["C2"].value, "='ClonedSheet'!A2 + 10")

    def test_registry_registration(self) -> None:
        """Verify all write catalog tools are registered in ToolRegistry."""
        registry = ToolRegistry()
        register_xlsx_write_catalog_tools(registry)
        tool_names = [t["name"] for t in registry.list_tools()]

        self.assertIn("xlsx.set_format", tool_names)
        self.assertIn("xlsx.set_validation", tool_names)
        self.assertIn("xlsx.set_conditional_format", tool_names)
        self.assertIn("xlsx.copy_sheet", tool_names)


if __name__ == "__main__":
    unittest.main()
