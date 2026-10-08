"""
tests.test_xlsx.test_sheet_cloner — Unit tests for Phase 2.4.2: Sheet Cloner Parity Engine.
Covers:
- TC-59: Parity recovery of Data Validations and Conditional Formattings (remedy for EV-15.1)
- TC-41: Sheet clone operations, view attributes, and self-referencing formula rewiring
"""

import unittest
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

from doctools.core.xlsx.mutate.sheet_cloner import SheetCloner, clone_sheet_with_parity, rewire_formula_self_refs


class TestSheetCloner(unittest.TestCase):
    """Verifies SheetCloner parity preservation and formula rewiring."""

    def test_tc59_parity_cloner_preserves_validations_and_cf(self) -> None:
        """
        TC-59 (EV-15.1 Fix):
        Verify clone_sheet_with_parity preserves 100% of Data Validations and CF rules.
        """
        wb = openpyxl.Workbook()
        ws_src = wb.active
        ws_src.title = "Example"

        # 1. Add Data Validation
        dv = DataValidation(type="list", formula1='"O,X,N/A"', allow_blank=True)
        ws_src.add_data_validation(dv)
        dv.add("C3:C10")

        # 2. Add Conditional Formatting
        red_fill = PatternFill(start_color="FFEE1111", end_color="FFEE1111", fill_type="solid")
        rule = CellIsRule(operator="equal", formula=['"FAIL"'], fill=red_fill)
        ws_src.conditional_formatting.add("D3:D10", rule)

        # 3. Clone with parity
        ws_tgt = clone_sheet_with_parity(wb, "Example", "Handover (Staff)")

        # Verify DataValidation is 100% preserved
        self.assertEqual(len(ws_tgt.data_validations.dataValidation), 1)
        cloned_dv = ws_tgt.data_validations.dataValidation[0]
        self.assertEqual(cloned_dv.type, "list")
        self.assertEqual(cloned_dv.formula1, '"O,X,N/A"')
        self.assertIn("C3:C10", str(cloned_dv.sqref))

        # Verify Conditional Formatting is 100% preserved
        self.assertEqual(len(ws_tgt.conditional_formatting), 1)
        cloned_cf = list(ws_tgt.conditional_formatting)[0]
        self.assertIn("D3:D10", str(cloned_cf.sqref))

    def test_tc41_rewire_self_referencing_formulas(self) -> None:
        """
        TC-41: Verify formulas referencing the source sheet name are rewired to the target sheet.
        """
        wb = openpyxl.Workbook()
        ws_src = wb.active
        ws_src.title = "Return(Staff)"

        ws_src["A1"] = 100
        ws_src["B1"] = "='Return(Staff)'!A1 * 2"
        ws_src["C1"] = "=Return(Staff)!A1 + 50"
        ws_src["D1"] = "='OtherSheet'!A1 + 10"  # External reference, should not be touched!

        # Set freeze panes
        ws_src.freeze_panes = "B2"

        # Clone
        ws_tgt = clone_sheet_with_parity(wb, ws_src, "Resolve Ticket (Staff)")

        # Freeze panes preserved
        self.assertEqual(ws_tgt.freeze_panes, "B2")

        # Verify formula rewiring
        self.assertEqual(ws_tgt["B1"].value, "='Resolve Ticket (Staff)'!A1 * 2")
        self.assertEqual(ws_tgt["C1"].value, "='Resolve Ticket (Staff)'!A1 + 50")
        self.assertEqual(ws_tgt["D1"].value, "='OtherSheet'!A1 + 10")

    def test_rewire_formula_self_refs_unit(self) -> None:
        """Unit test for rewire_formula_self_refs function."""
        f1 = "=SUM('Source Sheet'!A1:A10)"
        self.assertEqual(rewire_formula_self_refs(f1, "Source Sheet", "New Sheet"), "=SUM('New Sheet'!A1:A10)")

        f2 = "=Source!B5 + 10"
        self.assertEqual(rewire_formula_self_refs(f2, "Source", "Target"), "='Target'!B5 + 10")

        f3 = "plain string"
        self.assertEqual(rewire_formula_self_refs(f3, "Source", "Target"), "plain string")


if __name__ == "__main__":
    unittest.main()
