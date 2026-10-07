"""
tests.test_xlsx.test_preflight
Unit tests for PreflightScanner and xlsx.preflight MCP tool operation (FR-01, FR-02).
"""

import io
import tempfile
import unittest
from pathlib import Path
import zipfile

import openpyxl
from openpyxl.drawing.image import Image
from PIL import Image as PILImage

from doctools.contract import Severity
from doctools.core.xlsx.preflight import preflight_scanner
from doctools.infra.file_store import FileStore
from doctools.operations.xlsx.preflight_ops import register_preflight_ops, xlsx_preflight
from doctools.registry import ToolRegistry


def _create_plain_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1["A1"] = 10
    ws1["A2"] = 20
    ws1["A3"] = "=SUM(A1:A2)"
    ws1.merge_cells("B1:C2")

    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Tổng kết"
    ws2["B1"] = "=Sheet1!A3"

    wb.save(path)


def _create_workbook_with_image(path: Path, tmp_dir: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cover"

    # Create dummy 10x10 PNG
    img_path = tmp_dir / "logo.png"
    pil_img = PILImage.new("RGB", (10, 10), color="blue")
    pil_img.save(img_path)

    img = Image(str(img_path))
    ws.add_image(img, "B2")
    wb.save(path)


class TestPreflightScanner(unittest.TestCase):
    """Tests PreflightScanner and Fidelity Tier assignment."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.plain_path = self.base_dir / "plain.xlsx"
        _create_plain_workbook(self.plain_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_plain_workbook_t1(self):
        res = preflight_scanner.scan(self.plain_path)
        inv = res.inventory

        self.assertEqual(inv.fidelity_tier, "T1")
        self.assertEqual(len(inv.sheets), 2)
        self.assertEqual(inv.total_formulas, 2)
        self.assertEqual(inv.total_merges, 1)
        self.assertFalse(inv.has_drawingml)
        self.assertFalse(inv.has_vba_macros)
        self.assertFalse(inv.has_pivot_tables)

        errors = [i for i in res.issues if i.severity == Severity.ERROR]
        self.assertEqual(len(errors), 0)

    def test_drawingml_workbook_t2(self):
        img_wb_path = self.base_dir / "with_img.xlsx"
        _create_workbook_with_image(img_wb_path, self.base_dir)

        res = preflight_scanner.scan(img_wb_path)
        inv = res.inventory

        self.assertEqual(inv.fidelity_tier, "T2")
        self.assertTrue(inv.has_drawingml or inv.has_images)

    def test_hidden_sheet_warning(self):
        wb = openpyxl.load_workbook(self.plain_path)
        wb["Summary"].sheet_state = "hidden"
        hidden_path = self.base_dir / "hidden.xlsx"
        wb.save(hidden_path)

        res = preflight_scanner.scan(hidden_path)
        self.assertTrue(any(i.code == "W-PKG-HIDDEN" for i in res.issues))

    def test_corrupt_package(self):
        bad_path = self.base_dir / "corrupt.xlsx"
        bad_path.write_bytes(b"NOT_A_ZIP")

        res = preflight_scanner.scan(bad_path)
        self.assertTrue(any(i.code == "E-PKG-UNREADABLE" for i in res.issues))


class TestPreflightOperations(unittest.TestCase):
    """Tests xlsx.preflight MCP operation and registry execution."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.file_store = FileStore(base_dir=self.base_dir / "store")
        self.registry = ToolRegistry(file_store=self.file_store)
        register_preflight_ops(self.registry)

        self.wb_path = self.base_dir / "op_wb.xlsx"
        _create_plain_workbook(self.wb_path)
        self.file_ref = self.file_store.store_file(self.wb_path, engine="xlsx")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_xlsx_preflight_operation(self):
        res = xlsx_preflight(self.file_ref, file_store=self.file_store)
        self.assertTrue(res.success)
        self.assertIsNotNone(res.stats)
        self.assertEqual(res.stats.extra["fidelity_tier"], "T1")
        self.assertIn("inventory", res.stats.extra)

    def test_registry_tool_execution(self):
        res = self.registry.execute("xlsx.preflight", {"file_ref": self.file_ref})
        self.assertTrue(res.success)
        self.assertIn("inventory", res.stats.extra)


if __name__ == "__main__":
    unittest.main()
