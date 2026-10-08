"""
tests.test_xlsx.test_extended_gates — Unit tests for Phase 2.4.3: Extended Quality Gates (UG-14..16).
Verifies:
- UG-14: Validation & Object Parity Gate (E-XLSX-UG14-VALIDATION-DROPPED)
- UG-15: Table Border Consistency Gate (W-XLSX-UG15-INCONSISTENT-BORDERS)
- UG-16: Formula Deterministic Anomaly Gate (W-XLSX-UG16-EMPTY-CELL-REF)
"""

import unittest
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Border, Side

from doctools.contract.issues import Severity
from doctools.gates.xlsx.extended_gates import (
    check_ug14_validation_parity,
    check_ug15_table_border_consistency,
    check_ug16_formula_deterministic_anomaly,
    run_extended_gates,
)
from doctools.gates.xlsx.universal_gates import XlsxUniversalGates


class TestExtendedGates(unittest.TestCase):
    """Unit tests for UG-14, UG-15, and UG-16."""

    def test_ug14_validation_dropped_detected(self) -> None:
        """UG-14: Emits E-XLSX-UG14-VALIDATION-DROPPED when a cloned/sibling sheet drops DV."""
        wb = openpyxl.Workbook()
        ws_ref = wb.active
        ws_ref.title = "Example"

        # Ref sheet has DataValidation
        dv = DataValidation(type="list", formula1='"O,X"')
        ws_ref.add_data_validation(dv)
        dv.add("C3:C10")

        # Mutated sheet has data but 0 validations (EV-15 defect)
        ws_child = wb.create_sheet("Handover (Staff)")
        for r in range(1, 10):
            ws_child.cell(row=r, column=1, value=f"Row {r}")
            ws_child.cell(row=r, column=3, value="O")

        issues = check_ug14_validation_parity(wb, reference_sheet="Example")
        ug14_issues = [i for i in issues if i.code == "E-XLSX-UG14-VALIDATION-DROPPED"]

        self.assertGreaterEqual(len(ug14_issues), 1)
        self.assertEqual(ug14_issues[0].severity, Severity.ERROR)
        self.assertIn("Handover (Staff)", ug14_issues[0].message)

    def test_ug15_table_border_consistency_detected(self) -> None:
        """UG-15: Emits W-XLSX-UG15-INCONSISTENT-BORDERS when cells in table lack borders."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Statistics"

        hair_side = Side(border_style="hair", color="000000")
        hair_border = Border(top=hair_side, bottom=hair_side, left=hair_side, right=hair_side)

        # Prototype row (row 2) has full borders
        for col in range(2, 5):
            ws.cell(row=2, column=col, value=f"Header {col}").border = hair_border

        # Lower data rows lack borders
        for row in range(3, 7):
            for col in range(2, 5):
                ws.cell(row=row, column=col, value=50)  # No border!

        issues = check_ug15_table_border_consistency(wb)
        ug15_issues = [i for i in issues if i.code == "W-XLSX-UG15-INCONSISTENT-BORDERS"]

        self.assertGreaterEqual(len(ug15_issues), 1)
        self.assertEqual(ug15_issues[0].severity, Severity.WARNING)
        self.assertIn("Statistics", ug15_issues[0].location.sheet)

    def test_ug16_formula_deterministic_anomaly_detected(self) -> None:
        """UG-16: Emits W-XLSX-UG16-EMPTY-CELL-REF when arithmetic formula references empty cell."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Return(Staff)"

        ws["Q7"] = 22
        # AA7 is empty
        ws["O8"] = "=Q7-AA7"

        issues = check_ug16_formula_deterministic_anomaly(wb)
        ug16_issues = [i for i in issues if i.code == "W-XLSX-UG16-EMPTY-CELL-REF"]

        self.assertGreaterEqual(len(ug16_issues), 1)
        self.assertEqual(ug16_issues[0].severity, Severity.WARNING)
        self.assertIn("AA7", ug16_issues[0].evidence["target"])

    def test_universal_gates_integration(self) -> None:
        """Verify XlsxUniversalGates executes UG-14..16 alongside UG-01..13."""
        wb = openpyxl.Workbook()
        ws_ref = wb.active
        ws_ref.title = "Example"
        dv = DataValidation(type="list", formula1='"O"')
        ws_ref.add_data_validation(dv)
        dv.add("A1")

        ws_data = wb.create_sheet("DataSheet")
        for r in range(1, 8):
            ws_data.cell(row=r, column=1, value="val")
        ws_data["B2"] = "=A1-Z99"  # Empty cell reference

        gates = XlsxUniversalGates(reference_sheet="Example")
        issues = gates.validate(wb)

        issue_codes = [i.code for i in issues]
        self.assertIn("E-XLSX-UG14-VALIDATION-DROPPED", issue_codes)
        self.assertIn("W-XLSX-UG16-EMPTY-CELL-REF", issue_codes)


if __name__ == "__main__":
    unittest.main()
