"""
tests.test_xlsx.test_inspect_tiered
Test suite for Phase 2.3a Tiered Inspection & Coverage Matrix (TC-47, TC-54, TC-58).
Verifies:
- TC-47: Tiered inspection (summary, structure, objects, cells)
- TC-54: Capability coverage matrix (Annex I 24 features and live audit)
- TC-58: Token efficiency and deterministic truncation flags
"""

import io
from pathlib import Path
import tempfile
import unittest

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

from doctools.contract.envelope import ResultEnvelope
from doctools.core.xlsx.inspect.coverage_analyzer import (
    ANNEX_I_BASELINE,
    CoverageAnalyzer,
    CoverageState,
)
from doctools.core.xlsx.inspect.sheet_inspector import TieredSheetInspector
from doctools.operations.xlsx.inspect_ops import (
    register_xlsx_inspect_tools,
    xlsx_coverage_report,
    xlsx_inspect,
)
from doctools.registry import ToolRegistry


class TestInspectTiered(unittest.TestCase):
    """Unit tests for Phase 2.3a: Tiered Inspection and Coverage Matrix."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_file = Path(self.temp_dir.name) / "test_sample.xlsx"

        # Create a rich test workbook
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "SummarySheet"
        ws1["A1"] = "Metric"
        ws1["B1"] = "Value"
        ws1["A2"] = "Total Count"
        ws1["B2"] = 150
        ws1["A3"] = "Pass Rate"
        ws1["B3"] = "=B2/200"
        ws1["A4"] = "Status"
        ws1["B4"] = '=IF(B3>0.5, "PASS", "FAIL")'
        ws1["A5"] = "#REF!"  # simulated error

        # Merge cells
        ws1.merge_cells("A10:B11")
        ws1["A10"] = "Merged Zone"

        # Freeze pane
        ws1.freeze_panes = "A2"

        # Data Validation
        dv = DataValidation(type="list", formula1='"OptionA,OptionB"')
        ws1.add_data_validation(dv)
        dv.add("B20")

        # Conditional Formatting
        red_fill = PatternFill(start_color="FFEE1111", end_color="FFEE1111", fill_type="solid")
        ws1.conditional_formatting.add("B2:B10", CellIsRule(operator="lessThan", formula=["100"], fill=red_fill))

        # Sheet 2: Sibling Sheet
        ws2 = wb.create_sheet(title="DetailSheet")
        ws2["A1"] = "ID"
        ws2["B1"] = "Name"
        ws2["C1"] = "Score"
        for r in range(2, 25):
            ws2[f"A{r}"] = f"ID_{r}"
            ws2[f"B{r}"] = f"Item {r}"
            ws2[f"C{r}"] = r * 10

        wb.save(self.test_file)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_tc47_tiered_inspect_summary(self) -> None:
        """TC-47: Verify level='summary' returns high-level metrics without cell dump."""
        res: ResultEnvelope = xlsx_inspect(self.test_file, level="summary")
        self.assertTrue(res.success)
        extra = res.stats.extra

        self.assertEqual(extra["level"], "summary")
        self.assertEqual(extra["sheets_count"], 2)
        self.assertIn("sheets", extra)
        self.assertEqual(len(extra["sheets"]), 2)

        # Check sheet metadata
        s1 = extra["sheets"][0]
        self.assertEqual(s1["name"], "SummarySheet")
        self.assertGreaterEqual(s1["formula_cells"], 2)
        self.assertEqual(s1["error_cells"], 1)

        # Check functions
        self.assertIn("IF", extra["unique_functions"])
        self.assertEqual(extra["truncated"], False)

    def test_tc47_tiered_inspect_structure(self) -> None:
        """TC-47: Verify level='structure' returns layout, anchors, merges, freeze panes."""
        res: ResultEnvelope = xlsx_inspect(self.test_file, level="structure")
        self.assertTrue(res.success)
        extra = res.stats.extra

        self.assertEqual(extra["level"], "structure")
        self.assertIn("sheets_structure", extra)
        s1_struct = extra["sheets_structure"][0]
        self.assertEqual(s1_struct["sheet"], "SummarySheet")
        self.assertEqual(s1_struct["freeze_panes"], "A2")
        self.assertGreaterEqual(s1_struct["merged_cells_count"], 1)
        self.assertIn("A10:B11", s1_struct["merged_ranges_sample"])

    def test_tc47_tiered_inspect_objects(self) -> None:
        """TC-47: Verify level='objects' inventories validations, CF, and pagination."""
        res: ResultEnvelope = xlsx_inspect(self.test_file, level="objects", page=1, max_items=10)
        self.assertTrue(res.success)
        extra = res.stats.extra

        self.assertEqual(extra["level"], "objects")
        self.assertIn("objects", extra)
        types_found = {obj["type"] for obj in extra["objects"]}
        self.assertIn("data_validation", types_found)
        self.assertIn("conditional_formatting", types_found)

        # Check coverage_state is annotated
        for obj in extra["objects"]:
            self.assertIn("coverage_state", obj)

    def test_tc47_tiered_inspect_cells(self) -> None:
        """TC-47: Verify level='cells' drills down into exact cell contents with pagination."""
        res: ResultEnvelope = xlsx_inspect(
            self.test_file, level="cells", sheet="DetailSheet", page=1, max_items=10
        )
        self.assertTrue(res.success)
        extra = res.stats.extra

        self.assertEqual(extra["level"], "cells")
        self.assertEqual(extra["sheet"], "DetailSheet")
        self.assertEqual(len(extra["cells"]), 10)
        self.assertTrue(extra["truncated"])
        self.assertEqual(extra["page"], 1)
        self.assertEqual(extra["next_page"], 2)

        # Check first cell
        c1 = extra["cells"][0]
        self.assertEqual(c1["coord"], "A1")
        self.assertEqual(c1["value"], "ID")

    def test_tc54_coverage_matrix_report(self) -> None:
        """TC-54: Verify xlsx.coverage_report returns Annex I baseline & live file audit."""
        # Baseline only
        res_baseline = xlsx_coverage_report()
        self.assertTrue(res_baseline.success)
        extra_b = res_baseline.stats.extra
        self.assertEqual(extra_b["total_feature_families"], len(ANNEX_I_BASELINE))
        keys = [f["key"] for f in extra_b["baseline_matrix"]]
        self.assertIn("values_datatypes", keys)
        self.assertIn("formulas_basic", keys)
        self.assertIn("merged_cells", keys)
        self.assertIn("data_validation", keys)

        # File-correlated audit
        res_file = xlsx_coverage_report(self.test_file)
        self.assertTrue(res_file.success)
        extra_f = res_file.stats.extra
        self.assertIn("workbook_audit", extra_f)
        detected = extra_f["workbook_audit"]["detected_features"]
        self.assertTrue(detected["data_validation"]["present"])
        self.assertTrue(detected["conditional_formatting"]["present"])
        self.assertTrue(detected["merged_cells"]["present"])
        self.assertFalse(detected["vba_macros"]["present"])

    def test_tc58_token_efficiency_and_truncation(self) -> None:
        """TC-58: Verify token compactness of summary and deterministic truncation."""
        res_summary = xlsx_inspect(self.test_file, level="summary")
        summary_extra = res_summary.stats.extra

        # Summary must NOT contain cell-level dumps
        self.assertNotIn("cells", summary_extra)
        self.assertFalse(summary_extra["truncated"])

        # Page-bounded cell inspection must strictly enforce max_items and truncated flag
        res_page1 = xlsx_inspect(self.test_file, level="cells", sheet="DetailSheet", page=1, max_items=5)
        self.assertEqual(len(res_page1.stats.extra["cells"]), 5)
        self.assertTrue(res_page1.stats.extra["truncated"])

        # Final page where remaining items fit
        total_cells = res_page1.stats.extra["total_cells"]
        res_all = xlsx_inspect(self.test_file, level="cells", sheet="DetailSheet", page=1, max_items=total_cells + 10)
        self.assertEqual(len(res_all.stats.extra["cells"]), total_cells)
        self.assertFalse(res_all.stats.extra["truncated"])
        self.assertIsNone(res_all.stats.extra["next_page"])

    def test_mcp_registry_binding(self) -> None:
        """Verify xlsx.inspect and xlsx.coverage_report are registered in ToolRegistry."""
        registry = ToolRegistry()
        register_xlsx_inspect_tools(registry)
        tool_names = [t["name"] for t in registry.list_tools()]
        self.assertIn("xlsx.inspect", tool_names)
        self.assertIn("xlsx.coverage_report", tool_names)


if __name__ == "__main__":
    unittest.main()
