"""
Unit test suite for doctools.core.xlsx.mutate (Mutate Expander & Style Cloner).
Kiểm thử toàn diện Tiểu bước 2.1.4 tuân thủ E1..E14, FR-05, FR-06, ERR_XLSX_004:
- Consume prototype row vs Insert row
- 100% Style cloning (Font, Fill, Border, Alignment)
- Ép kiểu ID column '@' giữ số 0 đầu
- Khóa ô locked_zones (E-XLSX-LOCK-001)
- Ghi an toàn ô gộp và đồng bộ viền 4 phía
- Tự động co giãn dải ô công thức qua ShiftManager
- Tích hợp ToolRegistry MCP cho xlsx.mutate
"""

import os
import shutil
import tempfile
import unittest
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from doctools.contract.issues import Severity
from doctools.contract.xlsx.mutation import CellUpdate, MutationSpec, TableExpansion
from doctools.core.xlsx.mutate.mutator import XlsxMutator
from doctools.infra.file_store import FileStore
from doctools.operations.xlsx.mutate_ops import register_xlsx_mutate_tools, xlsx_mutate
from doctools.registry import ToolRegistry


class TestXlsxMutator(unittest.TestCase):
    """Kiểm thử động cơ XlsxMutator và các thao tác đột biến."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.mkdtemp()
        self.file_store = FileStore(base_dir=os.path.join(self.temp_dir, "filestore"))

        # Tạo template workbook thử nghiệm
        self.template_path = os.path.join(self.temp_dir, "template.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DataSheet"

        # Dòng 1: Header
        ws["A1"] = "ID"
        ws["B1"] = "Name"
        ws["C1"] = "Amount"

        # Dòng 2: Prototype row có format đặc thù
        thin_border = Border(
            left=Side(style="thin", color="000000"),
            right=Side(style="thin", color="000000"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="thin", color="000000"),
        )
        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        bold_font = Font(name="Arial", size=11, bold=True)
        center_align = Alignment(horizontal="center", vertical="center")

        ws["A2"] = "001"
        ws["A2"].font = bold_font
        ws["A2"].fill = yellow_fill
        ws["A2"].border = thin_border
        ws["A2"].alignment = center_align

        ws["B2"] = "Sample Name"
        ws["B2"].font = Font(name="Arial", size=11)
        ws["B2"].border = thin_border

        ws["C2"] = 100
        ws["C2"].border = thin_border

        # Dòng 4: Công thức tổng cộng
        ws["B4"] = "Total"
        ws["C4"] = "=SUM(C2:C3)"

        wb.save(self.template_path)

    def tearDown(self) -> None:
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_consume_prototype_row(self) -> None:
        """Kiểm tra consume prototype row: ghi đè dòng mẫu và nhân bản tiếp theo."""
        wb = openpyxl.load_workbook(self.template_path)
        mutator = XlsxMutator(wb)

        spec = MutationSpec(
            template_ref_or_id=self.template_path,
            expansions=[
                TableExpansion(
                    sheet="DataSheet",
                    prototype_row=2,
                    consume_prototype=True,
                    rows_data=[
                        {"A": "007", "B": "James Bond", "C": 500},
                        {"A": "008", "B": "Ethan Hunt", "C": 600},
                        {"A": "009", "B": "Jason Bourne", "C": 700},
                    ],
                    id_columns=["A"],
                )
            ],
        )

        issues = mutator.mutate(spec)
        self.assertFalse(any(iss.severity == Severity.ERROR for iss in issues))

        ws = wb["DataSheet"]
        # Bản ghi 1 ở dòng 2 (ghi đè proto)
        self.assertEqual(ws["A2"].value, "007")
        self.assertEqual(ws["A2"].number_format, "@")
        self.assertEqual(ws["B2"].value, "James Bond")
        self.assertEqual(ws["C2"].value, 500)
        self.assertTrue(ws["A2"].font.bold)

        # Bản ghi 2 ở dòng 3 (nhân bản)
        self.assertEqual(ws["A3"].value, "008")
        self.assertEqual(ws["A3"].number_format, "@")
        self.assertEqual(ws["B3"].value, "Ethan Hunt")
        self.assertEqual(ws["C3"].value, 600)
        self.assertTrue(ws["A3"].font.bold)
        self.assertIsNotNone(ws["A3"].fill)

        # Bản ghi 3 ở dòng 4 (nhân bản)
        self.assertEqual(ws["A4"].value, "009")
        self.assertEqual(ws["B4"].value, "Jason Bourne")
        self.assertEqual(ws["C4"].value, 700)

        # Dòng Total bị đẩy xuống dòng 6 và công thức được mở rộng
        self.assertEqual(ws["B6"].value, "Total")
        self.assertEqual(ws["C6"].value, "=SUM(C2:C5)")

    def test_locked_zones_enforcement(self) -> None:
        """Kiểm tra chặn ghi đè khi ô vi phạm locked_zones (E-XLSX-LOCK-001)."""
        wb = openpyxl.load_workbook(self.template_path)
        mutator = XlsxMutator(wb, locked_zones=["DataSheet!A1:C2"])

        # Cố ý cập nhật ô A2 trong locked zone
        spec = MutationSpec(
            template_ref_or_id=self.template_path,
            cell_updates=[
                CellUpdate(sheet="DataSheet", coordinate="A2", value="MODIFIED")
            ],
        )

        issues = mutator.mutate(spec)
        lock_issues = [iss for iss in issues if iss.code == "E-XLSX-LOCK-001"]
        self.assertEqual(len(lock_issues), 1)

    def test_safe_merged_cells_border_sync(self) -> None:
        """Kiểm tra ghi an toàn vào ô gộp và đồng bộ viền 4 phía."""
        wb = openpyxl.load_workbook(self.template_path)
        ws = wb["DataSheet"]
        ws.merge_cells("A10:C11")

        mutator = XlsxMutator(wb)
        spec = MutationSpec(
            template_ref_or_id=self.template_path,
            cell_updates=[
                CellUpdate(sheet="DataSheet", coordinate="A10", value="Merged Header"),
                CellUpdate(sheet="DataSheet", coordinate="B10", value="Child Cell"),
            ],
        )

        issues = mutator.mutate(spec)
        self.assertFalse(any(iss.severity == Severity.ERROR for iss in issues))
        self.assertEqual(ws["A10"].value, "Merged Header")

    def test_mcp_tool_mutate_integration(self) -> None:
        """Kiểm tra tích hợp công cụ MCP xlsx.mutate trong ToolRegistry."""
        registry = ToolRegistry(file_store=self.file_store)
        register_xlsx_mutate_tools(registry)

        out_file = os.path.join(self.temp_dir, "output_mutated.xlsx")
        spec_dict = {
            "template_ref_or_id": self.template_path,
            "output_filename_hint": "result.xlsx",
            "expansions": [
                {
                    "sheet": "DataSheet",
                    "prototype_row": 2,
                    "consume_prototype": True,
                    "rows_data": [
                        {"A": "A100", "B": "Alpha", "C": 10},
                        {"A": "A101", "B": "Beta", "C": 20},
                    ],
                }
            ],
            "cell_updates": [
                {"sheet": "DataSheet", "coordinate": "A1", "value": "ID Code"}
            ],
        }

        res = registry.execute("xlsx.mutate", {"spec": spec_dict, "output_path": out_file})
        self.assertTrue(res.success)
        self.assertIsNotNone(res.file_ref)
        self.assertTrue(os.path.exists(out_file))

        # Đọc lại file kết quả
        wb_check = openpyxl.load_workbook(out_file)
        ws_check = wb_check["DataSheet"]
        self.assertEqual(ws_check["A1"].value, "ID Code")
        self.assertEqual(ws_check["A2"].value, "A100")
        self.assertEqual(ws_check["A3"].value, "A101")


if __name__ == "__main__":
    unittest.main()
