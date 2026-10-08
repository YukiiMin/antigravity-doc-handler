"""
tests.test_xlsx.test_formula_analyzer
Test suite for Phase 2.3b: Formula Analyzer (TC-48, TC-49, TC-61).
Verifies:
- TC-48: Canonical R1C1 grouping and pattern-break detection (W-FORMULA-PATTERN-BREAK)
- TC-49: Circular dependency detection (W-FORMULA-CIRCULAR)
- TC-61: Deterministic empty-cell reference warning (W-FORMULA-EMPTY-CELL-REF)
"""

from pathlib import Path
import tempfile
import unittest

import openpyxl

from doctools.contract.envelope import ResultEnvelope
from doctools.core.xlsx.inspect.formula_profiler import (
    FormulaProfiler,
    a1_to_r1c1_coordinate,
    normalize_formula_to_r1c1,
)
from doctools.operations.xlsx.inspect_ops import (
    register_xlsx_inspect_tools,
    xlsx_analyze_formulas,
)
from doctools.registry import ToolRegistry


class TestFormulaAnalyzer(unittest.TestCase):
    """Unit tests for Phase 2.3b: Formula Analyzer."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_file = Path(self.temp_dir.name) / "test_formulas.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "MathSheet"

        # 1. Column B: 50 rows of =A{r}+C{r}, but row 25 is =A25*C25 (Pattern break)
        for r in range(1, 51):
            ws[f"A{r}"] = r * 10
            ws[f"C{r}"] = 5
            if r == 25:
                ws[f"B{r}"] = f"=A{r}*C{r}"  # Pattern Break!
            else:
                ws[f"B{r}"] = f"=A{r}+C{r}"

        # 2. EV-15 Bug Simulation: Empty-cell reference
        # D1 = A1 - AA7 (where AA7 is empty / None)
        ws["D1"] = "=A1-AA7"

        # 3. Sheet 2: Circular Dependency Loop
        ws2 = wb.create_sheet(title="CycleSheet")
        ws2["A1"] = "=B1+1"
        ws2["B1"] = "=C1*2"
        ws2["C1"] = "=A1-5"  # Cycle: A1 -> B1 -> C1 -> A1

        wb.save(self.test_file)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_r1c1_coordinate_conversion(self) -> None:
        """Verify individual coordinate conversion to relative R1C1."""
        self.assertEqual(a1_to_r1c1_coordinate("A2", base_row=2, base_col=2), "RC[-1]")
        self.assertEqual(a1_to_r1c1_coordinate("C2", base_row=2, base_col=2), "RC[1]")
        self.assertEqual(a1_to_r1c1_coordinate("B3", base_row=2, base_col=2), "R[1]C")
        self.assertEqual(a1_to_r1c1_coordinate("$A$1", base_row=5, base_col=5), "R1C1")

    def test_tc48_r1c1_pattern_grouping_and_break(self) -> None:
        """TC-48: Verify R1C1 grouping collapses repetitive formulas and flags breaks."""
        res: ResultEnvelope = xlsx_analyze_formulas(self.test_file, sheet="MathSheet")
        self.assertTrue(res.success)
        extra = res.stats.extra

        # Dominant pattern in column B should collapse 49 rows into 1 group
        col_b_groups = [g for g in extra["pattern_groups"] if g["column"] == "B"]
        self.assertEqual(len(col_b_groups), 1)
        self.assertEqual(col_b_groups[0]["count"], 49)
        self.assertEqual(col_b_groups[0]["r1c1_pattern"], "=RC[-1]+RC[1]")

        # Pattern break at B25 should be detected
        breaks = [b for b in extra["pattern_breaks"] if b["coord"] == "B25"]
        self.assertEqual(len(breaks), 1)
        self.assertEqual(breaks[0]["code"], "W-FORMULA-PATTERN-BREAK")
        self.assertIn("A25*C25", breaks[0]["formula"])

    def test_tc49_circular_reference_detection(self) -> None:
        """TC-49: Verify circular dependency loops are flagged with W-FORMULA-CIRCULAR."""
        res: ResultEnvelope = xlsx_analyze_formulas(self.test_file, sheet="CycleSheet")
        self.assertTrue(res.success)
        extra = res.stats.extra

        circulars = extra["circular_references"]
        self.assertGreaterEqual(len(circulars), 1)
        cycle_info = circulars[0]
        self.assertEqual(cycle_info["code"], "W-FORMULA-CIRCULAR")
        self.assertIn("CycleSheet", cycle_info["cycle"])

    def test_tc61_empty_cell_reference_warning(self) -> None:
        """TC-61 (EV-15): Verify formula referencing an empty cell raises deterministic warning."""
        res: ResultEnvelope = xlsx_analyze_formulas(self.test_file, sheet="MathSheet")
        self.assertTrue(res.success)
        extra = res.stats.extra

        empty_refs = extra["empty_cell_references"]
        self.assertGreaterEqual(len(empty_refs), 1)
        found_aa7 = any("AA7" in er["target"] for er in empty_refs)
        self.assertTrue(found_aa7, "Expected warning for empty cell reference AA7")
        self.assertEqual(empty_refs[0]["code"], "W-FORMULA-EMPTY-CELL-REF")

    def test_tool_registry_registration(self) -> None:
        """Verify xlsx.analyze_formulas is registered into ToolRegistry."""
        registry = ToolRegistry()
        register_xlsx_inspect_tools(registry)
        tool_names = [t["name"] for t in registry.list_tools()]
        self.assertIn("xlsx.analyze_formulas", tool_names)


if __name__ == "__main__":
    unittest.main()
