"""
tests.test_docx.test_docspec_builder
Unit tests for Path B DocSpecBuilder and docx.build_document MCP operation.
"""

import tempfile
import unittest
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT

from doctools.contract import Severity
from doctools.contract.docx.docspec import (
    CalloutBlock,
    DocSpec,
    HeadingBlock,
    ListBlock,
    PageBreakBlock,
    PageSetupSpec,
    ParagraphBlock,
    RunSpec,
    TableBlock,
)
from doctools.core.docx.build import docspec_builder
from doctools.infra.file_store import FileStore
from doctools.operations.docx.build_ops import docx_build_document, register_build_ops
from doctools.registry import ToolRegistry

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class TestDocSpecBuilder(unittest.TestCase):
    """Tests Path B deterministic DocSpec generation with OOXML guards."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_build_comprehensive_document(self):
        spec = DocSpec(
            docspec_version="1.0",
            title="Báo cáo kỹ thuật Antigravity",
            page_setup=PageSetupSpec(orientation="portrait", margin_top_cm=2.0),
            blocks=[
                HeadingBlock(level=1, text="1. Giới thiệu tổng quan"),
                ParagraphBlock(
                    runs=[
                        RunSpec(text="Hệ thống AI-Native ", bold=True),
                        RunSpec(text="vận hành tự động.", italic=True, color="003366"),
                    ]
                ),
                ListBlock(items=["Tính năng A", "Tính năng B", "Tính năng C"], ordered=False),
                HeadingBlock(level=2, text="1.1 Dữ liệu nghiệm thu"),
                TableBlock(
                    headers=["Mã", "Hạng mục", "Trạng thái"],
                    rows=[
                        ["IT-01", "Core Engine", "Đạt"],
                        ["IT-02", "Quality Gates", "Đạt"],
                    ],
                    col_widths=[2.5, 8.0, 3.5],
                ),
                CalloutBlock(
                    kind="note",
                    title="Lưu ý triển khai",
                    text="Cần kiểm định lại các chốt chặn OOXML trước khi phát hành.",
                ),
                PageBreakBlock(),
                HeadingBlock(level=1, text="2. Kết luận"),
                ParagraphBlock(text="Dự án hoàn thành đúng tiến độ đề ra."),
            ],
        )

        res = docspec_builder.build(spec)
        self.assertTrue(len(res.docx_bytes) > 0)
        self.assertEqual(len([i for i in res.issues if i.severity == Severity.ERROR]), 0)
        self.assertIn("docspec_build", res.guarantees_applied)
        self.assertIn("guards:cantSplit", res.guarantees_applied)
        self.assertIn("guards:tblHeader", res.guarantees_applied)

        # Inspect generated Word document structure
        out_file = self.base_dir / "generated_docspec.docx"
        out_file.write_bytes(res.docx_bytes)
        doc = docx.Document(out_file)

        # Check title and headings
        full_text = " ".join([p.text for p in doc.paragraphs])
        self.assertIn("Báo cáo kỹ thuật Antigravity", full_text)
        self.assertIn("1. Giới thiệu tổng quan", full_text)
        self.assertIn("Tính năng A", full_text)

        # Check table OOXML guards
        tbl = doc.tables[0]
        self.assertEqual(len(tbl.rows), 3)  # 1 header + 2 data rows
        # Check header
        trPr_0 = tbl.rows[0]._tr.find(f"{_W}trPr")
        self.assertIsNotNone(trPr_0.find(f"{_W}tblHeader"))

        # Check cantSplit and last_p on every row
        for row in tbl.rows:
            trPr = row._tr.find(f"{_W}trPr")
            self.assertIsNotNone(trPr.find(f"{_W}cantSplit"))
            for cell in row.cells:
                self.assertEqual(cell._tc[-1].tag, f"{_W}p")

        # Check callout box (second table in doc)
        self.assertTrue(len(doc.tables) >= 2)
        callout_tbl = doc.tables[1]
        self.assertIn("Lưu ý triển khai", callout_tbl.rows[0].cells[0].paragraphs[0].text)

    def test_docspec_validation_errors(self):
        invalid_spec_dict = {
            "docspec_version": "1.0",
            "blocks": [
                {
                    "type": "heading",
                    "level": 99,  # exceeds le=6
                    "text": "Invalid Heading",
                },
                {
                    "type": "unknown_block_type",
                },
            ],
        }
        res = docspec_builder.build(invalid_spec_dict)
        self.assertEqual(len(res.docx_bytes), 0)
        self.assertTrue(len(res.issues) > 0)
        self.assertTrue(any(i.code == "E-DOCX-DOCSPEC-INVALID" for i in res.issues))

    def test_docspec_landscape_page_setup(self):
        spec = DocSpec(
            page_setup=PageSetupSpec(orientation="landscape", margin_left_cm=3.0),
            blocks=[
                HeadingBlock(level=1, text="Bảng biểu khổ ngang"),
            ],
        )
        res = docspec_builder.build(spec)
        self.assertTrue(len(res.docx_bytes) > 0)

        out_file = self.base_dir / "landscape.docx"
        out_file.write_bytes(res.docx_bytes)
        doc = docx.Document(out_file)
        section = doc.sections[0]
        self.assertEqual(section.orientation, WD_ORIENT.LANDSCAPE)


class TestDocSpecOperations(unittest.TestCase):
    """Tests MCP docx.build_document dispatcher with source_type='docspec'."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.file_store = FileStore(base_dir=self.base_dir / "store")
        self.registry = ToolRegistry(file_store=self.file_store)
        register_build_ops(self.registry)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_build_document_docspec_dispatcher(self):
        docspec_data = {
            "docspec_version": "1.0",
            "title": "Tài liệu sinh qua MCP dispatcher",
            "blocks": [
                {"type": "heading", "level": 1, "text": "Chương 1"},
                {"type": "paragraph", "text": "Nội dung văn bản được tạo thành công."},
            ],
        }
        envelope = docx_build_document(
            source_type="docspec",
            docspec=docspec_data,
            file_store=self.file_store,
        )
        self.assertTrue(envelope.success)
        self.assertIsNotNone(envelope.file_ref)
        self.assertIn("docspec_build", envelope.guarantees_applied)

        resolved_path = self.file_store.resolve(envelope.file_ref)
        self.assertTrue(resolved_path.is_file())


if __name__ == "__main__":
    unittest.main()
