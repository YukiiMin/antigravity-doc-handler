"""
Unit test suite for docx template operations (doctools.operations.docx.template_ops).
Kiểm định việc tích hợp docx.lint_template và docx.normalize_template vào ToolRegistry,
đảm bảo hợp đồng dữ liệu ResultEnvelope, FileRef và AuditLogger.
"""

from __future__ import annotations
import shutil
import tempfile
import unittest
from pathlib import Path
import docx

from doctools.contract import ResultEnvelope
from doctools.core.docx.package_io import PackageWriter
from doctools.infra.audit import AuditLogger
from doctools.infra.file_store import FileStore
from doctools.operations.docx.template_ops import (
    docx_lint_template,
    docx_normalize_template,
    register_template_ops,
)
from doctools.registry import ToolRegistry


class TestTemplateOps(unittest.TestCase):
    """Kiểm tra hoạt động của MCP operations docx.lint_template và docx.normalize_template."""

    def setUp(self) -> None:
        self.test_dir = Path(tempfile.mkdtemp(prefix="test_ops_"))
        self.file_store = FileStore(base_dir=self.test_dir / "filestore")
        self.audit_logger = AuditLogger(log_file=self.test_dir / "audit.jsonl")
        self.registry = ToolRegistry(file_store=self.file_store, audit_logger=self.audit_logger)
        register_template_ops(self.registry)

    def tearDown(self) -> None:
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_registry_registration(self) -> None:
        self.assertTrue(self.registry.is_registered("docx.lint_template"))
        self.assertTrue(self.registry.is_registered("docx.normalize_template"))

    def test_execute_lint_template_via_registry(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("Kính gửi ông/bà {{ recipient_name }}")
        tpl_path = self.test_dir / "template.docx"
        PackageWriter.write_document_to_path(doc, tpl_path)

        res: ResultEnvelope = self.registry.execute("docx.lint_template", {"template_ref": str(tpl_path)})
        self.assertTrue(res.success)
        self.assertIn("read_only_template_inspection", res.guarantees_applied)
        self.assertIsNotNone(res.stats)
        self.assertIn("recipient_name", res.stats.extra.get("variables", []))

    def test_execute_normalize_template_via_registry(self) -> None:
        doc = docx.Document()
        p = doc.add_paragraph()
        p.add_run("{{ ")
        p.add_run("order_id")
        p.add_run(" }}")
        tpl_path = self.test_dir / "split_tpl.docx"
        PackageWriter.write_document_to_path(doc, tpl_path)

        res: ResultEnvelope = self.registry.execute("docx.normalize_template", {"template_ref": str(tpl_path)})
        self.assertTrue(res.success)
        self.assertIsNotNone(res.file_ref)
        self.assertTrue(res.file_ref.uri.startswith("resource://docx/files/"))
        self.assertIn("runs_consolidated", res.guarantees_applied)


if __name__ == "__main__":
    unittest.main()
