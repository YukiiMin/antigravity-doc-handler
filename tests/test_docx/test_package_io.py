"""
Unit test suite for doctools.core.docx.package_io (PackageReader and PackageWriter).
Kiểm định an toàn bảo mật khi đọc gói ZIP ECMA-376 và ghi tệp nguyên tử.
"""

from __future__ import annotations
import shutil
import tempfile
import unittest
from pathlib import Path
import docx

from doctools.core.docx.package_io import (
    PackageReader,
    safe_read_package,
    PackageWriter,
    safe_write_package,
)


class TestPackageIO(unittest.TestCase):
    """Kiểm tra đọc và ghi gói tệp DOCX an toàn."""

    def setUp(self) -> None:
        self.test_dir = Path(tempfile.mkdtemp(prefix="test_pkg_io_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_roundtrip_document_write_and_read(self) -> None:
        doc = docx.Document()
        doc.add_heading("Docx Engine Test Document", level=1)
        doc.add_paragraph("This is a paragraph verifying atomic package writing.")

        out_path = self.test_dir / "sample_output.docx"
        PackageWriter.write_document_to_path(doc, out_path)
        self.assertTrue(out_path.is_file())

        reader = safe_read_package(out_path)
        part_names = reader.get_part_names()
        self.assertIn("word/document.xml", part_names)
        self.assertIn("[Content_Types].xml", part_names)

        doc_xml = reader.get_part_xml("word/document.xml")
        self.assertIsNotNone(doc_xml)

        loaded_doc = reader.to_docx_document()
        self.assertEqual(len(loaded_doc.paragraphs), 2)
        self.assertEqual(loaded_doc.paragraphs[0].text, "Docx Engine Test Document")

    def test_parts_dictionary_write(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("Testing part dictionary output.")
        bytes_data = PackageWriter.write_document_to_bytes(doc)

        reader = safe_read_package(bytes_data)
        parts = {name: reader.get_part_bytes(name) for name in reader.get_part_names() if reader.get_part_bytes(name)}

        out_path = self.test_dir / "repackaged.docx"
        PackageWriter.write_parts_to_path(parts, out_path)
        self.assertTrue(out_path.is_file())

        re_reader = safe_read_package(out_path)
        self.assertIn("word/document.xml", re_reader.get_part_names())


if __name__ == "__main__":
    unittest.main()
