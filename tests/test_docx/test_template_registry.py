"""
tests.test_docx.test_template_registry
Unit and integration tests for Template Manifest schema, Template Registry, and MCP operations.
"""

import hashlib
import tempfile
import unittest
from pathlib import Path

import docx
from doctools.contract import Engine, FileRef, Severity
from doctools.contract.docx.manifest import LayoutPolicy, TemplateManifest, VariableDef
from doctools.core.docx.package_io import PackageWriter


def _create_sample_docx(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Template Sample", level=1)
    doc.add_paragraph("Sample content for template registry test.")
    PackageWriter.write_document_to_path(doc, path)
from doctools.core.docx.template.manifest_parser import (
    ManifestParseError,
    parse_manifest,
    validate_manifest_against_template,
)
from doctools.core.docx.template.template_registry import (
    TemplateRegistry,
    get_default_template_registry,
)
from doctools.infra.file_store import FileStore
from doctools.operations.docx.template_ops import (
    docx_get_template_manifest,
    docx_list_templates,
    docx_register_template,
    register_template_ops,
)
from doctools import ToolRegistry


class TestTemplateManifestSchema(unittest.TestCase):
    """Tests Pydantic models and parser for TemplateManifest."""

    def test_parse_manifest_from_dict(self):
        data = {
            "template_id": "contract_v1",
            "version": 1,
            "title": "Hợp đồng kinh tế",
            "variables": {
                "contract_no": {"type": "string", "immutable": True},
                "amount": {"type": "number", "immutable": True},
                "notes": {"type": "string", "immutable": False, "max_chars": 500},
            },
            "layout_policy_default": {
                "apply_guards": "apply",
                "field_update": "update_on_open",
            },
        }
        manifest = parse_manifest(data)
        self.assertEqual(manifest.template_id, "contract_v1")
        self.assertEqual(manifest.version, 1)
        self.assertTrue(manifest.variables["contract_no"].immutable)
        self.assertEqual(manifest.layout_policy_default.apply_guards, "apply")
        self.assertEqual(manifest.layout_policy_default.field_update, "update_on_open")

    def test_parse_manifest_from_yaml_string(self):
        yaml_content = """
        template_id: report_v2
        version: 2
        title: "Báo cáo tiến độ"
        variables:
          author:
            type: string
            immutable: false
          approved_by:
            type: string
            immutable: true
        """
        manifest = parse_manifest(yaml_content)
        self.assertEqual(manifest.template_id, "report_v2")
        self.assertEqual(manifest.version, 2)
        self.assertIn("author", manifest.variables)
        self.assertTrue(manifest.variables["approved_by"].immutable)

    def test_parse_manifest_from_json_string(self):
        json_content = '{"template_id": "json_tpl", "version": 3, "variables": {}}'
        manifest = parse_manifest(json_content)
        self.assertEqual(manifest.template_id, "json_tpl")
        self.assertEqual(manifest.version, 3)

    def test_parse_manifest_invalid_raises(self):
        with self.assertRaises(ManifestParseError):
            parse_manifest("invalid: [not a valid yaml structure")


class TestTemplateRegistry(unittest.TestCase):
    """Tests TemplateRegistry storage, validation, and versioning."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.file_store = FileStore(base_dir=self.base_dir / "store")
        self.registry = TemplateRegistry(file_store=self.file_store)

        # Create a valid sample docx package
        self.docx_path = self.base_dir / "sample.docx"
        _create_sample_docx(self.docx_path)
        self.docx_bytes = self.docx_path.read_bytes()
        self.docx_hash = hashlib.sha256(self.docx_bytes).hexdigest()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_register_template_success(self):
        manifest_data = {
            "template_id": "tpl_test",
            "version": 1,
            "variables": {
                "customer_name": {"type": "string"},
            },
        }
        record, issues = self.registry.register(self.docx_path, manifest_data)
        self.assertIsNotNone(record)
        self.assertEqual(len([i for i in issues if i.severity == Severity.ERROR]), 0)
        self.assertEqual(record.template_id, "tpl_test")
        self.assertEqual(record.version, 1)
        self.assertEqual(record.manifest.sha256, self.docx_hash)

        # Lookup
        found = self.registry.get("tpl_test")
        self.assertIsNotNone(found)
        self.assertEqual(found.version, 1)

    def test_register_template_sha256_mismatch(self):
        manifest_data = {
            "template_id": "tpl_mismatch",
            "sha256": "0000000000000000000000000000000000000000000000000000000000000000",
        }
        record, issues = self.registry.register(self.docx_path, manifest_data)
        self.assertIsNone(record)
        mismatch_issues = [i for i in issues if i.code == "E-DOCX-TPL-HASH-MISMATCH"]
        self.assertEqual(len(mismatch_issues), 1)

    def test_register_template_missing_file(self):
        missing_path = self.base_dir / "non_existent.docx"
        record, issues = self.registry.register(missing_path, {"template_id": "missing"})
        self.assertIsNone(record)
        self.assertTrue(any(i.code == "E-FILE-NOT-FOUND" for i in issues))

    def test_template_versioning_increment(self):
        record1, _ = self.registry.register(
            self.docx_path, {"template_id": "versioned_tpl", "version": 1}
        )
        self.assertEqual(record1.version, 1)

        # Register next version without explicit version number
        record2, _ = self.registry.register(
            self.docx_path, {"template_id": "versioned_tpl"}
        )
        self.assertEqual(record2.version, 2)

        # Check retrieval: latest is 2, specific is 1
        self.assertEqual(self.registry.get("versioned_tpl").version, 2)
        self.assertEqual(self.registry.get("versioned_tpl", version=1).version, 1)

    def test_list_templates(self):
        self.registry.register(self.docx_path, {"template_id": "alpha", "title": "Alpha Tpl"})
        self.registry.register(self.docx_path, {"template_id": "beta", "title": "Beta Tpl"})

        listing = self.registry.list_templates()
        self.assertEqual(len(listing), 2)
        ids = [item["template_id"] for item in listing]
        self.assertEqual(ids, ["alpha", "beta"])


class TestTemplateOperations(unittest.TestCase):
    """Tests MCP tool operation functions and Registry dispatch."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.file_store = FileStore(base_dir=self.base_dir / "store")
        self.template_registry = TemplateRegistry(file_store=self.file_store)

        self.docx_path = self.base_dir / "op_sample.docx"
        _create_sample_docx(self.docx_path)

        self.mcp_registry = ToolRegistry(file_store=self.file_store)
        register_template_ops(self.mcp_registry)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_operation_register_and_list(self):
        # Register via function
        res_reg = docx_register_template(
            template_ref=self.docx_path,
            manifest={"template_id": "op_test", "title": "Op Test Title"},
            template_registry=self.template_registry,
        )
        self.assertTrue(res_reg.success)
        self.assertIsNotNone(res_reg.file_ref)
        self.assertEqual(res_reg.stats.extra["template_id"], "op_test")

        # List via function
        res_list = docx_list_templates(template_registry=self.template_registry)
        self.assertTrue(res_list.success)
        self.assertEqual(res_list.stats.extra["count"], 1)

        # Get manifest via function
        res_manifest = docx_get_template_manifest("op_test", template_registry=self.template_registry)
        self.assertTrue(res_manifest.success)
        self.assertEqual(res_manifest.stats.extra["manifest"]["template_id"], "op_test")

    def test_operation_get_missing_manifest(self):
        res = docx_get_template_manifest("non_existent", template_registry=self.template_registry)
        self.assertFalse(res.success)
        self.assertTrue(any(i.code == "E-DOCX-TPL-NOT-FOUND" for i in res.diagnostics.errors))


if __name__ == "__main__":
    unittest.main()
