"""
tests.test_docx.test_field_manager
Unit tests for FieldManager and docx.update_fields MCP tool operation (FR-12).
"""

import io
import tempfile
import unittest
from pathlib import Path
import zipfile

import docx
from lxml import etree

from doctools.contract import Severity
from doctools.core.docx.fields import field_manager
from doctools.core.docx.package_io import PackageWriter
from doctools.infra.file_store import FileStore
from doctools.operations.docx.field_ops import docx_update_fields, register_field_ops
from doctools.registry import ToolRegistry

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _create_plain_doc(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Văn bản chuẩn", level=1)
    doc.add_paragraph("Không chứa trường động.")
    PackageWriter.write_document_to_path(doc, path)


def _create_doc_with_toc_field(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Tài liệu có Mục Lục", level=1)

    # Inject a w:fldSimple for TOC
    p = doc.add_paragraph()
    fldSimple = etree.Element(f"{_W}fldSimple", attrib={f"{_W}instr": 'TOC \\o "1-3" \\h \\z \\u'})
    p._p.append(fldSimple)

    doc.add_heading("Chương 1: Mở đầu", level=2)
    doc.add_paragraph("Nội dung chi tiết chương 1.")

    PackageWriter.write_document_to_path(doc, path)


class TestFieldManager(unittest.TestCase):
    """Tests dynamic field discovery and settings.xml injection."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.plain_path = self.base_dir / "plain.docx"
        self.toc_path = self.base_dir / "with_toc.docx"
        _create_plain_doc(self.plain_path)
        _create_doc_with_toc_field(self.toc_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_discover_plain_document(self):
        res = field_manager.discover_fields(self.plain_path)
        self.assertFalse(res.has_toc)
        self.assertFalse(res.requires_page_update)
        self.assertEqual(len(res.fields_found), 0)

    def test_discover_toc_document(self):
        res = field_manager.discover_fields(self.toc_path)
        self.assertTrue(res.has_toc)
        self.assertTrue(res.requires_page_update)
        self.assertEqual(len(res.fields_found), 1)
        self.assertEqual(res.fields_found[0]["type"], "TOC")

    def test_apply_update_policy_auto(self):
        raw_bytes = self.toc_path.read_bytes()
        out_bytes, issues, guarantees, stats = field_manager.apply_update_policy(
            raw_bytes=raw_bytes,
            field_update="auto",
            toc_mode="with_page_numbers",
        )
        self.assertTrue(stats["update_on_open_applied"])
        self.assertIn("fields:update_on_open_injected", guarantees)
        self.assertTrue(any(i.code == "W-FIELD-UPDATE-REQUIRED" for i in issues))

        # Inspect settings.xml inside zip
        with zipfile.ZipFile(io.BytesIO(out_bytes), "r") as zf:
            self.assertIn("word/settings.xml", zf.namelist())
            settings_xml = zf.read("word/settings.xml")
            root = etree.fromstring(settings_xml)
            uf = root.find(f"{_W}updateFields")
            self.assertIsNotNone(uf)
            self.assertEqual(uf.get(f"{_W}val"), "true")

    def test_apply_update_policy_none(self):
        raw_bytes = self.toc_path.read_bytes()
        out_bytes, issues, guarantees, stats = field_manager.apply_update_policy(
            raw_bytes=raw_bytes,
            field_update="none",
        )
        self.assertFalse(stats["update_on_open_applied"])
        self.assertTrue(any(i.code == "W-FIELD-UPDATE-REQUIRED" for i in issues))


class TestFieldOperations(unittest.TestCase):
    """Tests MCP docx.update_fields operation."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.file_store = FileStore(base_dir=self.base_dir / "store")
        self.registry = ToolRegistry(file_store=self.file_store)
        register_field_ops(self.registry)

        self.toc_path = self.base_dir / "op_toc.docx"
        _create_doc_with_toc_field(self.toc_path)
        self.file_ref = self.file_store.store_file(self.toc_path, engine="docx")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_docx_update_fields_operation(self):
        res = docx_update_fields(
            file_ref_or_path=self.file_ref,
            field_update="auto",
            file_store=self.file_store,
        )
        self.assertTrue(res.success)
        self.assertIsNotNone(res.file_ref)
        self.assertIn("fields:update_on_open_injected", res.guarantees_applied)

        # Verify output file
        resolved = self.file_store.resolve(res.file_ref)
        with zipfile.ZipFile(resolved, "r") as zf:
            settings_xml = zf.read("word/settings.xml")
            root = etree.fromstring(settings_xml)
            uf = root.find(f"{_W}updateFields")
            self.assertIsNotNone(uf)


if __name__ == "__main__":
    unittest.main()
