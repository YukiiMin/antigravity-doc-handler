"""
tests.test_docx.test_merger
Unit tests for DocxMerger and docx.merge MCP tool operation (FR-06).
"""

import tempfile
import unittest
from pathlib import Path

import docx

from doctools.contract import Severity
from doctools.core.docx.merge import docx_merger
from doctools.core.docx.package_io import PackageWriter
from doctools.core.docx.schema import schema_helper
from doctools.infra.file_store import FileStore
from doctools.operations.docx.merge_ops import docx_merge, register_merge_ops
from doctools.registry import ToolRegistry


def _create_doc(path: Path, title: str, body_text: str, has_table: bool = False) -> None:
    doc = docx.Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body_text)

    if has_table:
        tbl = doc.add_table(rows=2, cols=2)
        tbl.rows[0].cells[0].paragraphs[0].text = "Header 1"
        tbl.rows[0].cells[1].paragraphs[0].text = "Header 2"
        tbl.rows[1].cells[0].paragraphs[0].text = "Dữ liệu A"
        tbl.rows[1].cells[1].paragraphs[0].text = "Dữ liệu B"
        for r_idx, r in enumerate(tbl.rows):
            schema_helper.set_tr_cant_split(r, True)
            if r_idx == 0:
                schema_helper.set_tr_header(r, True)
            for c in r.cells:
                schema_helper.ensure_cell_last_p(c)

    PackageWriter.write_document_to_path(doc, path)


class TestDocxMerger(unittest.TestCase):
    """Tests multi-document merging and appendix composition."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.base_doc = self.base_dir / "base.docx"
        self.part1 = self.base_dir / "part1.docx"
        self.part2 = self.base_dir / "part2.docx"

        _create_doc(self.base_doc, "Chương 1: Tài liệu chính", "Nội dung phần mở đầu.", has_table=False)
        _create_doc(self.part1, "Phụ lục A: Bảng kê", "Dữ liệu phụ lục chi tiết.", has_table=True)
        _create_doc(self.part2, "Phụ lục B: Kết luận", "Nội dung kết luận báo cáo.", has_table=False)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_merge_two_documents(self):
        res = docx_merger.merge(
            base_input=self.base_doc,
            parts=[self.part1, self.part2],
            style_conflict_policy="master_wins",
        )
        self.assertTrue(len(res.docx_bytes) > 0)
        self.assertEqual(len([i for i in res.issues if i.severity == Severity.ERROR]), 0)
        self.assertIn("merge:master_wins", res.guarantees_applied)
        self.assertIn("guards:cantSplit", res.guarantees_applied)

        # Inspect resulting document
        out_file = self.base_dir / "merged_out.docx"
        out_file.write_bytes(res.docx_bytes)
        doc = docx.Document(out_file)

        full_text = " ".join([p.text for p in doc.paragraphs])
        self.assertIn("Chương 1: Tài liệu chính", full_text)
        self.assertIn("Phụ lục A: Bảng kê", full_text)
        self.assertIn("Phụ lục B: Kết luận", full_text)

        # Check table preserved from part 1
        self.assertTrue(len(doc.tables) >= 1)
        tbl = doc.tables[0]
        self.assertEqual(len(tbl.rows), 2)

    def test_merge_empty_parts_error(self):
        res = docx_merger.merge(base_input=self.base_doc, parts=[])
        self.assertEqual(len(res.docx_bytes), 0)
        self.assertTrue(any(i.code == "E-DOCX-MERGE-NO-PARTS" for i in res.issues))


class TestMergeOperations(unittest.TestCase):
    """Tests docx.merge MCP operation and registry dispatch."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.file_store = FileStore(base_dir=self.base_dir / "store")
        self.registry = ToolRegistry(file_store=self.file_store)
        register_merge_ops(self.registry)

        self.base_doc = self.base_dir / "op_base.docx"
        self.part1 = self.base_dir / "op_part1.docx"
        _create_doc(self.base_doc, "Chính", "Nội dung chính")
        _create_doc(self.part1, "Phụ", "Nội dung phụ")

        self.base_ref = self.file_store.store_file(self.base_doc, engine="docx")
        self.part1_ref = self.file_store.store_file(self.part1, engine="docx")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_docx_merge_operation(self):
        res = docx_merge(
            base_ref_or_path=self.base_ref,
            parts=[self.part1_ref],
            style_conflict_policy="master_wins",
            file_store=self.file_store,
        )
        self.assertTrue(res.success)
        self.assertIsNotNone(res.file_ref)
        self.assertIn("merge:master_wins", res.guarantees_applied)

        resolved = self.file_store.resolve(res.file_ref)
        doc = docx.Document(resolved)
        full_text = " ".join([p.text for p in doc.paragraphs])
        self.assertIn("Chính", full_text)
        self.assertIn("Phụ", full_text)


if __name__ == "__main__":
    unittest.main()
