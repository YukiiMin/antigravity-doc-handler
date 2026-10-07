"""
tests.test_docx.test_inspect_and_gates
Unit tests for StructureInspector, Universal Quality Gates (DG-01 to DG-06),
and docx.inspect_structure / docx.validate MCP operations.
"""

import io
import tempfile
import unittest
from pathlib import Path
import zipfile

import docx
from lxml import etree

from doctools.contract import Severity
from doctools.core.docx.inspect import structure_inspector
from doctools.core.docx.package_io import PackageWriter
from doctools.core.docx.schema import schema_helper
from doctools.gates.docx import docx_quality_gates
from doctools.infra.file_store import FileStore
from doctools.operations.docx.inspect_ops import (
    docx_inspect_structure,
    docx_validate,
    register_inspect_ops,
)
from doctools.registry import ToolRegistry

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _create_sample_doc(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Chương 1: Khởi đầu", level=1)
    doc.add_paragraph("Đoạn văn mở đầu tài liệu kiểm tra.")
    doc.add_heading("Mục 1.1: Chi tiết", level=2)

    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].paragraphs[0].text = "Header A"
    tbl.rows[0].cells[1].paragraphs[0].text = "Header B"
    tbl.rows[1].cells[0].paragraphs[0].text = "Data 1"
    tbl.rows[1].cells[1].paragraphs[0].text = "Data 2"

    # Apply guards
    for r_idx, r in enumerate(tbl.rows):
        schema_helper.set_tr_cant_split(r, True)
        if r_idx == 0:
            schema_helper.set_tr_header(r, True)
        for c in r.cells:
            schema_helper.ensure_cell_last_p(c)

    PackageWriter.write_document_to_path(doc, path)


class TestStructureInspector(unittest.TestCase):
    """Tests StructureInspector tree extraction and anchor stability."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.doc_path = self.base_dir / "sample.docx"
        _create_sample_doc(self.doc_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_inspect_structure_tree(self):
        tree = structure_inspector.inspect(self.doc_path)
        self.assertEqual(tree["root"], "/body")
        self.assertEqual(tree["stats"]["heading_count"], 2)
        self.assertEqual(tree["stats"]["table_count"], 1)

        # Check outline
        headings = tree["headings_outline"]
        self.assertEqual(len(headings), 2)
        self.assertEqual(headings[0]["anchor"], "/body/p[0]")
        self.assertEqual(headings[0]["level"], 1)
        self.assertEqual(headings[0]["text"], "Chương 1: Khởi đầu")

        # Check children nodes
        children = tree["children"]
        self.assertTrue(len(children) >= 4)
        p0 = children[0]
        self.assertEqual(p0["type"], "paragraph")
        self.assertEqual(p0["anchor"], "/body/p[0]")

        tbl = [c for c in children if c["type"] == "table"][0]
        self.assertEqual(tbl["anchor"], "/body/tbl[0]")
        self.assertEqual(tbl["row_count"], 2)
        self.assertEqual(tbl["col_count"], 2)
        self.assertTrue(tbl["has_header_guard"])

        # Check cell anchor
        first_cell = tbl["rows"][0]["cells"][0]
        self.assertEqual(first_cell["anchor"], "/body/tbl[0]/tr[0]/tc[0]")
        self.assertEqual(first_cell["text"], "Header A")


class TestDocxQualityGates(unittest.TestCase):
    """Tests Universal Quality Gates DG-01 to DG-06."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_clean_document_passes_all_gates(self):
        doc_path = self.base_dir / "clean.docx"
        _create_sample_doc(doc_path)

        issues = docx_quality_gates.validate(doc_path)
        errors = [i for i in issues if i.severity == Severity.ERROR]
        warnings = [i for i in issues if i.severity == Severity.WARNING]

        self.assertEqual(len(errors), 0)
        self.assertEqual(len(warnings), 0)

    def test_dg01_corrupt_package(self):
        corrupt_path = self.base_dir / "corrupt.docx"
        corrupt_path.write_bytes(b"NOT_A_VALID_ZIP_CONTENT")

        issues = docx_quality_gates.validate(corrupt_path)
        self.assertTrue(any(i.code == "E-DOCX-PACKAGE-CORRUPT" for i in issues))

    def test_dg02_last_paragraph_rule_violation(self):
        doc = docx.Document()
        tbl = doc.add_table(rows=1, cols=1)
        cell = tbl.rows[0].cells[0]
        # Remove paragraph to violate ERR_DOCX_001
        for p in list(cell._tc):
            cell._tc.remove(p)

        buf = io.BytesIO()
        doc.save(buf)
        issues = docx_quality_gates.validate(buf.getvalue())
        self.assertTrue(any(i.code == "E-DOCX-OPENXML-TC-P" for i in issues))

    def test_dg03_missing_table_guards(self):
        doc = docx.Document()
        # Create table without guards
        doc.add_table(rows=2, cols=2)

        buf = io.BytesIO()
        doc.save(buf)
        issues = docx_quality_gates.validate(buf.getvalue())
        self.assertTrue(any(i.code == "W-DOCX-MISSING-CANTSPLIT" for i in issues))
        self.assertTrue(any(i.code == "W-DOCX-MISSING-TBLHEADER" for i in issues))

    def test_dg05_unresolved_variable(self):
        doc = docx.Document()
        doc.add_paragraph("Kính gửi ông {{ customer_name }}, hợp đồng số {% ref %}.")

        buf = io.BytesIO()
        doc.save(buf)
        issues = docx_quality_gates.validate(buf.getvalue())
        self.assertTrue(any(i.code == "W-DOCX-UNRESOLVED-VARIABLE" for i in issues))

    def test_dg06_heading_hierarchy_skip(self):
        doc = docx.Document()
        doc.add_heading("Heading 1", level=1)
        # Skip level 2 and go straight to level 3
        doc.add_heading("Heading 3 direct", level=3)

        buf = io.BytesIO()
        doc.save(buf)
        issues = docx_quality_gates.validate(buf.getvalue())
        self.assertTrue(any(i.code == "W-DOCX-HEADING-LEVEL-SKIP" for i in issues))


class TestInspectOperations(unittest.TestCase):
    """Tests MCP inspect and validation operations."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.file_store = FileStore(base_dir=self.base_dir / "store")
        self.registry = ToolRegistry(file_store=self.file_store)
        register_inspect_ops(self.registry)

        self.doc_path = self.base_dir / "test_op.docx"
        _create_sample_doc(self.doc_path)
        self.file_ref = self.file_store.store_file(self.doc_path, engine="docx")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_docx_inspect_structure_operation(self):
        res = docx_inspect_structure(self.file_ref, file_store=self.file_store)
        self.assertTrue(res.success)
        self.assertIsNotNone(res.stats)
        self.assertIn("structure", res.stats.extra)
        tree = res.stats.extra["structure"]
        self.assertEqual(tree["root"], "/body")
        self.assertTrue(tree["stats"]["paragraph_count"] > 0)

    def test_docx_validate_operation(self):
        res = docx_validate(self.file_ref, file_store=self.file_store)
        self.assertTrue(res.success)
        self.assertEqual(res.diagnostics.total_count, 0)

    def test_registry_tool_execution(self):
        res = self.registry.execute("docx.inspect_structure", {"file_ref": self.file_ref})
        self.assertTrue(res.success)
        self.assertIsNotNone(res.stats)
        self.assertIn("structure", res.stats.extra)


if __name__ == "__main__":
    unittest.main()
