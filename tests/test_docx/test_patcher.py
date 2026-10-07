"""
tests.test_docx.test_patcher
Unit tests for DocxPatcher and docx.patch MCP tool operation (FR-11).
"""

import tempfile
import unittest
from pathlib import Path

import docx

from doctools.contract import Severity
from doctools.contract.docx.patch import PatchOp, PatchSpec
from doctools.core.docx.package_io import PackageWriter
from doctools.core.docx.patch import docx_patcher
from doctools.core.docx.schema import schema_helper
from doctools.infra.file_store import FileStore
from doctools.operations.docx.patch_ops import docx_patch, register_patch_ops
from doctools.registry import ToolRegistry


def _create_base_doc(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Tiêu đề gốc", level=1)
    p = doc.add_paragraph("Nội dung hợp đồng cũ trị giá 100.000.000 VND.")
    p.runs[0].bold = True

    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].paragraphs[0].text = "Header 1"
    tbl.rows[0].cells[1].paragraphs[0].text = "Header 2"
    tbl.rows[1].cells[0].paragraphs[0].text = "Ô dữ liệu cũ"
    tbl.rows[1].cells[1].paragraphs[0].text = "Trạng thái cũ"

    for r_idx, r in enumerate(tbl.rows):
        schema_helper.set_tr_cant_split(r, True)
        if r_idx == 0:
            schema_helper.set_tr_header(r, True)
        for c in r.cells:
            schema_helper.ensure_cell_last_p(c)

    PackageWriter.write_document_to_path(doc, path)


class TestDocxPatcher(unittest.TestCase):
    """Tests declarative anchor-based patching."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.doc_path = self.base_dir / "target.docx"
        _create_base_doc(self.doc_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_replace_text_preserves_run_formatting(self):
        spec = PatchSpec(
            operations=[
                PatchOp(
                    anchor="/body/p[1]",
                    op="replace_text",
                    target_text="100.000.000 VND",
                    payload="250.000.000 VND",
                )
            ]
        )
        res = docx_patcher.patch(self.doc_path, spec)
        self.assertTrue(len(res.docx_bytes) > 0)
        self.assertEqual(len([i for i in res.issues if i.severity == Severity.ERROR]), 0)

        # Inspect resulting document
        out_file = self.base_dir / "patched_text.docx"
        out_file.write_bytes(res.docx_bytes)
        doc = docx.Document(out_file)
        p1 = doc.paragraphs[1]
        self.assertIn("250.000.000 VND", p1.text)
        self.assertTrue(p1.runs[0].bold)

    def test_update_table_cell_preserves_invariants(self):
        spec = PatchSpec(
            operations=[
                PatchOp(
                    anchor="/body/tbl[0]/tr[1]/tc[0]",
                    op="update_cell",
                    payload="Dữ liệu mới cập nhật",
                )
            ]
        )
        res = docx_patcher.patch(self.doc_path, spec)
        self.assertTrue(len(res.docx_bytes) > 0)

        out_file = self.base_dir / "patched_cell.docx"
        out_file.write_bytes(res.docx_bytes)
        doc = docx.Document(out_file)
        cell = doc.tables[0].rows[1].cells[0]
        self.assertIn("Dữ liệu mới cập nhật", cell.text)
        # Check last paragraph rule
        self.assertEqual(cell._tc[-1].tag, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")

    def test_insert_paragraph_after(self):
        spec = PatchSpec(
            operations=[
                PatchOp(
                    anchor="/body/p[0]",
                    op="insert_paragraph_after",
                    payload="Đoạn văn được chèn thêm ngay sau tiêu đề.",
                )
            ]
        )
        res = docx_patcher.patch(self.doc_path, spec)
        self.assertTrue(len(res.docx_bytes) > 0)

        out_file = self.base_dir / "patched_insert.docx"
        out_file.write_bytes(res.docx_bytes)
        doc = docx.Document(out_file)
        self.assertEqual(doc.paragraphs[1].text, "Đoạn văn được chèn thêm ngay sau tiêu đề.")

    def test_delete_paragraph(self):
        spec = PatchSpec(
            operations=[
                PatchOp(
                    anchor="/body/p[1]",
                    op="delete",
                )
            ]
        )
        res = docx_patcher.patch(self.doc_path, spec)
        self.assertTrue(len(res.docx_bytes) > 0)

        out_file = self.base_dir / "patched_del.docx"
        out_file.write_bytes(res.docx_bytes)
        doc = docx.Document(out_file)
        full_text = " ".join([p.text for p in doc.paragraphs])
        self.assertNotIn("Nội dung hợp đồng cũ", full_text)


class TestPatchOperations(unittest.TestCase):
    """Tests MCP docx.patch operation and ToolRegistry integration."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.file_store = FileStore(base_dir=self.base_dir / "store")
        self.registry = ToolRegistry(file_store=self.file_store)
        register_patch_ops(self.registry)

        self.doc_path = self.base_dir / "op_patch.docx"
        _create_base_doc(self.doc_path)
        self.file_ref = self.file_store.store_file(self.doc_path, engine="docx")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_docx_patch_operation(self):
        ops = [
            {
                "anchor": "/body/p[1]",
                "op": "replace_text",
                "target_text": "100.000.000 VND",
                "payload": "500.000.000 VND",
            }
        ]
        res = docx_patch(
            file_ref_or_path=self.file_ref,
            operations=ops,
            file_store=self.file_store,
        )
        self.assertTrue(res.success)
        self.assertIsNotNone(res.file_ref)
        self.assertIn("format_preserving_patch", res.guarantees_applied)

        # Verify patched content
        resolved = self.file_store.resolve(res.file_ref)
        doc = docx.Document(resolved)
        self.assertIn("500.000.000 VND", doc.paragraphs[1].text)


if __name__ == "__main__":
    unittest.main()
