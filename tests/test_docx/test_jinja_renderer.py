"""
tests.test_docx.test_jinja_renderer
Unit tests for JinjaRenderer (Path A) and build operations (docx.render_template, docx.build_document).
"""

import tempfile
import unittest
from pathlib import Path

import docx
from lxml import etree

from doctools.contract import Severity
from doctools.contract.docx.manifest import LayoutPolicy, TemplateManifest, VariableDef
from doctools.core.docx.build import jinja_renderer
from doctools.core.docx.package_io import PackageWriter
from doctools.core.docx.template.template_registry import TemplateRegistry
from doctools.infra.file_store import FileStore
from doctools.operations.docx.build_ops import (
    docx_build_document,
    docx_render_template,
    register_build_ops,
)
from doctools.registry import ToolRegistry

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _create_template_docx(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Contract {{ contract_id }}", level=1)
    doc.add_paragraph("Customer: {{ customer_name }}")
    doc.add_paragraph("Summary: {{ summary }}")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header 1"
    table.rows[0].cells[1].text = "Header 2"
    table.rows[1].cells[0].text = "{{ item_1 }}"
    table.rows[1].cells[1].text = "{{ item_2 }}"

    PackageWriter.write_document_to_path(doc, path)


class TestJinjaRenderer(unittest.TestCase):
    """Tests Path A JinjaRenderer in sandboxed environment with OOXML guards."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.tpl_path = self.base_dir / "test_template.docx"
        _create_template_docx(self.tpl_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_render_basic_template_with_guards(self):
        context = {
            "contract_id": "HD-2026-001",
            "customer_name": "Antigravity Corp",
            "summary": "Short executive summary",
            "item_1": "Server Setup",
            "item_2": "Completed",
        }
        policy = LayoutPolicy(apply_guards="apply")
        result = jinja_renderer.render(
            template_input=self.tpl_path,
            context_data=context,
            layout_policy=policy,
        )

        self.assertTrue(len(result.docx_bytes) > 0)
        self.assertEqual(len(result.issues), 0)
        self.assertIn("sandboxed_jinja_render", result.guarantees_applied)
        self.assertIn("guards:cantSplit", result.guarantees_applied)
        self.assertIn("guards:tblHeader", result.guarantees_applied)

        # Inspect resulting document structure
        out_path = self.base_dir / "rendered.docx"
        out_path.write_bytes(result.docx_bytes)
        doc = docx.Document(out_path)

        # Check rendered text
        full_text = " ".join([p.text for p in doc.paragraphs])
        self.assertIn("Contract HD-2026-001", full_text)
        self.assertIn("Customer: Antigravity Corp", full_text)

        # Check table OOXML guards
        table = doc.tables[0]
        # cantSplit on every row
        for row in table.rows:
            trPr = row._tr.find(f"{_W}trPr")
            self.assertIsNotNone(trPr)
            self.assertIsNotNone(trPr.find(f"{_W}cantSplit"))

        # tblHeader on header row 0
        trPr_0 = table.rows[0]._tr.find(f"{_W}trPr")
        self.assertIsNotNone(trPr_0.find(f"{_W}tblHeader"))

        # vAlign and last_p on cells
        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.find(f"{_W}tcPr")
                self.assertIsNotNone(tcPr)
                valign = tcPr.find(f"{_W}vAlign")
                self.assertIsNotNone(valign)
                self.assertEqual(valign.get(f"{_W}val"), "center")
                # Last paragraph rule
                self.assertEqual(cell._tc[-1].tag, f"{_W}p")

        # Heading keepNext guard
        h1 = doc.paragraphs[0]
        pPr = h1._p.find(f"{_W}pPr")
        self.assertIsNotNone(pPr)
        self.assertIsNotNone(pPr.find(f"{_W}keepNext"))

    def test_render_immutable_constraint_violation(self):
        manifest = TemplateManifest(
            template_id="strict_contract",
            variables={
                "customer_name": VariableDef(type="string", immutable=True),
            },
        )
        # Attempt to pass empty value for immutable field
        context = {"customer_name": ""}
        result = jinja_renderer.render(
            template_input=self.tpl_path,
            context_data=context,
            manifest=manifest,
        )
        self.assertTrue(any(i.code == "E-DOCX-IMMUTABLE-VIOLATION" for i in result.issues))
        self.assertEqual(len(result.docx_bytes), 0)

    def test_render_max_chars_warning(self):
        manifest = TemplateManifest(
            template_id="brief_report",
            variables={
                "summary": VariableDef(type="string", max_chars=10),
            },
        )
        context = {
            "contract_id": "C01",
            "customer_name": "Test",
            "summary": "This is a very long summary that exceeds 10 chars",
            "item_1": "A",
            "item_2": "B",
        }
        result = jinja_renderer.render(
            template_input=self.tpl_path,
            context_data=context,
            manifest=manifest,
        )
        # Warning should be raised, but rendering succeeds
        self.assertTrue(len(result.docx_bytes) > 0)
        self.assertTrue(any(i.code == "W-DOCX-MAX-CHARS-EXCEEDED" for i in result.issues))


class TestBuildOperations(unittest.TestCase):
    """Tests docx.render_template and docx.build_document operations."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)
        self.file_store = FileStore(base_dir=self.base_dir / "store")
        self.template_registry = TemplateRegistry(file_store=self.file_store)

        self.tpl_path = self.base_dir / "op_template.docx"
        _create_template_docx(self.tpl_path)

        # Register template into registry
        self.template_registry.register(
            template_ref=self.tpl_path,
            manifest_input={
                "template_id": "registered_invoice",
                "variables": {
                    "contract_id": {"type": "string"},
                    "customer_name": {"type": "string"},
                },
                "layout_policy_default": {"apply_guards": "apply"},
            },
        )

        self.mcp_registry = ToolRegistry(file_store=self.file_store)
        register_build_ops(self.mcp_registry)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_docx_render_template_by_id(self):
        res = docx_render_template(
            template_ref_or_id="registered_invoice",
            context_data={
                "contract_id": "INV-100",
                "customer_name": "Global Tech",
                "summary": "OK",
                "item_1": "Consulting",
                "item_2": "Done",
            },
            file_store=self.file_store,
            template_registry=self.template_registry,
        )
        self.assertTrue(res.success)
        self.assertIsNotNone(res.file_ref)
        self.assertIn("sandboxed_jinja_render", res.guarantees_applied)

    def test_docx_build_document_unified(self):
        res = docx_build_document(
            source_type="template",
            template_id="registered_invoice",
            context_data={
                "contract_id": "INV-200",
                "customer_name": "DeepMind",
                "summary": "Research",
                "item_1": "Agent",
                "item_2": "Success",
            },
            file_store=self.file_store,
            template_registry=self.template_registry,
        )
        self.assertTrue(res.success)
        self.assertIsNotNone(res.file_ref)


if __name__ == "__main__":
    unittest.main()
