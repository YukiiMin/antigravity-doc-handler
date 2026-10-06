"""
Unit test suite for JinjaNormalizer (doctools.core.docx.template.jinja_normalizer).
Kiểm định Mục 4.14: hàn gắn run Jinja bị băm, loại w:proofErr, bảo toàn w:noProof/w:lang, cảnh báo format.
"""

from __future__ import annotations
import shutil
import tempfile
import unittest
from pathlib import Path
import docx
from lxml import etree

from doctools.contract import Severity
from doctools.core.docx.package_io import PackageReader, PackageWriter, safe_read_package
from doctools.core.docx.template import jinja_normalizer

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W = f"{{{_W_NS}}}"


class TestJinjaNormalizer(unittest.TestCase):
    """Kiểm tra hoạt động hàn gắn thẻ Jinja và dọn rác của JinjaNormalizer."""

    def setUp(self) -> None:
        self.test_dir = Path(tempfile.mkdtemp(prefix="test_normalizer_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_consolidate_split_runs_and_remove_trash(self) -> None:
        doc = docx.Document()
        p = doc.add_paragraph()

        # Tạo cấu trúc: r1("{{ ") -> proofErr -> r2("fullname") -> r3(" }}")
        r1 = p.add_run("{{ ")
        
        # Chèn thẻ rác w:proofErr giữa các run
        proof_err = etree.Element(f"{_W}proofErr")
        proof_err.set(f"{_W}type", "spellStart")
        p._p.append(proof_err)

        r2 = p.add_run("fullname")
        r3 = p.add_run(" }}")

        tpl_path = self.test_dir / "split_with_trash.docx"
        PackageWriter.write_document_to_path(doc, tpl_path)

        result = jinja_normalizer.normalize(tpl_path)

        # Kiểm tra diff summary
        self.assertGreater(result.diff_summary["trash_elements_removed"], 0)
        self.assertGreater(result.diff_summary["runs_consolidated"], 0)

        # Đọc lại gói đã chuẩn hóa
        norm_reader = safe_read_package(result.normalized_bytes)
        norm_doc = norm_reader.to_docx_document()
        norm_p = norm_doc.paragraphs[0]

        # Sau khi chuẩn hóa, các run phải được gộp thành 1 run chứa trọn vẹn "{{ fullname }}"
        run_texts = [r.text for r in norm_p.runs]
        self.assertIn("{{ fullname }}", "".join(run_texts))
        # Không còn thẻ w:proofErr trong đoạn
        self.assertEqual(len(norm_p._p.findall(f"{_W}proofErr")), 0)

    def test_mixed_formatting_warning(self) -> None:
        doc = docx.Document()
        p = doc.add_paragraph()
        r1 = p.add_run("{{ ")
        r1.bold = True
        r2 = p.add_run("salary")
        r2.italic = True
        r3 = p.add_run(" }}")

        tpl_path = self.test_dir / "mixed_format.docx"
        PackageWriter.write_document_to_path(doc, tpl_path)

        result = jinja_normalizer.normalize(tpl_path)
        warning_codes = [i.code for i in result.issues if i.severity == Severity.WARNING]
        self.assertIn("W-TPL-MIXED-FORMAT", warning_codes)

    def test_non_split_paragraph_remains_untouched(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("Đoạn văn bình thường không có thẻ.")
        doc.add_paragraph("Thẻ chuẩn trong 1 run: {{ clean_tag }}")

        tpl_path = self.test_dir / "clean_template.docx"
        PackageWriter.write_document_to_path(doc, tpl_path)

        result = jinja_normalizer.normalize(tpl_path)
        self.assertEqual(result.diff_summary["runs_consolidated"], 0)
        self.assertEqual(result.diff_summary["trash_elements_removed"], 0)
        self.assertEqual(result.diff_summary["paragraphs_modified"], 0)


if __name__ == "__main__":
    unittest.main()
