"""
Unit test suite for TemplateLinter (doctools.core.docx.template.template_linter).
Kiểm định tĩnh: cú pháp Jinja, trích xuất biến, phát hiện split tags, dynamic fields, SSTI.
"""

from __future__ import annotations
import shutil
import tempfile
import unittest
from pathlib import Path
import docx
from lxml import etree

from doctools.contract import Severity
from doctools.core.docx.package_io import PackageWriter
from doctools.core.docx.template import template_linter

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W = f"{{{_W_NS}}}"


class TestTemplateLinter(unittest.TestCase):
    """Kiểm tra hoạt động của TemplateLinter."""

    def setUp(self) -> None:
        self.test_dir = Path(tempfile.mkdtemp(prefix="test_linter_"))

    def tearDown(self) -> None:
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_valid_template_variables_extracted(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("Kính gửi: {{ customer_name }}")
        doc.add_paragraph("Hợp đồng số: {{ contract.number }} - Ngày: {{ date }}")

        tpl_path = self.test_dir / "valid_template.docx"
        PackageWriter.write_document_to_path(doc, tpl_path)

        report = template_linter.lint(tpl_path)
        self.assertTrue(report.valid)
        self.assertEqual(len([i for i in report.issues if i.severity == Severity.ERROR]), 0)
        self.assertIn("customer_name", report.variables)
        self.assertIn("contract", report.variables)
        self.assertIn("date", report.variables)

    def test_unclosed_jinja_tag_detected(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("Thẻ bị lỗi mở không đóng: {{ unclosed_variable ")

        tpl_path = self.test_dir / "unclosed_template.docx"
        PackageWriter.write_document_to_path(doc, tpl_path)

        report = template_linter.lint(tpl_path)
        self.assertFalse(report.valid)
        error_codes = [i.code for i in report.issues if i.severity == Severity.ERROR]
        self.assertIn("E-TPL-UNCLOSED-TAG", error_codes)

    def test_ssti_prohibited_pattern_detected(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("Tấn công SSTI: {{ ''.__class__.__mro__ }}")

        tpl_path = self.test_dir / "ssti_template.docx"
        PackageWriter.write_document_to_path(doc, tpl_path)

        report = template_linter.lint(tpl_path)
        self.assertFalse(report.valid)
        error_codes = [i.code for i in report.issues if i.severity == Severity.ERROR]
        self.assertIn("E-SEC-SSTI", error_codes)

    def test_split_tag_detected(self) -> None:
        doc = docx.Document()
        p = doc.add_paragraph()
        # Tạo 3 run tách rời thẻ Jinja
        r1 = p.add_run("{{ ")
        r2 = p.add_run("employee_id")
        r3 = p.add_run(" }}")

        tpl_path = self.test_dir / "split_template.docx"
        PackageWriter.write_document_to_path(doc, tpl_path)

        report = template_linter.lint(tpl_path)
        warning_codes = [i.code for i in report.issues if i.severity == Severity.WARNING]
        self.assertIn("W-TPL-SPLIT-TAG", warning_codes)

    def test_dynamic_toc_detected(self) -> None:
        doc = docx.Document()
        p = doc.add_paragraph("Mục lục báo cáo:")
        # Thêm field simple TOC
        fld = etree.SubElement(p._p, f"{_W}fldSimple")
        fld.set(f"{_W}instr", "TOC \\o \"1-3\" \\h \\z \\u")

        tpl_path = self.test_dir / "toc_template.docx"
        PackageWriter.write_document_to_path(doc, tpl_path)

        report = template_linter.lint(tpl_path)
        self.assertTrue(report.fields_detected["toc"])
        info_codes = [i.code for i in report.issues if i.severity == Severity.INFO]
        self.assertIn("W-FIELD-TOC-DETECTED", info_codes)


if __name__ == "__main__":
    unittest.main()
