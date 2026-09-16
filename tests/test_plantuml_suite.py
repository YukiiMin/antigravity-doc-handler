"""
test_plantuml_suite.py — Comprehensive end-to-end test suite for PlantUML Diagram Engine.

Tests:
1. ERD (Entity-Relationship Diagram) with Crow's foot notation.
2. OOP Class Diagram rendering.
3. C4 Architecture (Context) using internal Standard Library (<C4/C4_Context>).
4. Activity Diagram with Swimlanes.
5. Use Case Diagram rendering.
6. Mind Map diagram rendering.
7. Vietnamese encoding trap verification (accented characters: đ, ư, ơ, ã, ê, ô).
8. Aspect ratio scaling verification (wide vs tall diagrams).
9. DOCX Placeholder injection & keepNext chain.
10. Atomic Protection on PlantUML syntax error.
11. Unified Dispatcher (`ai_tools_cli.py diagram-render`) with auto-inject into DOCX.
"""

from __future__ import annotations

import os
import shutil
import tempfile
import argparse
import sys
from docx import Document
from docx.oxml.ns import qn

TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.abspath(os.path.join(TESTS_DIR, ".."))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from plantuml_renderer import (
    calculate_aspect_dimensions,
    render_plantuml_to_png,
    render_plantuml,
)
from ai_tools_cli import cmd_diagram_render


def run_tests() -> bool:
    print("=" * 65)
    print("STARTING PLANTUML DIAGRAM ENGINE END-TO-END TEST SUITE")
    print("=" * 65)

    test_dir = tempfile.mkdtemp(prefix="plantuml_test_")
    all_passed = True

    try:
        # -------------------------------------------------------------
        # Test 1: ERD with Crow's Foot notation
        # -------------------------------------------------------------
        print("\n[TEST 1] ERD Rendering (Crow's foot PK/FK)...")
        spec_erd = {
            "diagram_name": "erd_sample",
            "code": """@startuml
hide circle
skinparam linetype ortho

entity "User" as user {
  * user_id : INTEGER <<PK>>
  --
  username : VARCHAR(50)
  email : VARCHAR(100)
  created_at : TIMESTAMP
}

entity "Order" as ord {
  * order_id : INTEGER <<PK>>
  --
  * user_id : INTEGER <<FK>>
  total_price : DECIMAL(10,2)
  status : ENUM
  order_date : DATE
}

entity "OrderItem" as item {
  * item_id : INTEGER <<PK>>
  --
  * order_id : INTEGER <<FK>>
  product_name : VARCHAR(100)
  quantity : INTEGER
  price : DECIMAL(10,2)
}

user ||--o{ ord : "places"
ord ||--|{ item : "contains"
@enduml""",
            "output_path": os.path.join(test_dir, "erd.png"),
            "dpi": 300,
        }
        res_erd = render_plantuml_to_png(spec_erd)
        assert os.path.isfile(res_erd["png_path"]), "ERD PNG not created"
        assert res_erd["file_size_bytes"] > 1024, "ERD PNG too small"
        assert res_erd["width_cm"] <= 14.0, "Width exceeds 14.0cm"
        assert res_erd["height_cm"] <= 20.0, "Height exceeds 20.0cm"
        print(f"  --> PASS: {res_erd['png_path']} ({res_erd['dimensions_px']}) | Doc: {res_erd['width_cm']}x{res_erd['height_cm']}cm")

        # -------------------------------------------------------------
        # Test 2: OOP Class Diagram
        # -------------------------------------------------------------
        print("\n[TEST 2] Class Diagram Rendering (OOP)...")
        spec_class = {
            "diagram_name": "class_sample",
            "code": """@startuml
skinparam classAttributeIconSize 0

interface IRepository<T> {
  + getById(id: int): T
  + save(entity: T): void
}

class UserRepository implements IRepository {
  - dbContext: DatabaseContext
  + getById(id: int): User
  + save(entity: User): void
}

class User {
  - id: int
  - name: String
  - email: String
  + getId(): int
  + getEmail(): String
}

UserRepository --> User : "manages"
@enduml""",
            "output_path": os.path.join(test_dir, "class.png"),
            "dpi": 300,
        }
        res_class = render_plantuml_to_png(spec_class)
        assert os.path.isfile(res_class["png_path"]), "Class PNG not created"
        assert res_class["file_size_bytes"] > 1024, "Class PNG too small"
        print(f"  --> PASS: {res_class['png_path']} ({res_class['dimensions_px']}) | Doc: {res_class['width_cm']}x{res_class['height_cm']}cm")

        # -------------------------------------------------------------
        # Test 3: C4 Context Diagram (Standard Library Offline)
        # -------------------------------------------------------------
        print("\n[TEST 3] C4 Architecture Rendering (Standard Library Offline <C4/C4_Context>)...")
        spec_c4 = {
            "diagram_name": "c4_sample",
            "code": """@startuml
!include <C4/C4_Context>

Person(user, "Người dùng", "Khách hàng sử dụng ứng dụng di động")
System(app, "Hệ thống SCORT", "Hệ thống sao chép đối tượng dữ liệu SAP")
System_Ext(sap_core, "SAP S/4HANA", "Hệ thống ERP trung tâm")
System_Ext(vnpay, "VNPay Gateway", "Cổng thanh toán điện tử")

Rel(user, app, "Thao tác trên ứng dụng", "HTTPS")
Rel(app, sap_core, "Đồng bộ dữ liệu qua", "OData V4 / RFC")
Rel(app, vnpay, "Thanh toán qua", "REST API")
@enduml""",
            "output_path": os.path.join(test_dir, "c4.png"),
            "dpi": 300,
        }
        res_c4 = render_plantuml_to_png(spec_c4)
        assert os.path.isfile(res_c4["png_path"]), "C4 PNG not created"
        assert res_c4["file_size_bytes"] > 1024, "C4 PNG too small"
        print(f"  --> PASS: {res_c4['png_path']} ({res_c4['dimensions_px']}) | Doc: {res_c4['width_cm']}x{res_c4['height_cm']}cm")

        # -------------------------------------------------------------
        # Test 4: Activity Diagram with Swimlanes
        # -------------------------------------------------------------
        print("\n[TEST 4] Activity Diagram with Swimlanes...")
        spec_swim = {
            "diagram_name": "swimlane_sample",
            "code": """@startuml
|Khách hàng|
start
:Mở ứng dụng;
:Chọn món ăn;
:Bấm "Đặt đơn hàng";

|Hệ thống|
:Tạo đơn hàng;
:Gửi yêu cầu thanh toán;

|Cổng thanh toán|
:Xử lý thẻ ngân hàng;
if (Thành công?) then (có)
  :Gửi mã giao dịch;
  |Hệ thống|
  :Cập nhật trạng thái Đã thanh toán;
  |Khách hàng|
  :Nhận hóa đơn điện tử;
else (không)
  :Báo lỗi thanh toán;
  |Khách hàng|
  :Thử lại phương thức khác;
endif
stop
@enduml""",
            "output_path": os.path.join(test_dir, "swimlane.png"),
            "dpi": 300,
        }
        res_swim = render_plantuml_to_png(spec_swim)
        assert os.path.isfile(res_swim["png_path"]), "Swimlane PNG not created"
        print(f"  --> PASS: {res_swim['png_path']} ({res_swim['dimensions_px']}) | Doc: {res_swim['width_cm']}x{res_swim['height_cm']}cm")

        # -------------------------------------------------------------
        # Test 5: Use Case Diagram
        # -------------------------------------------------------------
        print("\n[TEST 5] Use Case Diagram...")
        spec_uc = {
            "diagram_name": "usecase_sample",
            "code": """@startuml
left to right direction
actor "Khách hàng" as customer
actor "Quản trị viên" as admin

rectangle "Hệ thống Bán lẻ" {
  usecase "Xem danh mục" as UC1
  usecase "Thêm vào giỏ" as UC2
  usecase "Thanh toán" as UC3
  usecase "Quản trị kho" as UC4
}

customer --> UC1
customer --> UC2
customer --> UC3
admin --> UC4
@enduml""",
            "output_path": os.path.join(test_dir, "usecase.png"),
            "dpi": 300,
        }
        res_uc = render_plantuml_to_png(spec_uc)
        assert os.path.isfile(res_uc["png_path"]), "Use Case PNG not created"
        print(f"  --> PASS: {res_uc['png_path']} ({res_uc['dimensions_px']}) | Doc: {res_uc['width_cm']}x{res_uc['height_cm']}cm")

        # -------------------------------------------------------------
        # Test 6: Mind Map
        # -------------------------------------------------------------
        print("\n[TEST 6] Mind Map Diagram...")
        spec_mm = {
            "diagram_name": "mindmap_sample",
            "code": """@startmindmap
* SCORT System
** DB Core
*** CDS View Entities
*** Behavior Definitions
*** DDIC Tables
** API Services
*** Service Definitions
*** Service Bindings OData V4
** Frontend UI5
*** List Report
*** Object Page
@endmindmap""",
            "output_path": os.path.join(test_dir, "mindmap.png"),
            "dpi": 300,
        }
        res_mm = render_plantuml_to_png(spec_mm)
        assert os.path.isfile(res_mm["png_path"]), "Mindmap PNG not created"
        print(f"  --> PASS: {res_mm['png_path']} ({res_mm['dimensions_px']}) | Doc: {res_mm['width_cm']}x{res_mm['height_cm']}cm")

        # -------------------------------------------------------------
        # Test 7: Vietnamese Encoding Verification (Trap 2)
        # -------------------------------------------------------------
        print("\n[TEST 7] Vietnamese Encoding Trap Verification (-charset UTF-8)...")
        vn_code = """@startuml
title Tiếng Việt hoàn hảo: Đảng, Nước, Đường, Giấy, Ốc, Ứng dụng
class "Bảng Khách Hàng" as Table1 {
  + mã_định_danh : INT
  + họ_và_tên : VARCHAR
  + địa_chỉ_thường_trú : VARCHAR
  + ngày_sinh : DATE
}
@enduml"""
        vn_out = os.path.join(test_dir, "vietnamese_test.png")
        res_vn = render_plantuml(vn_code, vn_out)
        assert os.path.isfile(vn_out), "Vietnamese test PNG not created"
        assert res_vn["file_size_bytes"] > 1024, "Vietnamese PNG too small"
        print(f"  --> PASS: Vietnamese PNG rendered properly ({res_vn['dimensions_px']})")

        # -------------------------------------------------------------
        # Test 8: Aspect Ratio Scaling Validation
        # -------------------------------------------------------------
        print("\n[TEST 8] Aspect Ratio Scaling Validation...")
        w, h = calculate_aspect_dimensions(2000, 800, max_w_cm=14.0, max_h_cm=20.0)
        assert w == 14.0 and h == 5.6, f"Expected (14.0, 5.6), got ({w}, {h})"
        w2, h2 = calculate_aspect_dimensions(1000, 3000, max_w_cm=14.0, max_h_cm=20.0)
        assert h2 == 20.0 and w2 == 6.67, f"Expected (6.67, 20.0), got ({w2}, {h2})"
        print(f"  --> PASS: Wide -> ({w}cm, {h}cm) | Tall -> ({w2}cm, {h2}cm)")

        # -------------------------------------------------------------
        # Test 9: DOCX Injection with Placeholder & keepNext Chain
        # -------------------------------------------------------------
        print("\n[TEST 9] DOCX Placeholder Injection & keepNext Chain...")
        test_docx_path = os.path.join(test_dir, "test_document.docx")
        doc = Document()
        doc.add_heading("3. Thiết kế Cơ sở Dữ liệu", level=1)
        doc.add_paragraph("Mô tả tổng quan về kiến trúc dữ liệu.")
        p_place = doc.add_paragraph()
        p_place.add_run("{{DIAGRAM_")
        p_place.add_run("DATABASE_")
        p_place.add_run("ERD}}")
        doc.add_paragraph("Đoạn văn kết thúc sau biểu đồ ERD.")
        doc.save(test_docx_path)

        spec_inject_placeholder = {
            "code": spec_erd["code"],
            "diagram_name": "erd_inject",
            "inject_into": test_docx_path,
            "placeholder": "{{DIAGRAM_DATABASE_ERD}}",
            "caption": "Hình 3.1: Sơ đồ ERD Cơ sở Dữ liệu Hệ thống",
            "caption_template": {
                "caption": "Hình 3.1: Sơ đồ ERD Cơ sở Dữ liệu Hệ thống",
                "description": "Biểu đồ thực thể thể hiện mối quan hệ giữa người dùng, đơn hàng và chi tiết sản phẩm.",
            },
        }
        res_inj = render_plantuml_to_png(spec_inject_placeholder, base_dir=test_dir)
        assert res_inj["injected"] is True, "Injection flag is not True"

        # Verify modified DOCX
        doc_mod = Document(test_docx_path)
        all_texts = [p.text for p in doc_mod.paragraphs]

        for t in all_texts:
            assert "{{DIAGRAM_DATABASE_ERD}}" not in t, f"Placeholder still present: {t}"
            assert "{{DIAGRAM_" not in t, f"Fragmented run still present: {t}"

        assert any("Hình 3.1: Sơ đồ ERD" in t for t in all_texts), "Caption not found in DOCX"
        assert any("Biểu đồ thực thể" in t for t in all_texts), "Description not found in DOCX"

        # Check keepNext
        img_p_found = False
        cap_p_found = False
        desc_p_found = False

        for p in doc_mod.paragraphs:
            drawings = p._element.findall(".//" + qn("w:drawing"))
            if drawings:
                img_p_found = True
                pPr = p._element.find(qn("w:pPr"))
                assert pPr is not None and pPr.find(qn("w:keepNext")) is not None, "Image missing keepNext"

            if "Hình 3.1:" in p.text:
                cap_p_found = True
                pPr = p._element.find(qn("w:pPr"))
                assert pPr is not None and pPr.find(qn("w:keepNext")) is not None, "Caption missing keepNext"

            if "Biểu đồ thực thể" in p.text:
                desc_p_found = True
                pPr = p._element.find(qn("w:pPr"))
                has_keep = (pPr is not None and pPr.find(qn("w:keepNext")) is not None)
                assert not has_keep, "Final description should NOT have keepNext"

        assert img_p_found and cap_p_found and desc_p_found, "Not all injected paragraphs found"
        print("  --> PASS: Placeholder replaced cleanly, keepNext chain verified!")

        # -------------------------------------------------------------
        # Test 10: Atomic Protection on PlantUML syntax error
        # -------------------------------------------------------------
        print("\n[TEST 10] Atomic Protection on PlantUML Syntax Error...")
        mod_time_before = os.path.getmtime(test_docx_path)
        spec_fail = {
            "code": "this is total invalid plantuml syntax !!! (((",
            "diagram_name": "fail_sample",
            "inject_into": test_docx_path,
            "placeholder": "nonexistent",
        }
        caught_error = False
        try:
            render_plantuml_to_png(spec_fail, base_dir=test_dir)
        except (RuntimeError, ValueError, Exception) as exc:
            caught_error = True
            print(f"  Caught expected error: {type(exc).__name__}")

        assert caught_error, "Atomic Protection FAILED: Expected exception was not raised"
        mod_time_after = os.path.getmtime(test_docx_path)
        assert mod_time_before == mod_time_after, "Atomic Protection FAILED: DOCX was touched on failure"
        print("  --> PASS: Atomic Protection passed (DOCX was untouched on compilation error)")

        # -------------------------------------------------------------
        # Test 11: Unified Dispatcher (ai_tools_cli.py diagram-render)
        # -------------------------------------------------------------
        print("\n[TEST 11] Unified Dispatcher (diagram-render CLI)...")
        test_cli_docx = os.path.join(test_dir, "cli_test_doc.docx")
        doc_cli = Document()
        doc_cli.add_heading("2. Kiến trúc C4", level=1)
        p_c4_place = doc_cli.add_paragraph()
        p_c4_place.add_run("[[DIAGRAM_C4_CONTEXT]]")
        doc_cli.save(test_cli_docx)

        unified_spec_path = os.path.join(test_dir, "unified_c4_spec.json")
        import json
        with open(unified_spec_path, "w", encoding="utf-8") as f:
            json.dump({
                "engine": "plantuml",
                "diagram_type": "c4_context",
                "code": spec_c4["code"],
                "inject_into": test_cli_docx,
                "placeholder": "[[DIAGRAM_C4_CONTEXT]]",
                "caption": "Hình 2.1: Sơ đồ C4 Context Hệ thống SCORT",
                "width_cm": 14.0,
            }, f, ensure_ascii=False, indent=2)

        cli_args = argparse.Namespace(
            spec=unified_spec_path,
            output=None,
            inject=None,
            scale=None,
        )
        exit_code = cmd_diagram_render(cli_args)
        assert exit_code == 0, f"cmd_diagram_render failed with exit code {exit_code}"

        doc_cli_mod = Document(test_cli_docx)
        cli_texts = [p.text for p in doc_cli_mod.paragraphs]
        assert any("Hình 2.1: Sơ đồ C4 Context" in t for t in cli_texts), "C4 caption not injected via CLI"
        print("  --> PASS: Unified Dispatcher rendered and auto-injected into DOCX seamlessly!")

    except Exception as e:
        print(f"\n[FAILURE] Test suite encountered error: {e}")
        import traceback
        traceback.print_exc()
        all_passed = False
    finally:
        shutil.rmtree(test_dir, ignore_errors=True)

    print("\n" + "=" * 65)
    if all_passed:
        print("ALL 11 PLANTUML & UNIFIED DISPATCHER TESTS PASSED!")
    else:
        print("TEST SUITE FAILED!")
    print("=" * 65)
    return all_passed


if __name__ == "__main__":
    success = run_tests()
    import sys
    sys.exit(0 if success else 1)
