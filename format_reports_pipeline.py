import sys
import os
import re
import shutil
import copy
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

sys.stdout.reconfigure(encoding='utf-8')

NAVY_COLOR = RGBColor(31, 73, 125)      # #1F497D (Heading 1, Heading 2, Major Titles)
DARK_NAVY = RGBColor(44, 62, 80)       # #2C3E50 (Heading 3)
DARK_SLATE = RGBColor(52, 73, 94)      # #34495E (Heading 4)
BODY_COLOR = RGBColor(0, 0, 0)         # #000000
MUTED_COLOR = RGBColor(100, 116, 139)  # #64748B (Muted text, metadata)

def apply_font_to_run(run, font_name="Calibri", size_pt=11.0, bold=None, italic=None, color_rgb=None):
    run.font.name = font_name
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color_rgb is not None:
        run.font.color.rgb = color_rgb
        
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def apply_font_to_paragraph(p, font_name="Calibri", size_pt=11.0, bold=None, italic=None, color_rgb=None):
    for r in p.runs:
        apply_font_to_run(r, font_name=font_name, size_pt=size_pt, bold=bold, italic=italic, color_rgb=color_rgb)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcMar'):
            tcPr.remove(child)
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D5DD", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    borders = parse_xml(f'''<w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="none"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>''')
    tblPr.append(borders)

def format_all_tables(doc):
    for tbl in doc.tables:
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, color="CBD5E1", sz="4")
        for r_idx, row in enumerate(tbl.rows):
            trPr = row._tr.get_or_add_trPr()
            if r_idx == 0:
                if not trPr.xpath('.//w:tblHeader'):
                    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            if not trPr.xpath('.//w:cantSplit'):
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
            for c_idx, cell in enumerate(row.cells):
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                if r_idx == 0:
                    set_cell_background(cell, "1F497D")
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(0)
                        for r in p.runs:
                            apply_font_to_run(r, font_name="Calibri", size_pt=10.0, bold=True, color_rgb=RGBColor(255, 255, 255))
                else:
                    bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(0)
                        for r in p.runs:
                            apply_font_to_run(r, font_name="Calibri", size_pt=10.0, color_rgb=BODY_COLOR)

def inject_dynamic_page_numbers(doc, font_name="Calibri", font_size_pt=10.0):
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)
        rPr.append(rFonts)
        
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(font_size_pt * 2)))
        rPr.append(sz)
        
        t = OxmlElement('w:t')
        t.text = "1"
        r.append(rPr)
        r.append(t)
        fldSimple.append(r)
        
        p._p.append(fldSimple)

def set_header_text(doc, header_text="SmartMarketBot – Capstone Project", font_name="Calibri", font_size_pt=9.5):
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(4.0)
        
        r = p.add_run(header_text)
        apply_font_to_run(r, font_name=font_name, size_pt=font_size_pt, italic=True, color_rgb=BODY_COLOR)

def set_tab_stop(paragraph, pos_twips=9050, align_val='right', leader_val='dot'):
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag.endswith('tabs'):
            pPr.remove(child)
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), align_val)
    tab.set(qn('w:leader'), leader_val)
    tab.set(qn('w:pos'), str(pos_twips))
    tabs.append(tab)
    pPr.append(tabs)

def create_toc_paragraph(doc, level, title, page_str, tab_pos_twips=9050):
    p = doc.add_paragraph()
    p.text = ""
    
    if level == 1:
        indent = Inches(0.0)
        is_bold = True
        size_pt = 11.0
        color_rgb = NAVY_COLOR
        space_before = 6.0
        space_after = 2.0
    elif level == 2:
        indent = Inches(0.18)
        is_bold = True
        size_pt = 11.0
        color_rgb = DARK_NAVY
        space_before = 3.0
        space_after = 2.0
    elif level == 3:
        indent = Inches(0.36)
        is_bold = False
        size_pt = 10.5
        color_rgb = BODY_COLOR
        space_before = 1.5
        space_after = 1.5
    else: # level 4
        indent = Inches(0.54)
        is_bold = False
        size_pt = 10.0
        color_rgb = BODY_COLOR
        space_before = 1.0
        space_after = 1.0
        
    p.paragraph_format.left_indent = indent
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    set_tab_stop(p, pos_twips=tab_pos_twips, align_val='right', leader_val='dot')
    
    r_title = p.add_run(title)
    apply_font_to_run(r_title, font_name="Calibri", size_pt=size_pt, bold=is_bold, color_rgb=color_rgb)
    
    r_tab = p.add_run('\t')
    apply_font_to_run(r_tab, font_name="Calibri", size_pt=size_pt, bold=False, color_rgb=MUTED_COLOR)
    
    r_page = p.add_run(page_str)
    apply_font_to_run(r_page, font_name="Calibri", size_pt=size_pt, bold=is_bold, color_rgb=color_rgb)
    
    return p

def format_figure_captions(doc):
    fig_pattern = re.compile(r'^(Figure|Table)\s+\d+(\.\d+)*', re.IGNORECASE)
    count = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if fig_pattern.match(t):
            # Normalize to Figure X.Y – Title
            t = re.sub(r'^(Figure\s+\d+(\.\d+)*)\s*[:\-–—]\s*', r'\1 – ', t, flags=re.IGNORECASE)
            t = re.sub(r'^(Table\s+\d+(\.\d+)*)\s*[:\-–—]\s*', r'\1 – ', t, flags=re.IGNORECASE)
            p.text = t
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4.0)
            p.paragraph_format.space_after = Pt(12.0)
            for r in p.runs:
                apply_font_to_run(r, font_name="Calibri", size_pt=10.0, bold=True, color_rgb=NAVY_COLOR)
            count += 1
    return count

print('Pipeline module loaded.')
