"""
Advanced DOCX Post-Processing & Normalization Engine
===================================================
Provides robust, production-grade document manipulation capabilities:
1. Dynamic OpenXML Page Numbering Injection (w:fldSimple w:instr="PAGE")
2. Academic & Enterprise Figure/Table Caption Normalization (Times New Roman 10.5pt, centered, En-dash)
3. Hierarchical Table of Contents (TOC) Construction with Dot Leaders & Dynamic Indentation
4. Context-Preserving Machine Translation with Quoted Term Protection (__QUOTED_N__)
5. Universal Emoji / Special Icon Sanitization
6. Windows File Lock Resilience (Graceful fallback when file is open in MS Word)
"""

import os
import re
import sys
import json
import time
import shutil
import urllib.request
import urllib.parse
from typing import List, Tuple, Dict, Optional, Any

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls


# ==========================================
# 1. DYNAMIC PAGE NUMBERING INJECTION
# ==========================================

def inject_dynamic_page_numbers(doc: docx.Document, font_name: str = "Times New Roman", font_size_pt: float = 10.0) -> None:
    """
    Injects dynamic OpenXML page numbering field (<w:fldSimple w:instr="PAGE"/>)
    into the footers of all sections in the document.
    """
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = ""
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
        
        # Add dynamic PAGE field
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)
        rPr.append(rFonts)
        
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(font_size_pt * 2)))
        rPr.append(sz)
        
        t = OxmlElement('w:t')
        t.text = "1"
        r.append(rPr)
        r.append(t)
        fldSimple.append(r)
        
        p._p.append(fldSimple)


# ==========================================
# 2. FIGURE & TABLE CAPTION NORMALIZATION
# ==========================================

def normalize_caption_text(text: str) -> str:
    """Standardizes Figure/Table caption format with En-dash ('–')."""
    t = text.strip()
    # Normalize Figure X.Y - Foo or Figure X.Y: Foo -> Figure X.Y – Foo
    t = re.sub(r'^(Figure\s+\d+(\.\d+)*)\s*[:\-–—]\s*', r'\1 – ', t, flags=re.IGNORECASE)
    t = re.sub(r'^(Table\s+\d+(\.\d+)*)\s*[:\-–—]\s*', r'\1 – ', t, flags=re.IGNORECASE)
    return t

def format_figure_captions(
    doc: docx.Document,
    font_name: str = "Times New Roman",
    font_size_pt: float = 10.5,
    space_before_pt: float = 4.0,
    space_after_pt: float = 12.0,
    color_rgb: Tuple[int, int, int] = (30, 41, 59)
) -> int:
    """
    Finds and formats all figure and table captions across the document.
    Ensures center alignment, correct typography, and controlled spacing.
    """
    count = 0
    fig_pattern = re.compile(r'^(Figure|Table)\s+\d+(\.\d+)*', re.IGNORECASE)
    
    for p in doc.paragraphs:
        t = p.text.strip()
        if fig_pattern.match(t):
            norm_text = normalize_caption_text(t)
            p.text = norm_text
            p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(space_before_pt)
            p.paragraph_format.space_after = Pt(space_after_pt)
            for r in p.runs:
                r.font.name = font_name
                r.font.size = Pt(font_size_pt)
                r.font.color.rgb = RGBColor(*color_rgb)
                r.font.italic = False
                # Ensure XML font elements
                rPr = r._r.get_or_add_rPr()
                rFonts = rPr.find(qn('w:rFonts'))
                if rFonts is None:
                    rFonts = OxmlElement('w:rFonts')
                    rPr.append(rFonts)
                rFonts.set(qn('w:ascii'), font_name)
                rFonts.set(qn('w:hAnsi'), font_name)
                rFonts.set(qn('w:cs'), font_name)
            count += 1
    return count


# ==========================================
# 3. TABLE OF CONTENTS REBUILDER
# ==========================================

def set_tab_stop(paragraph: docx.text.paragraph.Paragraph, pos_twips: int = 10469, align_val: str = 'right', leader_val: str = 'dot') -> None:
    """Sets a right-aligned tab stop with dot leaders on a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag.endswith('tabs'):
            pPr.remove(child)
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), align_val)
    tab.set(qn('w:leader'), leader_val)
    tab.set(qn('w:pos'), str(pos_twips))
    tabs.append(tab)
    pPr.append(tabs)

def rebuild_table_of_contents(
    doc: docx.Document,
    start_p_idx: int,
    toc_entries: List[Tuple[int, str, str]],  # [(level, title, page_str), ...]
    font_name: str = "Times New Roman",
    font_size_pt: float = 11.0,
    tab_pos_twips: int = 10469
) -> None:
    """
    Constructs a clean, perfectly aligned Table of Contents with dot leaders.
    Level 1: Indent 0.00in, Bold
    Level 2: Indent 0.15in, Bold
    Level 3: Indent 0.31in, Regular
    """
    for idx, (level, title, page) in enumerate(toc_entries):
        p_idx = start_p_idx + idx
        if p_idx < len(doc.paragraphs):
            p = doc.paragraphs[p_idx]
        else:
            p = doc.add_paragraph()
            
        p.text = ""
        if level == 1:
            indent = Inches(0.0)
            is_bold = True
        elif level == 2:
            indent = Inches(0.15)
            is_bold = True
        else:
            indent = Inches(0.31)
            is_bold = False
            
        p.paragraph_format.left_indent = indent
        p.paragraph_format.space_before = Pt(0.0)
        p.paragraph_format.space_after = Pt(5.0)
        
        set_tab_stop(p, pos_twips=tab_pos_twips, align_val='right', leader_val='dot')
        
        r_title = p.add_run(title)
        r_title.font.name = font_name
        r_title.font.size = Pt(font_size_pt)
        r_title.font.bold = is_bold
        
        r_tab = p.add_run('\t')
        r_tab.font.name = font_name
        r_tab.font.size = Pt(font_size_pt)
        r_tab.font.bold = False
        
        r_page = p.add_run(page)
        r_page.font.name = font_name
        r_page.font.size = Pt(font_size_pt)
        r_page.font.bold = is_bold


# ==========================================
# 4. EMOJI & SPECIAL SYMBOL SANITIZATION
# ==========================================

EMOJI_REGEX = re.compile(
    r"[\U00010000-\U0010ffff"  # SMP plane (emojis)
    r"\u2600-\u27bf"           # Misc symbols & Dingbats
    r"\u2300-\u23ff"           # Misc Technical
    r"\u2b50\u2b55\u2b06\u2b07"
    r"\u200d\ufe0f"            # Zero-width joiner & variation selectors
    r"]",
    re.UNICODE
)

DEFAULT_SUBSTITUTIONS = {
    '("✕")': '("X")',
    '"✕"': '"X"',
    '✕': 'X',
    '★': '',
    '🛒': '',
    '✨': '',
    '⭐': '',
    '👤': '',
    '🎙️': '',
    '🎙': '',
    '✓': '',
    '⏸': '',
    '⚠️': '',
    '⚠': '',
    '🔌': '',
    '🕒': '',
    '⛶': '',
}

def clean_string_symbols(text: str, custom_subs: Optional[Dict[str, str]] = None) -> str:
    """Removes emojis and symbols while preserving formatting and grammatical integrity."""
    if not text:
        return text
    
    subs = {**DEFAULT_SUBSTITUTIONS, **(custom_subs or {})}
    t = text
    for k, v in subs.items():
        t = t.replace(k, v)
        
    t = EMOJI_REGEX.sub('', t)
    
    # Clean up empty quotes and whitespace artifacts
    t = t.replace('""', '')
    t = re.sub(r'\"\s+', '"', t)
    t = re.sub(r'\s+\"', '"', t)
    t = re.sub(r'\(\s+\"', '("', t)
    t = re.sub(r'\"\s+\)', '")', t)
    t = re.sub(r'[ \t]{2,}', ' ', t)
    return t

def sanitize_emojis_and_symbols(doc: docx.Document, custom_subs: Optional[Dict[str, str]] = None) -> int:
    """Sanitizes emojis and symbols across all paragraphs and tables in the document."""
    cleaned_count = 0
    for p in doc.paragraphs:
        old_t = p.text
        new_t = clean_string_symbols(old_t, custom_subs)
        if new_t != old_t:
            if len(p.runs) == 1:
                p.runs[0].text = new_t
            else:
                for r in p.runs:
                    r.text = clean_string_symbols(r.text, custom_subs)
            cleaned_count += 1
            
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    old_t = p.text
                    new_t = clean_string_symbols(old_t, custom_subs)
                    if new_t != old_t:
                        if len(p.runs) == 1:
                            p.runs[0].text = new_t
                        else:
                            for r in p.runs:
                                r.text = clean_string_symbols(r.text, custom_subs)
                        cleaned_count += 1
    return cleaned_count


# ==========================================
# 5. QUOTE-SAFE TRANSLATION ENGINE
# ==========================================

VN_CHAR_REGEX = re.compile(r'[àáảãạăắằẳẵặâấầẩẫậđèéẻẽẹêếềểễệìíỉĩịòóỏõọôốồổỗộơớờởỡợùúủũụưứừửữựỳýỷỹỵÀÁẢÃẠĂẮẰẲẴẶÂẤẦẨẪẬĐÈÉẺẼẸÊẾỀỂỄỆÌÍỈĨỊÒÓỎÕỌÔỐỒỔỖỘƠỚỜỞỠỢÙÚỦŨỤƯỨỪỬỮỰỲÝỶỸỴ]')

def strip_quotes(text: str) -> str:
    """Strips all double and single quotes to detect unquoted non-English text."""
    t = re.sub(r'\"[^\"]*\"|“[^”]*”', '', text)
    t = re.sub(r'\'[^\']*\'|‘[^’]*’', '', t)
    return t

def translate_vn_to_en_protected(text: str, timeout_sec: float = 12.0) -> str:
    """
    Translates Vietnamese narrative text to English while strictly protecting
    all quoted substrings (\"...\" / '...') via __QUOTED_N__ placeholder masking.
    """
    if not text or not text.strip():
        return text
    
    # If no Vietnamese outside quotes, return as-is
    if not VN_CHAR_REGEX.search(strip_quotes(text)):
        return text
        
    quotes = []
    def repl_quote(match):
        quotes.append(match.group(0))
        return f"__QUOTED_{len(quotes)-1}__"
    
    # Mask double and single quotes
    t_protected = re.sub(r'\"[^\"]*\"|“[^”]*”', repl_quote, text)
    t_protected = re.sub(r'\'[^\']*\'|‘[^’]*’', repl_quote, t_protected)
    
    url = 'https://translate.googleapis.com/translate_a/single?client=gtx&sl=vi&tl=en&dt=t&q=' + urllib.parse.quote(t_protected)
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=timeout_sec) as resp:
                data = json.loads(resp.read().decode('utf-8'))
                translated = ''.join([item[0] for item in data[0] if item[0]])
                
                # Unmask placeholders
                for idx, q in enumerate(quotes):
                    translated = re.sub(rf'__QUOTED_{idx}__|_ _QUOTED_{idx}_ _|_QUOTED_{idx}_|__QUOTED_ {idx}__', q, translated)
                return translated
        except Exception:
            time.sleep(0.5)
            if attempt == 2:
                return text
    return text


# ==========================================
# 6. RESILIENT FILE SAVING (WORD LOCK HANDLING)
# ==========================================

def safe_save_docx(doc: docx.Document, target_path: str, fallback_suffix: str = "_updated") -> str:
    """
    Saves the Word document. If the file is locked exclusively by Microsoft Word
    (PermissionError Errno 13), gracefully saves to [target_path][fallback_suffix].docx
    to prevent data loss.
    """
    base, ext = os.path.splitext(target_path)
    fallback_path = f"{base}{fallback_suffix}{ext}"
    
    try:
        doc.save(target_path)
        return target_path
    except PermissionError:
        doc.save(fallback_path)
        print(f"[WARN] Target file '{target_path}' is locked by another process (e.g., Microsoft Word).")
        print(f"[INFO] Successfully saved updated document to fallback path: '{fallback_path}'")
        return fallback_path
