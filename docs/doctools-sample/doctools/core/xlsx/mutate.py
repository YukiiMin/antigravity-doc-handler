"""Sửa ô trên BẢN SAO; rào chắn nằm trong code."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, range_boundaries

from ...contract.models import Issue
from ..errors import DocToolsError


def _rc(coord: str) -> tuple[int, int]:
    col, row = coordinate_from_string(coord)
    return row, column_index_from_string(col)


def _locked(coord: str, sheet: str, zones: list[str]) -> str | None:
    row, col = _rc(coord)
    for z in zones:
        z_sheet, _, rng = z.rpartition("!")
        if z_sheet and z_sheet.strip("'") not in (sheet, "*"):
            continue
        c1, r1, c2, r2 = range_boundaries(rng)
        if r1 <= row <= r2 and c1 <= col <= c2:
            return z
    return None


def set_cells(src: Path, out: Path, sheet: str, updates: dict, locked_zones: list[str], allow_formulas: bool) -> list[Issue]:
    if out == src:
        raise DocToolsError("E-IO-OVERWRITE", "Không được ghi đè file nguồn.", fixable_by="ai",
                            suggested_action="Bỏ out_path hoặc chọn đường dẫn khác.")
    wb = load_workbook(src)
    if sheet not in wb.sheetnames:
        raise DocToolsError("E-SPEC-SHEET", f"Không có sheet '{sheet}'.", evidence={"available": wb.sheetnames},
                            fixable_by="ai", suggested_action="Chọn một tên trong 'available'.")
    ws = wb[sheet]

    errors: list[Issue] = []
    for coord in updates:                                   # kiểm toàn bộ trước, ghi sau (all-or-nothing)
        try:
            _rc(coord)
        except Exception:
            errors.append(Issue(code="E-SPEC-COORD", severity="error", message=f"Toạ độ ô không hợp lệ: {coord}",
                                location={"sheet": sheet, "cell": coord}, fixable_by="ai"))
            continue
        zone = _locked(coord, sheet, locked_zones)
        if zone:
            errors.append(Issue(code="E-LOCK-001", severity="error", message=f"Ô {coord} nằm trong vùng khóa {zone}.",
                                location={"sheet": sheet, "cell": coord}, evidence={"zone": zone}, fixable_by="human",
                                suggested_action="Không tự ghi; hỏi người dùng có mở khóa vùng này không."))
        if isinstance(ws[coord], MergedCell):
            top_left = next(r for r in ws.merged_cells.ranges if coord in r).start_cell.coordinate
            errors.append(Issue(code="E-MERGE-001", severity="error", message=f"{coord} là ô con của vùng gộp.",
                                location={"sheet": sheet, "cell": coord}, evidence={"top_left": top_left},
                                fixable_by="ai", suggested_action=f"Ghi vào ô trên-trái {top_left}."))
    if errors:
        return errors

    as_text: list[str] = []
    for coord, value in updates.items():
        cell = ws[coord]
        cell.value = value
        if isinstance(value, str) and value.startswith("=") and not allow_formulas:
            cell.data_type = "s"                            # openpyxl mặc định coi mọi chuỗi '=' là công thức
            as_text.append(coord)
    wb.calculation.fullCalcOnLoad = True                    # Excel tính lại khi mở (file openpyxl lưu không có cache)
    wb.save(out)

    info: list[Issue] = []
    if as_text:
        info.append(Issue(code="I-TEXT-EQ", severity="info", message="Chuỗi bắt đầu bằng '=' được lưu như văn bản.",
                          evidence={"cells": as_text}, suggested_action="Đặt allow_formulas=true nếu thật sự muốn công thức."))
    return info
