import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from helpers import add_zip_entry, make_fixture  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
from doctools.core.guards import make_file_ref  # noqa: E402
from doctools.registry import REGISTRY, run, tool_schemas  # noqa: E402
from doctools import agent_docs  # noqa: E402


class Base(unittest.TestCase):
    """Mỗi test tự có thư mục riêng và quyền truy cập riêng: không phụ thuộc thứ tự chạy."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name).resolve()
        self._env = mock.patch.dict(os.environ, {"DOCTOOLS_ROOTS": str(self.dir)})
        self._env.start()
        self.fx = make_fixture(self.dir)

    def tearDown(self):
        self._env.stop()
        self._tmp.cleanup()

    def codes(self, r):
        return {i.code for i in r.issues}


class TestContract(Base):
    def test_registry_is_complete(self):
        self.assertGreaterEqual(len(REGISTRY), 3)
        for s in tool_schemas():
            self.assertTrue(s["description"].strip())
            self.assertIn("properties", s["inputSchema"])
            self.assertIn("readOnlyHint", s["annotations"])

    def test_schema_error_is_structured(self):
        r = run("set_cells_xlsx", {"path": 123})
        self.assertFalse(r.success)
        self.assertEqual(r.issues[0].code, "E-SPEC-SCHEMA")
        self.assertIn("path", {i.location["path"] for i in r.issues})

    def test_unknown_tool_lists_available(self):
        r = run("khong_ton_tai", {})
        self.assertEqual(r.issues[0].code, "E-SPEC-UNKNOWN-TOOL")
        self.assertIn("preflight_xlsx", r.issues[0].evidence["available"])

    def test_agent_docs_in_sync(self):
        self.assertEqual(agent_docs.check(ROOT / ".agent"), [])


class TestPreflight(Base):
    def test_inventory(self):
        r = run("preflight_xlsx", {"path": str(self.fx)})
        inv = r.data["inventory"]
        self.assertTrue(r.success)
        self.assertEqual(r.data["fidelity_tier"], "T1")
        self.assertEqual((inv["images"], inv["formulas"], inv["merged_ranges"],
                          inv["conditional_formats"], inv["data_validations"]), (1, 1, 1, 1, 1))

    def test_macro_rejected(self):
        bad = add_zip_entry(self.fx, self.dir / "macro.xlsx", "xl/vbaProject.bin")
        r = run("preflight_xlsx", {"path": str(bad)})
        self.assertFalse(r.success)
        self.assertEqual(r.data["fidelity_tier"], "REJECT")
        self.assertIn("E-SEC-MACRO", self.codes(r))

    def test_zip_path_traversal_rejected(self):
        bad = add_zip_entry(self.fx, self.dir / "evil.xlsx", "../evil.txt")
        self.assertIn("E-SEC-ZIP", self.codes(run("preflight_xlsx", {"path": str(bad)})))

    def test_not_a_zip(self):
        junk = self.dir / "junk.xlsx"
        junk.write_text("không phải zip")
        self.assertIn("E-PKG-ZIP", self.codes(run("preflight_xlsx", {"path": str(junk)})))

    def test_path_outside_roots(self):
        with tempfile.TemporaryDirectory() as other:
            outside = make_fixture(Path(other))
            self.assertIn("E-SEC-PATH", self.codes(run("preflight_xlsx", {"path": str(outside)})))


class TestSetCells(Base):
    def test_preserves_components_and_original(self):
        sha_before = make_file_ref(self.fx).sha256
        r = run("set_cells_xlsx", {"path": str(self.fx), "sheet": "Data", "updates": {"B2": 10, "A1": "Họ tên"}})
        self.assertTrue(r.success, r.issues)
        self.assertNotIn("E-DIFF-UNDECLARED", self.codes(r))
        self.assertEqual(make_file_ref(self.fx).sha256, sha_before)          # nguồn không đổi
        self.assertEqual(load_workbook(r.file_ref.path)["Data"]["B2"].value, 10)
        self.assertIn("W-CALC-NO-CACHE", self.codes(r))

    def test_locked_zone_blocks_everything(self):
        r = run("set_cells_xlsx", {"path": str(self.fx), "sheet": "Data", "updates": {"A1": "x", "B2": 1},
                                   "locked_zones": ["Data!B1:B3"], "out_path": str(self.dir / "o.xlsx")})
        self.assertFalse(r.success)
        self.assertIn("E-LOCK-001", self.codes(r))
        self.assertFalse((self.dir / "o.xlsx").exists())                      # all-or-nothing

    def test_equals_string_is_text_unless_allowed(self):
        r = run("set_cells_xlsx", {"path": str(self.fx), "sheet": "Data", "updates": {"A2": "=cmd|' /C calc'!A0"}})
        cell = load_workbook(r.file_ref.path)["Data"]["A2"]
        self.assertEqual(cell.data_type, "s")
        self.assertIn("I-TEXT-EQ", self.codes(r))
        r2 = run("set_cells_xlsx", {"path": str(self.fx), "sheet": "Data", "updates": {"A2": "=1+1"},
                                    "allow_formulas": True, "out_path": str(self.dir / "f.xlsx")})
        self.assertEqual(load_workbook(r2.file_ref.path)["Data"]["A2"].data_type, "f")

    def test_merged_child_cell_gets_hint(self):
        r = run("set_cells_xlsx", {"path": str(self.fx), "sheet": "Data", "updates": {"B10": "x"}})
        issue = next(i for i in r.issues if i.code == "E-MERGE-001")
        self.assertEqual(issue.evidence["top_left"], "A10")

    def test_never_overwrites_input(self):
        r = run("set_cells_xlsx", {"path": str(self.fx), "sheet": "Data", "updates": {"A1": "x"}, "out_path": str(self.fx)})
        self.assertIn("E-IO-OVERWRITE", self.codes(r))

    def test_unknown_sheet_lists_available(self):
        r = run("set_cells_xlsx", {"path": str(self.fx), "sheet": "Nope", "updates": {"A1": 1}})
        self.assertEqual(r.issues[0].evidence["available"], ["Data"])


class TestDiff(Base):
    def test_detects_loss_and_accepts_declared(self):
        empty = self.dir / "empty.xlsx"
        wb = Workbook()
        wb.active.title = "Data"
        wb.save(empty)
        r = run("diff_inventory_xlsx", {"before": str(self.fx), "after": str(empty)})
        self.assertFalse(r.success)
        keys = {i.evidence["key"] for i in r.issues}
        self.assertTrue({"images", "conditional_formats", "data_validations", "merged_ranges", "formulas"} <= keys)
        declared = {k: r.data["after"][k] - r.data["before"][k] for k in keys}
        self.assertTrue(run("diff_inventory_xlsx", {"before": str(self.fx), "after": str(empty),
                                                    "declared_changes": declared}).success)


class TestAdapters(Base):
    def test_cli_matches_library(self):
        payload = json.dumps({"path": str(self.fx)})
        env = {**os.environ, "PYTHONPATH": str(ROOT)}
        p = subprocess.run([sys.executable, "-m", "doctools", "run", "preflight_xlsx", "--input", payload],
                           capture_output=True, text=True, env=env, cwd=str(ROOT))
        self.assertEqual(p.returncode, 0, p.stderr)
        self.assertEqual(json.loads(p.stdout)["data"], run("preflight_xlsx", {"path": str(self.fx)}).data)

    def test_cli_exit_code_two_on_error(self):
        env = {**os.environ, "PYTHONPATH": str(ROOT)}
        p = subprocess.run([sys.executable, "-m", "doctools", "run", "set_cells_xlsx", "--input",
                            json.dumps({"path": str(self.fx), "sheet": "Nope", "updates": {}})],
                           capture_output=True, text=True, env=env, cwd=str(ROOT))
        self.assertEqual(p.returncode, 2)


if __name__ == "__main__":
    unittest.main()
