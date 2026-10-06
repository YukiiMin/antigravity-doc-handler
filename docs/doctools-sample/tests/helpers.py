import os
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image as PILImage


def make_fixture(dirpath: Path) -> Path:
    """Workbook mẫu: công thức, merge, CF, DV, ảnh."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"], ws["B1"] = "Tên", "Điểm"
    ws["A2"], ws["B2"] = "An", 7
    ws["A3"], ws["B3"] = "Bình", 9
    ws["B4"] = "=SUM(B2:B3)"
    ws.merge_cells("A10:C11")
    ws["A10"] = "Ghi chú"
    ws.conditional_formatting.add("B2:B3", CellIsRule(operator="greaterThan", formula=["8"],
                                  fill=PatternFill("solid", start_color="FFFF0000", end_color="FFFF0000")))
    dv = DataValidation(type="list", formula1='"Đạt,Rớt"')
    dv.add("C2:C3")
    ws.add_data_validation(dv)
    png = dirpath / "logo.png"
    PILImage.new("RGB", (8, 8), "navy").save(png)
    ws.add_image(Image(str(png)), "E2")
    path = dirpath / "fixture.xlsx"
    wb.save(path)
    return path


def add_zip_entry(src: Path, dst: Path, name: str, data: bytes = b"x") -> Path:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for i in zin.infolist():
            zout.writestr(i, zin.read(i.filename))
        zout.writestr(name, data)
    return dst
