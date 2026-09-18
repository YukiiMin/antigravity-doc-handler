# -*- coding: utf-8 -*-
"""
create_example_test01.py — Precision Sheet Generator for example_test01 (NAV-SVC data).

Extracts 100% format tokens from 'Example' sheet in the template workbook.
Injects real test data extracted from 'NAV-SVC' sheet in input workbook.
Enforces Invariants E1, E3, E5, E6, E7, E8.

Target workbook: dataset1/Template/Report5_Unit Test-template.xlsx
Input workbook:  dataset1/input/Report5_Unit Test.xlsx (sheet NAV-SVC)
New sheet name:  example_test01
"""

from copy import copy
import datetime
import os
import sys
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding='utf-8')

TMPL_PATH = r"d:\Minh\For_myself\ZSCORT_GSU26_SAP05\tool\AI-Docx-Testing-Product\dataset1\Template\Report5_Unit Test-template.xlsx"
INPUT_PATH = r"d:\Minh\For_myself\ZSCORT_GSU26_SAP05\tool\AI-Docx-Testing-Product\dataset1\input\Report5_Unit Test.xlsx"


def safe_copy(obj):
    return copy(obj) if obj is not None else None


def extract_all_tokens(ws_ref):
    """E1: Extract all formatting tokens directly from Example reference sheet."""
    r9_tc      = ws_ref.cell(9, 6)    # Row 9 TC header (navy, 180° rot, double border top)
    r9_a       = ws_ref.cell(9, 1)    # Row 9 col A (navy fill + border l=double, t=double, align_v=bottom)
    r9_b       = ws_ref.cell(9, 2)    # Row 9 col B (navy fill, double border top)
    r9_c       = ws_ref.cell(9, 3)    # Row 9 col C (navy fill, double border top)
    r9_d       = ws_ref.cell(9, 4)    # Row 9 col D (navy fill, double border top)

    a10        = ws_ref.cell(10, 1)   # Col A Condition title (navy, bold, align_h=center, align_v=center)
    a15        = ws_ref.cell(15, 1)   # Col A mid-Condition fill row (navy)
    a33        = ws_ref.cell(33, 1)   # Col A Confirm title (navy, bold)
    a34        = ws_ref.cell(34, 1)   # Col A mid-Confirm fill row (navy)

    b10        = ws_ref.cell(10, 2)   # Col B Condition group header (bold, Tahoma 8pt, align_h=left, align_v=top)
    b15        = ws_ref.cell(15, 2)   # Col B Condition sub-row (bold, Tahoma 8pt, align_h=left, align_v=top)
    b33        = ws_ref.cell(33, 2)   # Col B Confirm Return header
    b34        = ws_ref.cell(34, 2)   # Col B Confirm sub-label
    b40        = ws_ref.cell(40, 2)   # Col B Exception header

    d15        = ws_ref.cell(15, 4)   # Col D Condition data value (Tahoma 8pt non-bold, align_h=right, align_v=top)
    d33        = ws_ref.cell(33, 4)   # Col D Return header row
    d35        = ws_ref.cell(35, 4)   # Col D Confirm data value (Tahoma 8pt non-bold)
    d40        = ws_ref.cell(40, 4)   # Col D Exception header row

    b45        = ws_ref.cell(45, 2)   # Col B Type label (Tahoma 8pt non-bold, align_h=left, v=bottom, border T=double, B=thin)
    b46        = ws_ref.cell(46, 2)   # Col B PF label (Tahoma 8pt non-bold, align_h=left, v=bottom, border T=thin, B=thin)
    b47        = ws_ref.cell(47, 2)   # Col B Date label (Tahoma 8pt non-bold, align_h=left, v=top, border T=thin, B=thin)
    b48        = ws_ref.cell(48, 2)   # Col B Defect label (Tahoma 8pt non-bold, align_h=left, v=top, border T=thin, B=double)

    mark_cell  = ws_ref.cell(15, 6)   # O-mark matrix cell (Courier New 12pt bold, align center, thin border)
    type_r     = ws_ref.cell(45, 6)   # Type row cell (Courier New 8pt non-bold)
    pf_r       = ws_ref.cell(46, 6)   # PF row cell (Courier New 8pt non-bold)
    date_r     = ws_ref.cell(47, 6)   # Date row cell (Tahoma 8pt, textRotation=255, mm/dd)

    return {
        # Row 9 full tokens
        'r9_font':       safe_copy(r9_tc.font),
        'r9_fill':       safe_copy(r9_tc.fill),
        'r9_align':      safe_copy(r9_tc.alignment),
        'r9_border':     safe_copy(r9_tc.border),

        'r9_a_font':     safe_copy(r9_a.font),
        'r9_a_fill':     safe_copy(r9_a.fill),
        'r9_a_align':    safe_copy(r9_a.alignment),
        'r9_a_border':   safe_copy(r9_a.border),

        'r9_b_font':     safe_copy(r9_b.font),
        'r9_b_fill':     safe_copy(r9_b.fill),
        'r9_b_align':    safe_copy(r9_b.alignment),
        'r9_b_border':   safe_copy(r9_b.border),

        'r9_c_font':     safe_copy(r9_c.font),
        'r9_c_fill':     safe_copy(r9_c.fill),
        'r9_c_align':    safe_copy(r9_c.alignment),
        'r9_c_border':   safe_copy(r9_c.border),

        'r9_d_align':    safe_copy(r9_d.alignment),

        # Col A tokens
        'a_cond_font':       safe_copy(a10.font),
        'a_cond_fill':       safe_copy(a10.fill),
        'a_cond_align':      safe_copy(a10.alignment),
        'a_cond_border':     safe_copy(a10.border),

        'a_mid_cond_font':   safe_copy(a15.font),
        'a_mid_cond_fill':   safe_copy(a15.fill),
        'a_mid_cond_align':  safe_copy(a15.alignment),
        'a_mid_cond_border': safe_copy(a15.border),

        'a_conf_font':       safe_copy(a33.font),
        'a_conf_fill':       safe_copy(a33.fill),
        'a_conf_align':      safe_copy(a33.alignment),
        'a_conf_border':     safe_copy(a33.border),

        'a_mid_conf_font':   safe_copy(a34.font),
        'a_mid_conf_fill':   safe_copy(a34.fill),
        'a_mid_conf_align':  safe_copy(a34.alignment),
        'a_mid_conf_border': safe_copy(a34.border),

        # Col B label tokens
        'b_top_font':        safe_copy(b10.font),
        'b_top_fill':        safe_copy(b10.fill),
        'b_top_align':       safe_copy(b10.alignment),

        'b_conf_head_font':  safe_copy(b33.font),
        'b_conf_head_fill':  safe_copy(b33.fill),
        'b_conf_head_align': safe_copy(b33.alignment),

        'b_sub_font':        safe_copy(b34.font),
        'b_sub_fill':        safe_copy(b34.fill),
        'b_sub_align':       safe_copy(b34.alignment),

        'b_ex_head_font':    safe_copy(b40.font),
        'b_ex_head_fill':    safe_copy(b40.fill),
        'b_ex_head_align':   safe_copy(b40.alignment),

        # Col B Result labels (Example B45..B48: Tahoma 8pt non-bold, align_h=left)
        'res_type_label_font':   safe_copy(b45.font),
        'res_type_label_fill':   safe_copy(b45.fill),
        'res_type_label_align':  safe_copy(b45.alignment),
        'res_type_label_border': safe_copy(b45.border),

        'res_pf_label_font':     safe_copy(b46.font),
        'res_pf_label_fill':     safe_copy(b46.fill),
        'res_pf_label_align':    safe_copy(b46.alignment),
        'res_pf_label_border':   safe_copy(b46.border),

        'res_date_label_font':   safe_copy(b47.font),
        'res_date_label_fill':   safe_copy(b47.fill),
        'res_date_label_align':  safe_copy(b47.alignment),
        'res_date_label_border': safe_copy(b47.border),

        'res_defect_label_font': safe_copy(b48.font),
        'res_defect_label_fill': safe_copy(b48.fill),
        'res_defect_label_align':safe_copy(b48.alignment),
        'res_defect_label_border':safe_copy(b48.border),

        # Col D label tokens
        'd_top_font':        safe_copy(d15.font),
        'd_top_fill':        safe_copy(d15.fill),
        'd_top_align':       safe_copy(d15.alignment),

        'd_conf_head_font':  safe_copy(d33.font),
        'd_conf_head_fill':  safe_copy(d33.fill),
        'd_conf_head_align': safe_copy(d33.alignment),

        'd_sub_font':        safe_copy(d35.font),
        'd_sub_fill':        safe_copy(d35.fill),
        'd_sub_align':       safe_copy(d35.alignment),

        'd_ex_head_font':    safe_copy(d40.font),
        'd_ex_head_fill':    safe_copy(d40.fill),
        'd_ex_head_align':   safe_copy(d40.alignment),

        # Matrix & Result tokens
        'mark_font':   safe_copy(mark_cell.font),
        'mark_align':  safe_copy(mark_cell.alignment),
        'mark_border': safe_copy(mark_cell.border),

        'type_font':   safe_copy(type_r.font),
        'type_align':  safe_copy(type_r.alignment),
        'type_border': safe_copy(type_r.border),

        'pf_font':     safe_copy(pf_r.font),
        'pf_align':    safe_copy(pf_r.alignment),
        'pf_border':   safe_copy(pf_r.border),

        'date_font':   safe_copy(date_r.font),
        'date_align':  safe_copy(date_r.alignment),
        'date_border': safe_copy(date_r.border),
        'date_numfmt': date_r.number_format,
    }


def apply(cell, font=None, fill=None, align=None, border=None, numfmt=None):
    if font:   cell.font = copy(font)
    if fill:   cell.fill = copy(fill)
    if align:  cell.alignment = copy(align)
    if border: cell.border = copy(border)
    if numfmt: cell.number_format = numfmt


def generate_example_test01(wb_tmpl, wb_in):
    sheet_name = 'example_test01'
    if sheet_name in wb_tmpl.sheetnames:
        del wb_tmpl[sheet_name]

    ws_ex = wb_tmpl['Example']
    ws_in = wb_in['NAV-SVC']
    t = extract_all_tokens(ws_ex)

    # Clone Example sheet
    ws = wb_tmpl.copy_worksheet(ws_ex)
    ws.title = sheet_name

    # Remove stale merges below row 8
    for m in list(ws.merged_cells.ranges):
        if m.min_row > 8:
            ws.merged_cells.remove(m)

    # Clear all cells below row 8
    for r in range(9, 65):
        for c in range(1, 25):
            cell = ws.cell(r, c)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            cell.alignment = Alignment()
            cell.font = Font()
            cell.number_format = 'General'

    # Column widths from Example (E1 / E5)
    for c in range(1, 22):
        col = get_column_letter(c)
        ex_w = ws_ex.column_dimensions[col].width
        ws.column_dimensions[col].width = ex_w if ex_w else 8.43

    # Metadata rows 2-5 extracted from NAV-SVC
    ws.cell(2, 3).value  = ws_in['C2'].value    # 'NAV-SVC'
    ws.cell(2, 12).value = ws_in['L2'].value    # 'NavigationService (PlanRouteAsync)'
    ws.cell(3, 3).value  = ws_in['C3'].value    # 'SmartMarketBot Team'
    ws.cell(3, 6).value  = ws_in['F3'].value or 'SmartMarketBot Team'
    ws.cell(4, 3).value  = ws_in['C4'].value    # 260
    ws.cell(5, 3).value  = ws_in['C5'].value    # Requirement / Purpose

    # NAV-SVC has 5 TCs: F(6)..J(10)
    TC_COUNT = 5
    MAX_TC_COL = 6 + TC_COUNT - 1   # Col 10 (J)
    LAST_TC = get_column_letter(MAX_TC_COL)

    # Section row boundaries
    COND_START = 10
    COND_END   = 19     # 10 rows: 10..19
    CONF_START = 20
    CONF_END   = 28     # 9 rows: 20..28
    RES_START  = 29
    TYPE_ROW   = 29
    PF_ROW     = 30
    DATE_ROW   = 31
    DEFECT_ROW = 32

    # Row 7 Live KPI Formulas (E3)
    ws.cell(7, 1).value  = f'=COUNTIF(F{PF_ROW}:{LAST_TC}{PF_ROW},"P")'
    ws.cell(7, 3).value  = f'=COUNTIF(F{PF_ROW}:{LAST_TC}{PF_ROW},"F")'
    ws.cell(7, 6).value  = '=SUM(O7,-A7,-C7)'
    ws.cell(7, 12).value = f'=COUNTIF(F{TYPE_ROW}:{LAST_TC}{TYPE_ROW},"N")'
    ws.cell(7, 13).value = f'=COUNTIF(F{TYPE_ROW}:{LAST_TC}{TYPE_ROW},"A")'
    ws.cell(7, 14).value = f'=COUNTIF(F{TYPE_ROW}:{LAST_TC}{TYPE_ROW},"B")'
    ws.cell(7, 15).value = f'=COUNTA(F9:{LAST_TC}9)'

    # Row 9 Header formatting
    ws.row_dimensions[9].height = 45.0
    apply(ws.cell(9, 1), font=t['r9_a_font'], fill=t['r9_a_fill'],
          align=t['r9_a_align'], border=t['r9_a_border'])
    apply(ws.cell(9, 2), font=t['r9_b_font'], fill=t['r9_b_fill'],
          align=t['r9_b_align'], border=t['r9_b_border'])
    for c in [3, 5]:
        apply(ws.cell(9, c), font=t['r9_c_font'], fill=t['r9_c_fill'],
              align=t['r9_c_align'], border=t['r9_c_border'])
    apply(ws.cell(9, 4), font=t['r9_c_font'], fill=t['r9_c_fill'],
          align=t['r9_d_align'], border=t['r9_c_border'])

    TC_IDS = [f'UTCID{i:02d}' for i in range(1, TC_COUNT + 1)]
    for i, tcid in enumerate(TC_IDS, start=6):
        cell = ws.cell(9, i)
        cell.value = tcid
        apply(cell, font=t['r9_font'], fill=t['r9_fill'],
              align=t['r9_align'], border=t['r9_border'])

    # Helper: Col A section styling
    def set_col_a(r, value=None, zone='cond'):
        cell = ws.cell(r, 1)
        cell.value = value
        if zone == 'cond':
            if value:
                apply(cell, font=t['a_cond_font'], fill=t['a_cond_fill'],
                      align=t['a_cond_align'], border=t['a_cond_border'])
            else:
                apply(cell, font=t['a_mid_cond_font'], fill=t['a_mid_cond_fill'],
                      align=t['a_mid_cond_align'], border=t['a_mid_cond_border'])
        else:
            if value:
                apply(cell, font=t['a_conf_font'], fill=t['a_conf_fill'],
                      align=t['a_conf_align'], border=t['a_conf_border'])
            else:
                apply(cell, font=t['a_mid_conf_font'], fill=t['a_mid_conf_fill'],
                      align=t['a_mid_conf_align'], border=t['a_mid_conf_border'])

    def set_b(r, text, style='top'):
        cell = ws.cell(r, 2)
        cell.value = text
        if style == 'top':
            apply(cell, font=t['b_top_font'], fill=t['b_top_fill'], align=t['b_top_align'])
        elif style == 'conf_head':
            apply(cell, font=t['b_conf_head_font'], fill=t['b_conf_head_fill'], align=t['b_conf_head_align'])
        elif style == 'ex_head':
            apply(cell, font=t['b_ex_head_font'], fill=t['b_ex_head_fill'], align=t['b_ex_head_align'])
        elif style == 'res_type':
            apply(cell, font=t['res_type_label_font'], fill=t['res_type_label_fill'],
                  align=t['res_type_label_align'], border=t['res_type_label_border'])
        elif style == 'res_pf':
            apply(cell, font=t['res_pf_label_font'], fill=t['res_pf_label_fill'],
                  align=t['res_pf_label_align'], border=t['res_pf_label_border'])
        elif style == 'res_date':
            apply(cell, font=t['res_date_label_font'], fill=t['res_date_label_fill'],
                  align=t['res_date_label_align'], border=t['res_date_label_border'])
        elif style == 'res_defect':
            apply(cell, font=t['res_defect_label_font'], fill=t['res_defect_label_fill'],
                  align=t['res_defect_label_align'], border=t['res_defect_label_border'])
        else:
            apply(cell, font=t['b_sub_font'], fill=t['b_sub_fill'], align=t['b_sub_align'])

    def set_d(r, text, style='top'):
        cell = ws.cell(r, 4)
        cell.value = text
        if style == 'top':
            apply(cell, font=t['d_top_font'], fill=t['d_top_fill'], align=t['d_top_align'])
        elif style == 'conf_head':
            apply(cell, font=t['d_conf_head_font'], fill=t['d_conf_head_fill'], align=t['d_conf_head_align'])
        elif style == 'ex_head':
            apply(cell, font=t['d_ex_head_font'], fill=t['d_ex_head_fill'], align=t['d_ex_head_align'])
        else:
            apply(cell, font=t['d_sub_font'], fill=t['d_sub_fill'], align=t['d_sub_align'])

    def mark_o(r, *cols):
        for c in cols:
            cell = ws.cell(r, c)
            cell.value = 'O'
            apply(cell, font=t['mark_font'], align=t['mark_align'], border=t['mark_border'])

    def grid(r, c_start, c_end):
        for c in range(c_start, c_end + 1):
            if ws.cell(r, c).value != 'O':
                apply(ws.cell(r, c), border=t['mark_border'])

    def row_h(r, h=13.5):
        ws.row_dimensions[r].height = h

    # ════════════════════════════════════════════════════════════════════
    # CONDITION block: rows 10-19 (E7 & E8 compliant)
    # ════════════════════════════════════════════════════════════════════
    set_col_a(COND_START, 'Condition', zone='cond')
    for r in range(COND_START, CONF_START):
        if r != COND_START:
            set_col_a(r, zone='cond')
        row_h(r)

        # Invariant E7: Initialize B-C-D unified 3-column box
        is_last = (r == CONF_START - 1)
        b_side = Side(style='double') if is_last else Side(style='thin')
        apply(ws.cell(r, 2), font=t['b_top_font'], fill=t['b_top_fill'], align=t['b_top_align'],
              border=Border(left=Side(style='thin'), top=Side(style='thin'), bottom=b_side))
        apply(ws.cell(r, 3), font=t['b_top_font'], fill=t['b_top_fill'], align=t['b_top_align'],
              border=Border(top=Side(style='thin'), bottom=b_side))
        apply(ws.cell(r, 4), font=t['d_top_font'], fill=t['d_top_fill'], align=t['d_top_align'],
              border=Border(right=Side(style='thin'), top=Side(style='thin'), bottom=b_side))
        if is_last:
            apply(ws.cell(r, 5), border=Border(bottom=b_side))

    # Precondition group (rows 10-14)
    set_b(10, 'Precondition', style='top')
    set_d(10, '', style='top')
    set_d(11, 'InMemory AppDbContext with Guid isolation', style='top')
    set_d(12, 'ILocalizationService mocked (key passthrough)', style='top')
    set_d(13, 'NavigationNodes seeded (2-3 nodes per test)', style='top')
    set_d(14, 'NavigationEdges seeded (bidirectional)', style='top')

    # Input Parameters group (rows 15-19) — following Function 2 / Example convention
    set_b(15, 'Input Parameters', style='top')
    set_d(15, 'StartNodeId=1, EndNodeId=3', style='top')
    set_d(16, 'StartNodeId=1, EndNodeId=999 (non-existent)', style='top')
    set_d(17, 'Node 1 IsBlocked=true', style='top')
    set_d(18, 'No edges between nodes', style='top')
    set_d(19, 'Two Maps (MapId=1 & 9999); FloorId=1 filter', style='top')

    # 'O' marks from NAV-SVC input
    mark_o(11, 6, 7, 8, 9, 10)
    mark_o(12, 6, 7, 8, 9, 10)
    mark_o(13, 6, 7, 8, 9, 10)
    mark_o(14, 6, 9)
    mark_o(15, 6)
    mark_o(16, 7)
    mark_o(17, 8)
    mark_o(18, 9)
    mark_o(19, 10)

    for r in range(COND_START, CONF_START):
        grid(r, 6, MAX_TC_COL)

    # Condition block outer closing border: double bottom on row 19 (cols B..J)
    _dbl = Side(style='double')
    for c in range(2, MAX_TC_COL + 1):
        cell = ws.cell(CONF_START - 1, c)
        ex = cell.border
        cell.border = Border(left=ex.left, right=ex.right, top=ex.top, bottom=_dbl)

    # ════════════════════════════════════════════════════════════════════
    # CONFIRM block: rows 20-28 (E7 & E8 compliant)
    # ════════════════════════════════════════════════════════════════════
    set_col_a(CONF_START, 'Confirm', zone='conf')
    for r in range(CONF_START, RES_START):
        if r != CONF_START:
            set_col_a(r, zone='conf')
        row_h(r)

        # Invariant E7: Initialize B-C-D unified 3-column box
        is_first = (r == CONF_START)
        t_side = None if is_first else Side(style='thin')
        apply(ws.cell(r, 2), font=t['b_sub_font'], fill=t['b_sub_fill'], align=t['b_sub_align'],
              border=Border(left=Side(style='thin'), top=t_side, bottom=Side(style='thin')))
        apply(ws.cell(r, 3), font=t['b_sub_font'], fill=t['b_sub_fill'], align=t['b_sub_align'],
              border=Border(top=t_side, bottom=Side(style='thin')))
        apply(ws.cell(r, 4), font=t['d_sub_font'], fill=t['d_sub_fill'], align=t['d_sub_align'],
              border=Border(right=Side(style='thin'), top=t_side, bottom=Side(style='thin')))
        apply(ws.cell(r, 5), border=Border(left=Side(style='thin'), right=Side(style='thin'), top=t_side, bottom=Side(style='thin')))

    # Return group (rows 20-25)
    set_b(20, 'Return', style='conf_head')
    set_d(20, '', style='conf_head')
    set_d(21, 'TotalDistance=20.0; Nodes=[1,2,3]', style='sub')
    set_d(22, "Throw InvalidOperationException('StartEndNodeNotExist')", style='sub')
    set_d(23, "Throw InvalidOperationException('StartEndNodeBlocked')", style='sub')
    set_d(24, 'TotalDistance=0; Nodes=[]', style='sub')
    set_d(25, 'Route=[1,2]; TotalDistance=1', style='sub')

    # Exception group (rows 26-28)
    set_b(26, 'Exception', style='ex_head')
    set_d(26, '', style='ex_head')
    set_d(27, 'InvalidOperationException - StartEndNodeNotExist', style='sub')
    set_d(28, 'InvalidOperationException - StartEndNodeBlocked', style='sub')

    # 'O' marks from NAV-SVC input
    mark_o(21, 6)
    mark_o(22, 7)
    mark_o(23, 8)
    mark_o(24, 9)
    mark_o(25, 10)
    mark_o(27, 7)
    mark_o(28, 8)

    for r in range(CONF_START, RES_START):
        grid(r, 6, MAX_TC_COL)

    # ════════════════════════════════════════════════════════════════════
    # RESULT block: rows 29-32
    # ════════════════════════════════════════════════════════════════════
    set_col_a(RES_START, 'Result', zone='conf')
    for r in range(RES_START, DEFECT_ROW + 1):
        if r != RES_START:
            set_col_a(r, zone='conf')
        row_h(r)

    # Result footer labels (Example B45..B48: Tahoma 8pt non-bold, align_h=left)
    set_b(TYPE_ROW,   'Type(N : Normal, A : Abnormal, B : Boundary)', style='res_type')
    set_b(PF_ROW,     'Passed/Failed', style='res_pf')
    set_b(DATE_ROW,   'Executed Date', style='res_date')
    set_b(DEFECT_ROW, 'Defect ID', style='res_defect')

    # Result merges B:D for each footer row + borders on C & D
    ws.cell(TYPE_ROW, 3).border = Border(top=Side(style='double'), bottom=Side(style='thin'))
    ws.cell(TYPE_ROW, 4).border = Border(right=Side(style='thin'), top=Side(style='double'), bottom=Side(style='thin'))

    ws.cell(PF_ROW, 3).border = Border(top=Side(style='thin'), bottom=Side(style='thin'))
    ws.cell(PF_ROW, 4).border = Border(right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    ws.cell(DATE_ROW, 3).border = Border(top=Side(style='thin'), bottom=Side(style='thin'))
    ws.cell(DATE_ROW, 4).border = Border(right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    ws.cell(DEFECT_ROW, 3).border = Border(top=Side(style='thin'), bottom=Side(style='double'))
    ws.cell(DEFECT_ROW, 4).border = Border(right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='double'))

    for fr in [TYPE_ROW, PF_ROW, DATE_ROW, DEFECT_ROW]:
        ws.merge_cells(f'B{fr}:D{fr}')

    # Type values from NAV-SVC input
    type_vals = ['N', 'A', 'A', 'N', 'B']
    for i, v in enumerate(type_vals):
        cell = ws.cell(TYPE_ROW, 6 + i)
        cell.value = v
        apply(cell, font=t['type_font'], align=t['type_align'], border=t['type_border'])

    # PF values from NAV-SVC input
    pf_vals = ['P', 'P', 'P', 'P', 'P']
    for i, v in enumerate(pf_vals):
        cell = ws.cell(PF_ROW, 6 + i)
        cell.value = v
        apply(cell, font=t['pf_font'], align=t['pf_align'], border=t['pf_border'])

    # Executed Dates from NAV-SVC input
    dates = [
        datetime.datetime(2026, 9, 4),
        datetime.datetime(2026, 9, 1),
        datetime.datetime(2026, 9, 9),
        datetime.datetime(2026, 9, 8),
        datetime.datetime(2026, 9, 8),
    ]
    for i, d in enumerate(dates):
        cell = ws.cell(DATE_ROW, 6 + i)
        cell.value = d
        apply(cell, font=t['date_font'], align=t['date_align'],
              border=t['date_border'], numfmt=t['date_numfmt'])

    # Defect ID row (empty, grid borders)
    for c in range(6, MAX_TC_COL + 1):
        apply(ws.cell(DEFECT_ROW, c), border=t['date_border'])

    # Result block outer closing double border on DEFECT_ROW (cols A..J)
    for c in range(1, MAX_TC_COL + 1):
        cell = ws.cell(DEFECT_ROW, c)
        ex = cell.border
        cell.border = Border(left=ex.left, right=ex.right, top=ex.top, bottom=_dbl)

    # ── Data Validations ──────────────────────────────────────────────────
    dv_o = DataValidation(type='list', formula1='"O"', allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_o)
    dv_o.add(f'F{COND_START}:{LAST_TC}{RES_START - 1}')

    dv_nab = DataValidation(type='list', formula1='"N,A,B"', allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_nab)
    dv_nab.add(f'F{TYPE_ROW}:{LAST_TC}{TYPE_ROW}')

    dv_pf = DataValidation(type='list', formula1='"P,F"', allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_pf)
    dv_pf.add(f'F{PF_ROW}:{LAST_TC}{PF_ROW}')

    return ws, ws_ex, MAX_TC_COL


def main():
    print(f"Loading template: {TMPL_PATH}")
    wb_tmpl = openpyxl.load_workbook(TMPL_PATH, data_only=False)

    print(f"Loading input: {INPUT_PATH}")
    wb_in = openpyxl.load_workbook(INPUT_PATH, data_only=False)

    print("Generating sheet 'example_test01' with NAV-SVC data...")
    ws_tgt, ws_ref, max_tc = generate_example_test01(wb_tmpl, wb_in)

    # E6: Run 12-Gate Diff Engine before saving
    print("\n[E6] Running 12-Gate Format Diff Engine...")
    try:
        from tools.format_diff_excel import run_diff
    except ImportError:
        from format_diff_excel import run_diff

    diffs = run_diff(ws_tgt, ws_ref, verbose=True)

    if diffs:
        print(f"\n✗ {len(diffs)} format diffs found — Aborting save:")
        for d in diffs:
            print(d)
        return False

    print(f"\n✓ 12-Gate Format Diff Clean — 0 issues found. Proceeding to save.")

    try:
        wb_tmpl.save(TMPL_PATH)
    except PermissionError:
        import subprocess, time
        subprocess.run(['taskkill', '/F', '/IM', 'excel.exe'], capture_output=True)
        time.sleep(0.5)
        wb_tmpl.save(TMPL_PATH)

    print(f"\nSUCCESS: Sheet 'example_test01' delivered in '{TMPL_PATH}'.")
    print(f"Sheet names: {wb_tmpl.sheetnames}")
    return True


if __name__ == '__main__':
    main()
