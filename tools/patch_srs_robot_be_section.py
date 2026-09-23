import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

MASTER_DOCX = r"e:\Do_an_SU26\SuperMarketBot-BE\docs\Report3_Software Requirement Specification.docx"
MIRROR_DOCX = r"e:\Do_an_SU26\SuperMarketBot-Android-Robot\screenShot\Report3_Software Requirement Specification.docx"
IMG_DIR = r"e:\Do_an_SU26\SuperMarketBot-Android-Robot\screenShot"

def img(name):
    p = os.path.join(IMG_DIR, name)
    if not os.path.exists(p):
        print(f"WARNING: Image not found at {p}")
    return p

print("=" * 70)
print(f"Opening Master Document: {MASTER_DOCX}")
doc = docx.Document(MASTER_DOCX)
print(f"Total Paragraphs: {len(doc.paragraphs)}")

# 1. Update Table of Contents (paragraphs 28-41)
print("\n[Step 1] Synchronizing Table of Contents...")
toc_items = [
    "3. Functional Requirements\t8",
    "3.1 System Functional Overview\t8",
    "3.2 App for Robot\t10",
    "    3.2.1 System Startup & Idle Kiosk Mode\t10",
    "    3.2.2 User Onboarding & Face ID Authentication\t10",
    "    3.2.3 Guest Shopping Experience & Daily Promotions\t11",
    "    3.2.4 Member Personalization & Special Offers\t11",
    "    3.2.5 Smart Search & AI Voice Assistant\t12",
    "    3.2.6 Product Detail & Instant Cart Operations\t13",
    "    3.2.7 Smart Cart & Autonomous Robot Navigation Guide\t13",
    "    3.2.8 Contextual Overlays & Autonomous Advertising\t14",
    "3.3 Android Customer App\t15",
    "3.4 Admin Web Dashboard\t17",
    "3.5 Android Staff App\t21"
]

for idx, item in enumerate(toc_items):
    p_idx = 28 + idx
    if p_idx < len(doc.paragraphs):
        p = doc.paragraphs[p_idx]
        p.text = item
        if item.startswith("    "):
            p.paragraph_format.left_indent = Inches(0.2)
        else:
            p.paragraph_format.left_indent = Inches(0.0)

print("Table of Contents synchronized.")

# 2. Locate Section 3.2 and Section 3.3
print("\n[Step 2] Locating Section 3.2 and Section 3.3 boundaries...")
idx_3_2 = -1
idx_3_3 = -1

for i in range(50, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    txt = p.text.strip()
    if txt == "3.2 App for Robot":
        idx_3_2 = i
    elif idx_3_2 != -1 and ("3.3 Android Customer App" in txt or "3.3 App for Customer" in txt):
        idx_3_3 = i
        break

print(f"Found Section 3.2 at index: {idx_3_2}")
print(f"Found Section 3.3 at index: {idx_3_3}")

if idx_3_2 == -1 or idx_3_3 == -1:
    print("FATAL ERROR: Could not locate Section 3.2 or 3.3")
    sys.exit(1)

# Remove all existing paragraphs strictly between 3.2 and 3.3
p_3_3 = doc.paragraphs[idx_3_3]
paragraphs_to_remove = []
for k in range(idx_3_2 + 1, idx_3_3):
    paragraphs_to_remove.append(doc.paragraphs[k])

for p_rem in paragraphs_to_remove:
    p_elem = p_rem._p
    p_parent = p_elem.getparent()
    if p_parent is not None:
        p_parent.remove(p_elem)

print(f"Cleared {len(paragraphs_to_remove)} flawed/misplaced paragraphs in Section 3.2.")

# Helper functions for structured styling
def add_h4(title):
    p = p_3_3.insert_paragraph_before(title, style='Heading 4')
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_p(text="", bold_prefix=None, red_marker=False, italic=False):
    p = p_3_3.insert_paragraph_before(style='normal')
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(2.5)
    if red_marker:
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(220, 38, 38)
    elif bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        if text:
            p.add_run(" " + text)
    else:
        r = p.add_run(text)
        if italic:
            r.italic = True
    return p

def add_images(image_paths, width_inches=2.5):
    p = p_3_3.insert_paragraph_before(style='normal')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    for idx, path in enumerate(image_paths):
        if os.path.exists(path):
            r = p.add_run()
            r.add_picture(path, width=Inches(width_inches))
            if idx < len(image_paths) - 1:
                p.add_run("   ")
        else:
            print(f"Warning: Image not found: {path}")
    return p

print("\n[Step 3] Rebuilding Section 3.2 with verified screenshots & enhanced business logic...")

# ==============================================================================
# 3.2.1 System Startup & Idle Kiosk Mode
# ==============================================================================
add_h4("3.2.1 System Startup & Idle Kiosk Mode")
add_images([img("01_Welcome_Kiosk_Idle.jpg")], width_inches=2.6)
add_p("[Figure 3.2.1: Welcome Screen & Standby Kiosk Idle Mode]", italic=True)

add_h4("Function Trigger")
add_p("System initiates automatically when the robot hardware powers on or the Android Robot Kiosk application (com.truongduyyyy.mobilerobot) boots up. Following a 3-second diagnostic sequence, the system enters the Welcome Screen. Standby Advertising Mode triggers automatically after 60 seconds of continuous user inactivity.", bold_prefix="System Startup & Standby Trigger:")
add_p("Tap anywhere on the central interactive screen area or click the \"Khuyến mãi hôm nay\" (Today's Promotions) button at the bottom to initiate the shopping interaction.", bold_prefix="User Tap Trigger:")

add_h4("Timing / Frequency")
add_p("Executed once upon application process initialization.", bold_prefix="System Startup:")
add_p("Operates continuously 24/7 during robot docking station standby and between customer shopping sessions.", bold_prefix="Standby & Welcome:")

add_h4("Function Description")
add_p("Supermarket shoppers, Technical staff / Robot fleet operators.", bold_prefix="Actors / Roles:")
add_p("Android Robot Kiosk Application (React Native / Expo SDK 57), Text-to-Speech Engine (RobotVoiceService), Customer Session Manager (CustomerSessionContext), Device Telemetry Service.", bold_prefix="System:")
add_p("Deliver a high-tech welcome greeting, present real-time date/time information, and activate intelligent shopping workflows.", bold_prefix="Purpose:")
add_p("Real-time digital clock (22:24) and Vietnamese date (Thứ Năm, 17/09/2026); animated 2D robot avatar with radar pulse effect; title badge \"TRỢ LÝ SIÊU THỊ THÔNG MINH\"; prominent call-to-action button \"[ CHẠM ĐỂ BẮT ĐẦU ]\"; promotional shortcut button \"Khuyến mãi hôm nay\" with gift box icon.", bold_prefix="Interface (Based on UI Mockups):")

add_h4("Normal Flow")
add_p("1. The application starts up, displaying the robot logo with an initialization spinner and diagnostic status text: '⚡ Đang khởi động hệ thống robot...', '🔍 Đang kiểm tra Camera AI & Cảm biến...', '🎙️ Đang kích hoạt cổng nhận diện giọng nói...'.")
add_p("2. After 3 seconds of hardware self-diagnostics, the Welcome Screen fades in smoothly. Robot Assistant broadcasts a welcoming greeting via speaker: \"Chào mừng quý khách đến với Smart Market Bot! Tôi có thể giúp gì cho bạn?\" (Welcome to Smart Market Bot! How may I assist you?).")
add_p("3. The customer views the real-time clock and central touch area.")
add_p("4. The customer taps anywhere on the screen -> The system emits an audible confirmation chime \"Tuyệt vời! Chúng ta bắt đầu thôi.\" and transitions immediately to Role Selection Screen (RoleSelectionScreen).")
add_p("5. If the customer taps the \"Khuyến mãi hôm nay\" button, the system navigates directly to the Super Promotion Screen (GuestCampaignScreen).")

add_h4("Abnormal Cases")
add_p("When the device loses network or SignalR Hub connectivity, a red network warning indicator appears on the top status bar while local clock display and touch navigation buttons remain functional.", bold_prefix="Network Disconnected:")
add_p("If a customer taps the screen but performs no subsequent action within 60 seconds, the system displays a countdown dialog and returns to the Welcome Screen to protect session integrity.", bold_prefix="Idle Timeout:")

add_h4("Data Processing")
add_p("Initializes CustomerSessionId in format session-{ROBOT_ID}-{timestamp}; monitors tablet battery telemetry via BatteryService and transmits heartbeat packets every 10 seconds to REST API endpoint /api/v1/robot-operations/devices/{ROBOT_CODE}/heartbeat.", bold_prefix="Session & Telemetry Pipeline:")

add_h4("Business Rules")
add_p("Enforces useKeepAwake to prevent tablet screen sleep throughout Kiosk operation; maintains robot greeting volume at standard 70% master volume.", bold_prefix="Display & Audio Policy:")

# ==============================================================================
# 3.2.2 User Onboarding & Face ID Authentication
# ==============================================================================
add_h4("3.2.2 User Onboarding & Face ID Authentication")
add_images([img("03_Role_Selection.jpg"), img("08_Face_Scan_Success.jpg")], width_inches=2.5)
add_p("[Figure 3.2.2: Role Selection Screen and Successful Face ID Biometric Recognition]", italic=True)

add_h4("Function Trigger")
add_p("From the Welcome Screen, customer taps to transition to Role Selection Screen. The customer taps \"QUÉT FACE ID ĐĂNG NHẬP\" (Scan Face ID to Login) on the Member card to initiate facial biometric login, or taps \"BẮT ĐẦU NGAY\" (Start Now) under Guest card to proceed as a guest.", bold_prefix="Trigger Role Selection & Face Login:")

add_h4("Timing / Frequency")
add_p("Invoked at the beginning of each customer shopping session on the Robot Kiosk.", bold_prefix="Timing:")

add_h4("Function Description")
add_p("Guest shopper (Khách vãng lai), Registered Supermarket Member (Khách hàng thành viên).", bold_prefix="Actors / Roles:")
add_p("Front-facing tablet camera, FastAPI Face AI Microservice (YuNet landmark detection + SFace 128-D embedding extraction via ONNX runtime), Backend ASP.NET Core Identity & Database.", bold_prefix="System:")
add_p("Segment user shopping tiers; provide ultra-fast (< 1.5s) contactless biometric authentication to identify members, load spending budgets, personal health tags, and purchase history.", bold_prefix="Purpose:")
add_p("Role Selection screen featuring white \"Khách vãng lai\" card and primary green \"Khách hàng thành viên\" card with crown badge \"KHUYÊN DÙNG\"; Face Scan screen with circular camera viewfinder, pulsing laser border, personalized greeting notification, and primary button \"Bắt đầu mua sắm\".", bold_prefix="Interface (Based on UI Mockups):")

add_h4("Normal Flow")
add_p("1. The customer selects the \"Khách hàng thành viên\" card and clicks \"QUÉT FACE ID ĐĂNG NHẬP\".")
add_p("2. The app activates the front camera; Robot Assistant speaks: \"Xin chào! Vui lòng nhìn thẳng vào máy ảnh để nhận diện.\" (Hello! Please look directly into the camera for identification).")
add_p("3. The customer aligns their face within the circular viewfinder. The camera captures an optimized frame, Base64-encodes it, and posts it to /api/v1/auth/login-face.")
add_p("4. Backend AI Service executes SFace Cosine Similarity matching against enrolled face vectors. When similarity exceeds the threshold (> 0.72), the system issues a JWT token and loads member profile (Full Name, Member Code, Tier, Shopping Budget).")
add_p("5. The viewfinder border turns emerald green with success text: \"Chào [Tên khách hàng], rất vui gặp lại! Cửa hàng có nhiều ưu đãi mới, mời anh/chị ghé xem!\".")
add_p("6. Robot Assistant speaks a personalized greeting and automatically navigates to Member Home Screen (MemberHomeScreen).")

add_h4("Abnormal Cases")
add_p("Viewfinder border turns red with message \"Không nhận diện được khuôn mặt\" and automatically schedules a retry in 2.5 seconds. If failure exceeds 3 attempts, provides a quick button to switch to Guest mode.", bold_prefix="Face Match Failed:")
add_p("Displays a system dialog instructing staff to grant camera permission in Android settings.", bold_prefix="Camera Permission Denied:")
add_p("If the authenticated account lacks the Member role, access is denied and the screen returns to Role Selection.", bold_prefix="Unauthorized Role:")

add_h4("Data Processing")
add_p("Downscales camera frames to maximum width 640px prior to AI service transmission; 128-D vector matching completes in memory (< 25ms); stores credentials securely in RobotAuthContext.", bold_prefix="Biometric Pipeline:")

add_h4("Business Rules")
add_p("Only accounts with role 'Member' are authorized to access Member Home; active session is cleared upon 60-second idle timeout.", bold_prefix="Authentication Policy:")

# ==============================================================================
# 3.2.3 Guest Shopping Experience & Daily Promotions
# ==============================================================================
add_h4("3.2.3 Guest Shopping Experience & Daily Promotions")
add_images([img("02_Daily_Promotions.jpg"), img("04_Guest_Home_Dashboard.jpg")], width_inches=2.5)
add_p("[Figure 3.2.3: Daily Promotions Carousel and Guest Home Dashboard]", italic=True)

add_h4("Function Trigger")
add_p("Customer clicks \"BẮT ĐẦU NGAY\" on the \"Khách vãng lai\" card on Role Selection Screen, or clicks \"Khuyến mãi hôm nay\" from Welcome Screen.", bold_prefix="Trigger:")

add_h4("Timing / Frequency")
add_p("Activated whenever an unregistered customer interacts with the Robot Kiosk.", bold_prefix="Timing:")

add_h4("Function Description")
add_p("Unregistered walk-in customer.", bold_prefix="Actors / Roles:")
add_p("Android Robot Frontend, Campaign & Advertising Service (AdService), Catalog Product Service.", bold_prefix="System:")
add_p("Provide a seamless shopping experience without requiring login, showcase daily high-discount campaigns, provide QR code to download personal mobile app, and offer fast search shortcuts.", bold_prefix="Purpose:")
add_p("Welcome banner with action buttons \"Đăng ký Thành viên\" and \"Tìm hiểu thêm\"; QR code card \"Tải App nhận ưu đãi\"; autonomous simulation shortcut \"Mở màn hình robot / mô phỏng tự hành\"; 3 primary shortcut cards: \"Tìm bằng Giọng nói\", \"Sản phẩm toàn hệ thống\", \"Tìm kiếm văn bản\"; discount product grid featuring percentage badges (-12%, HOT), product photography, and green promotional prices.", bold_prefix="Interface (Based on UI Mockups):")

add_h4("Normal Flow")
add_p("1. The guest enters Guest Home and views available shopping capabilities.")
add_p("2. The guest can scan the on-screen QR code with their mobile phone camera to receive the SmartMarketBot mobile app installation link.")
add_p("3. The guest clicks \"Tìm bằng Giọng nói\" for quick voice search or \"Sản phẩm toàn hệ thống\" to browse catalog categories.")
add_p("4. The guest browses active daily discounted products (e.g., Dầu hào nấm hương, Hạt nêm nấm hương chay, Khăn ướt Bumbo, Mì Koreno Jumbo...).")
add_p("5. The guest taps any product card to view its corresponding shelf and aisle location within the supermarket.")

add_h4("Abnormal Cases")
add_p("The system displays locally cached promotional product lists to ensure uninterrupted browsing during network downtime.", bold_prefix="Network Disconnect:")

add_h4("Data Processing")
add_p("Fetches active campaigns in real-time from REST endpoint /api/v1/ads/campaigns/active and product catalog from /api/v1/products.", bold_prefix="Data Pipeline:")

add_h4("Business Rules")
add_p("Guest shoppers do not accumulate loyalty reward points; upon completing a journey, system displays a prompt inviting customer to register for member benefits.", bold_prefix="Guest Policy:")

# ==============================================================================
# 3.2.4 Member Personalization & Special Offers
# ==============================================================================
add_h4("3.2.4 Member Personalization & Special Offers")
add_images([img("09_Member_Home_Dashboard.jpg"), img("10_Member_Special_Offers.jpg")], width_inches=2.5)
add_images([img("13_Recipe_Detail_Cook_Guide.jpg")], width_inches=2.5)
add_p("[Figure 3.2.4: Member Home Dashboard, Special Offers, and Recommended Recipe Detail View]", italic=True)

add_h4("Function Trigger")
add_p("Automatically routed following successful Face ID authentication, or when customer clicks \"Dành riêng cho bạn\" or \"Ưu đãi cho bạn\" on the navigation bar.", bold_prefix="Trigger:")

add_h4("Timing / Frequency")
add_p("Active throughout the member's authenticated shopping session.", bold_prefix="Timing:")

add_h4("Function Description")
add_p("Authenticated supermarket member.", bold_prefix="Actors / Roles:")
add_p("Recommender AI Engine, Nutrition & Recipe Recommendation Service, Member Health Profile Database.", bold_prefix="System:")
add_p("Personalize shopping experience according to health profiles, dietary regimens (Eat Clean, Keto, Low-carb, Vegan), track daily grocery budget, and recommend complete meal preparation packages.", bold_prefix="Purpose:")
add_p("Header displaying customer avatar, member name (\"Nguyễn Anh Hùng\"), badge \"MEMBER\", budget progress bar (0đ / 2.500.000đ), green cart button; 4 large interactive cards: \"Tìm kiếm sản phẩm\", \"Tìm bằng Giọng nói\", \"Dành riêng cho bạn\", \"Ưu đãi cho bạn\"; Recipe detail screen showing estimated total cost (209.000đ), nutritional advice card, ingredients checklist, and primary green button \"Mua tất cả nguyên liệu\".", bold_prefix="Interface (Based on UI Mockups):")

add_h4("Normal Flow")
add_p("1. The member enters Member Home and checks their daily grocery spending budget progress.")
add_p("2. The member clicks \"Ưu đãi cho bạn\" to view exclusive discounts tailored to their membership tier.")
add_p("3. The member scrolls down to \"Món ngon gợi ý\" (Suggested Dishes) and selects \"Salad rau muống, cải bó xôi tươi giòn sốt chanh\".")
add_p("4. The recipe detail screen opens, presenting estimated total cost (209.000đ), cooking advice, and required ingredients checklist (Rau muống sạch, Rau xà lách búp, Dầu Olive nguyên chất...).")
add_p("5. The member clicks \"Mua tất cả nguyên liệu\" -> The system automatically adds all available ingredients into Cart and calculates optimal picking route.")

add_h4("Abnormal Cases")
add_p("Out-of-stock items display a red \"Hết hàng\" badge, and system suggests equivalent substitutes (e.g., organic cider vinegar instead of fresh lime).", bold_prefix="Out of Stock:")
add_p("If total cart value exceeds the 2,500,000 VND daily limit, the progress bar turns orange with an advisory notification.", bold_prefix="Budget Exceeded:")

add_h4("Data Processing")
add_p("Calculates budget expenditure percentage (currentSpend / budget) * 100; filters catalog products by active Health Tags from /api/v1/members/{id}/recommendations.", bold_prefix="Recommendation Engine:")

add_h4("Business Rules")
add_p("Budget preferences are synchronized directly from member mobile application settings; real-time recalculation triggers upon any cart item addition or removal.", bold_prefix="Budget Rules:")

# ==============================================================================
# 3.2.5 Smart Search & AI Voice Assistant
# ==============================================================================
add_h4("3.2.5 Smart Search & AI Voice Assistant")
add_images([img("05_Voice_Search_Listening.jpg"), img("06_Category_Catalog.jpg")], width_inches=2.5)
add_images([img("07_Product_Search_Result.jpg")], width_inches=2.5)
add_p("[Figure 3.2.5: Voice Assistant Listening, Category Grid Catalog, and Search Results List]", italic=True)

add_h4("Function Trigger")
add_p("Customer clicks the floating microphone button at the bottom corner or taps any text search input across the application.", bold_prefix="Trigger:")

add_h4("Timing / Frequency")
add_p("Invoked on-demand when looking up products or aisle locations.", bold_prefix="Timing:")

add_h4("Function Description")
add_p("Customer seeking product information or shelf location.", bold_prefix="Actors / Roles:")
add_p("@react-native-voice/voice recognition module, Web Speech STT API, Vietnamese Semantic Search Engine.", bold_prefix="System:")
add_p("Provide hands-free natural Vietnamese voice querying and visual browsing across 18 supermarket product categories.", bold_prefix="Purpose:")
add_p("Voice listening interface featuring dynamic sound wave visualizer, central listening robot avatar, large circular mic button, voice prompt tips; two-column pastel rounded category grid (\"Bánh quy ngọt\", \"Bát sứ\", \"Củ\", \"Đường, Tiêu & Gia vị\", \"Gạo & Ngũ cốc\", \"Hải sản tươi\"...); search results list showing product image, title, price, in-stock badge, aisle/shelf location, and \"Thêm vào giỏ\" button.", bold_prefix="Interface (Based on UI Mockups):")

add_h4("Normal Flow")
add_p("1. The customer clicks the microphone button from screen corner or home dashboard.")
add_p("2. The voice recognition interface activates with dynamic audio waves; Robot Assistant chimes ready.")
add_p("3. The customer speaks: \"Tìm cho tôi gạo hạt dài\" (Find me long-grain rice).")
add_p("4. The system converts speech to text (STT) and displays the query in the search bar.")
add_p("5. Semantic search matches database items: \"Gạo Basmati hạt dài Ấn Độ túi 1kg\", \"Bún gạo khô sạch túi 500g\".")
add_p("6. The customer inspects price and shelf location, and clicks \"Thêm vào giỏ\" or the speaker button to hear audio details.")

add_h4("Abnormal Cases")
add_p("Displays notification \"Không nghe rõ. Vui lòng nói lại\" and keeps microphone active for an additional 5 seconds.", bold_prefix="Voice Unrecognized:")
add_p("Displays similar category recommendations or popular keyword chips to assist customer query reformulation.", bold_prefix="Product Not Found:")

add_h4("Data Processing")
add_p("Normalizes Vietnamese diacritics for fuzzy string matching; queries REST API /api/v1/products/search?q={query}&voice=true.", bold_prefix="Search Pipeline:")

add_h4("Business Rules")
add_p("Automatically closes microphone and terminates listening state after 5 seconds of silence to conserve power and avoid background supermarket chatter.", bold_prefix="Voice Timeout Policy:")

# ==============================================================================
# 3.2.6 Product Detail & Instant Cart Operations
# ==============================================================================
add_h4("3.2.6 Product Detail & Instant Cart Operations")
add_images([img("11_Product_Detail_View.jpg"), img("12_Product_Detail_Added_Cart.jpg")], width_inches=2.5)
add_p("[Figure 3.2.6: Product Detail View and Added-to-Cart State with Direct Navigation Action]", italic=True)

add_h4("Function Trigger")
add_p("Customer clicks any product card in search results, category catalog, or promotional recommendation lists.", bold_prefix="Trigger:")

add_h4("Timing / Frequency")
add_p("Frequent usage when customer wants in-depth information about a specific item.", bold_prefix="Timing:")

add_h4("Function Description")
add_p("Customer evaluating product purchase.", bold_prefix="Actors / Roles:")
add_p("Product Inventory Service, Nutrition & Health Tagging Engine, Shopping Cart Service.", bold_prefix="System:")
add_p("Display comprehensive product specifications, discounted pricing, dietary health tags, and provide one-tap cart addition or immediate robot navigation dispatch.", bold_prefix="Purpose:")
add_p("High-resolution product image; bold title (\"Quả bơ sáp Đắk Lắk tươi 1kg\"); promotional price 49.000đ with -11% badge and strikethrough original price 55.000đ; \"Đặc tính sức khoẻ\" pill badges (Eat Clean, Keto, Low-carb, Vegan); product description block; quantity stepper [-] [1] [+]; circular green add button \"Thêm vào giỏ\"; large pill navigation button \"Dẫn đường\"; top toast notification on success.", bold_prefix="Interface (Based on UI Mockups):")

add_h4("Normal Flow")
add_p("1. The customer selects \"Quả bơ sáp Đắk Lắk\" from recommendation list.")
add_p("2. The application opens Product Detail view, checking dietary alignment with customer profile.")
add_p("3. The customer adjusts purchase quantity using \"+\" or \"-\" stepper buttons.")
add_p("4. The customer clicks \"Thêm vào giỏ\" -> The system sounds confirmation chime, updates cart state, and displays toast notification: \"Đã thêm 1 Quả bơ sáp Đắk Lắk tươi 1kg vào giỏ hàng\".")
add_p("5. If the customer desires immediate guidance to the shelf, they click \"Dẫn đường\" -> Robot calculates route to Shelf 3.")

add_h4("Abnormal Cases")
add_p("Disables \"+\" button and displays maximum available shelf stock warning.", bold_prefix="Out of Stock Limit:")
add_p("Highlights dietary conflict warning in prominent red/yellow badges to caution customer before purchase.", bold_prefix="Allergy Conflict:")

add_h4("Data Processing")
add_p("Updates CustomerSessionContext with selected item ID, quantity, and recalculates total cart price in real-time.", bold_prefix="Cart Synchronization:")

add_h4("Business Rules")
add_p("Customer can add items to cart without triggering robot movement; supports quantity adjustments directly on screen. Direct guidance dispatch initiates single-target navigation to the product's mapped shelf.", bold_prefix="Cart Rules:")

# ==============================================================================
# 3.2.7 Smart Cart & Autonomous Robot Navigation Guide
# ==============================================================================
add_h4("3.2.7 Smart Cart & Autonomous Robot Navigation Guide")
add_images([img("14_Member_Cart_List.jpg"), img("15_Cart_Guide_Dispatching.jpg")], width_inches=2.5)
add_images([img("16_Cart_Guide_Timeout.png")], width_inches=2.5)
add_p("[Figure 3.2.7: Smart Cart Item List, Navigation Route Dispatching, and Navigation Error Alert]", italic=True)

add_h4("Function Trigger")
add_p("On Cart Screen -> Customer clicks green button \"Robot dẫn theo giỏ hàng\" (Robot guides cart route), or on Product Detail Screen / Ad Multi-Product Selection Screen -> Customer clicks \"Bắt đầu dẫn đường\".", bold_prefix="Trigger:")

add_h4("Timing / Frequency")
add_p("Triggered when customer is ready to navigate the supermarket aisle network with the robot.", bold_prefix="Timing:")

add_h4("Function Description")
add_p("Supermarket shopper, Robot autonomous navigation controller (ROS2 Navigation Stack / ESP32).", bold_prefix="Actors / Roles:")
add_p(".NET Route Optimization Engine (Dijkstra TSP), SignalR Hub (RobotHub), ROS2 Navigation Stack, Lidar sensors, AdInterruptionService, and Mecanum wheel odometry.", bold_prefix="System:")
add_p("Optimize shelf visiting sequence to minimize travel distance; render real-time stage progress, guide item retrieval at each shelf with countdown timer, and either resume interrupted ad missions or autonomously return to docking station.", bold_prefix="Purpose:")
add_p("Cart list showing items, unit prices, quantity stepper, red delete icon, total cost (128.000đ); Cart guide interface with robot status header (\"RB001 Sẵn sàng\"), guidance state card (DISPATCHED, NAVIGATING, ARRIVED), next destination card (\"Kệ 3 - Thực Phẩm Tươi Sống\") with shelf item checklist; multi-stop route timeline (Origin Waypoint 7 -> Shelf 3 -> Return Waypoint 7); red cancel button.", bold_prefix="Interface (Based on UI Mockups):")

add_h4("Normal Flow")
add_p("1. The customer verifies items in Cart (or selects multiple promoted products from advertising interruption) and clicks \"Robot dẫn theo giỏ hàng\" / \"Bắt đầu dẫn đường\".")
add_p("2. Backend receives product list and maps each item to corresponding Shelf ID and navigation Node ID.")
add_p("3. Dijkstra Traveling Salesperson algorithm computes optimal shortest path visiting all required shelves and dispatches navigation command to robot via SignalR.")
add_p("4. Tablet switches to CartGuideMapScreen, displaying DISPATCHED state with voice prompt: \"Tôi bắt đầu dẫn đường. Xin vui lòng đi theo tôi!\" (I am starting navigation. Please follow me!).")
add_p("5. Robot travels forward (MOVING), updating progress Leg 1/N.")
add_p("6. Upon arrival at target shelf (ARRIVED), screen displays items to pick at this shelf with a 30-second countdown timer.")
add_p("7. Customer picks items into physical cart and clicks \"Đã lấy xong\" (Finished picking) or lets timer expire to proceed.")
add_p("8. Robot continues guiding to remaining shelves until all items are gathered.")
add_p("9. Upon completing the final shelf pick:")
add_p("   - If navigation originated from an interrupted ad mission (AdInterruptionService.hasInterruptedMission() == true), the robot announces: \"Đã hoàn thành dẫn đường. Robot tiếp tục lịch trình quảng cáo!\" and automatically resumes the remaining advertising patrol route.")
add_p("   - If navigation was initiated from an ordinary shopping cart, the robot autonomously returns to docking station (Waypoint 7).")

add_h4("Abnormal Cases")
add_p("Displays red error alert \"Robot không gửi trạng thái xác nhận trong 20 giây\" (Robot did not confirm status within 20s) with retry and cancel options.", bold_prefix="Navigation Timeout:")
add_p("Front Lidar senses obstacle; robot transitions to PAUSED with voice prompt: \"Phía trước có vật cản, tôi đang đợi đường thông thoáng\" (Obstacle detected, waiting for clear path).", bold_prefix="Obstacle Detected:")
add_p("Immediate motor power cutoff and brake lock in ESTOP emergency stop state.", bold_prefix="Emergency Stop:")

add_h4("Data Processing")
add_p("Multi-waypoint traveling salesperson routing; converts real-world metric coordinates to tablet screen pixels; streams real-time pose telemetry via WebSocket port 81.", bold_prefix="Navigation Pipeline:")

add_h4("Business Rules")
add_p("Always appends start waypoint (Waypoint 7) and home return waypoint (Waypoint 7) to ordinary shopping routes; prevents automatic idle session logout while robot is actively navigating; integrates bidirectionally with AdInterruptionService to support seamless mission pause and resumption.", bold_prefix="Navigation Invariant:")

# ==============================================================================
# 3.2.8 Contextual Overlays & Autonomous Advertising
# ==============================================================================
add_h4("3.2.8 Contextual Overlays & Autonomous Advertising")
add_images([img("17_Zone_Ad_Modal.png")], width_inches=2.6)
add_p("[Figure 3.2.8: High-Resolution Contextual Zone Advertising Slide-Up Modal]", italic=True)

add_h4("Function Trigger")
add_p("Autonomous Trigger: Robot navigates into geographical radius of a merchandise zone/shelf (GeofencingContext receives zoneEntered), or Store Administrator broadcasts an advertising patrol mission from Admin Web Dashboard. Interruption Trigger: Customer near robot touches the ad screen or clicks promotional action buttons.", bold_prefix="Trigger:")

add_h4("Timing / Frequency")
add_p("Operates 24/7 during autonomous store patrol hours and when traversing sponsored aisles.", bold_prefix="Timing:")

add_h4("Function Description")
add_p("Walk-in shoppers (Khách vãng lai), Registered members (Khách hàng thành viên), Brand Sponsors, Supermarket Administrators.", bold_prefix="Actors / Roles:")
add_p("Ad Package & Priority Engine, Geofencing Engine, Dynamic Ad Playlist Service, AdInterruptionService, SignalR RobotHub, RobotVoiceService.", bold_prefix="System:")
add_p("Deliver location-based contextual promotions directly at points of interest, monetize promotional broadcasts via Brand Ad Packages (AdScore priority scoring), enable interactive ad interruption with differentiated Guest vs Member guidance, and seamlessly preserve and resume advertising patrol upon guidance completion.", bold_prefix="Purpose:")
add_p("Slide-up modal sheet with 16:9 hero banner, campaign headline, product photography, discounted price, countdown progress bar (12 seconds to 0), call-to-action buttons (\"Thêm vào giỏ ngay\", \"Dẫn tôi mua món này\", \"Xem tất cả sản phẩm khuyến mãi\"), and close button [X]; Multi-product selection modal (AdMultiProductSelectScreen) featuring prioritized product cards, package score badges, \"Chọn tất cả\" / \"Bỏ chọn\" toggles, and primary navigation button \"Bắt đầu dẫn đường\".", bold_prefix="Interface (Based on UI Mockups):")

add_h4("Detailed Business Architecture & Workflow")

add_p("Brands subscribe to commercial advertising packages (AdPackage: VIP, Pro, Standard, Basic). Each package defines flexible billing unit prices (ZoneUnitPrice, ShelfUnitPrice, RouteUnitPrice, ClickFee) and a system priority weight (AdScore, e.g. VIP = 100 points, Pro = 50 points, Standard = 20 points).", bold_prefix="1. Brand Ad Package & Priority Scoring:")
add_p("Store administrators broadcast ad missions from the Admin Web Dashboard targeted by Shelf (shelfId), by Zone (zoneId), or by Patrol Route (fixed waypoints / free-roam). Upon receiving an ad mission, the robot aggregates active campaigns, deduplicates items, and sorts the broadcast sequence strictly by package points: OrderByDescending(AdScore).ThenByDescending(Priority). System brand campaigns receive an additional priority bonus (+50 points). High-scoring packages receive prime screen exposure and voice broadcast frequency.", bold_prefix="2. Admin Dispatch & Priority Sorting:")
add_p("While the robot is actively patrolling or stopped at an aisle broadcasting advertisements, shoppers can interrupt the broadcast to request product guidance:", bold_prefix="3. Ad Interruption & Shopper Tier Segmentation:")
add_p("   • Guest Shoppers (Khách vãng lai): Do not need to log in. The guest is permitted to select exactly ONE (01) product currently displayed on the ad modal. Tapping \"Dẫn tôi mua món này\" immediately dispatches single-shelf navigation to escort the guest directly to that item.")
add_p("   • Registered Members (Khách hàng thành viên): The customer authenticates quickly via Face ID (< 1.5s) or login credentials, unlocking the Multi-Product Selection Screen (AdMultiProductSelectScreen). The member can browse all active sponsored products sorted by AdScore, tick multiple items simultaneously (with \"Chọn tất cả\" / \"Bỏ chọn\"), and tap \"Bắt đầu dẫn đường\". Robot computes an optimal multi-stop traveling salesperson route (Dijkstra TSP) to visit all chosen shelves.")
add_p("The instant a customer triggers product guidance, AdInterruptionService freezes and stores the interrupted ad mission context (remainingNodeIds, remainingShelfIds, isFreeRoam, campaignId, savedTimestamp). The active ad mission on the backend is safely paused/cancelled to avoid motor contention.", bold_prefix="4. State Preservation (AdInterruptionService):")
add_p("The robot guides the customer to each shelf with a 30-second picking timer. Upon the customer clicking \"Đã lấy xong\" at the final destination, the system inspects AdInterruptionService.hasInterruptedMission(). Detecting the paused ad mission, the robot announces: \"Đã hoàn thành dẫn đường. Robot tiếp tục lịch trình quảng cáo!\" and automatically resumes the remaining advertising patrol route without needing to return to the docking station.", bold_prefix="5. Automatic Ad Patrol Resumption:")

add_h4("Normal Flow")
add_p("1. Admin dispatches an advertising patrol mission to robot RB001 covering Zone 1 and Shelf 3.")
add_p("2. Robot loads active campaigns, sorts products by AdScore (VIP packages first), and begins autonomous patrol.")
add_p("3. Upon entering Shelf 3 radius, ZoneAdOverlay slides up; speaker broadcasts product audio endorsement: \"[Tên sản phẩm] - Giá ưu đãi chỉ [Giá] đồng. Quý khách có thể chạm vào màn hình để thêm vào giỏ hoặc yêu cầu tôi dẫn đường nhé!\".")
add_p("4. Countdown progress bar animates from 12 seconds to 0.")
add_p("5. A walk-in customer taps \"Dẫn tôi mua món này\" -> Robot immediately pauses ad patrol, caches remaining waypoints into AdInterruptionService, and guides customer to Shelf 3.")
add_p("6. Alternatively, a registered member taps \"Xem tất cả sản phẩm khuyến mãi\", logs in via Face ID, selects 3 promoted items on AdMultiProductSelectScreen, and taps \"Bắt đầu dẫn đường\" -> Robot computes Dijkstra TSP route covering all 3 shelves.")
add_p("7. At each shelf, the customer retrieves items and taps \"Đã lấy xong\".")
add_p("8. Once all items are collected, Robot detects the paused ad mission, speaks confirmation: \"Đã hoàn thành dẫn đường. Robot tiếp tục lịch trình quảng cáo!\", and smoothly resumes patrolling the remaining ad waypoints.")

add_h4("Abnormal Cases")
add_p("Automatically mutes ad audio and collapses overlay when urgent navigation guidance or safety warnings take precedence.", bold_prefix="High-priority Guidance:")
add_p("Serves pre-cached local banners and slogans stored on tablet storage during network disconnection.", bold_prefix="Offline Fallback:")
add_p("Enforces 1.5-second debounce on ad click interactions and flags repetitive clicks in session to prevent fraudulent sponsor billing.", bold_prefix="Click Fraud Protection:")

add_h4("Data Processing")
add_p("Logs ad impressions and click events to /api/v1/ads/impressions and /api/v1/ads/clicks with sessionId, robotCode, and campaignId for sponsor budget reconciliation; calculates Cosine Similarity in memory for biometric role elevation.", bold_prefix="Ad Analytics Pipeline:")

add_h4("Business Rules")
add_p("Advertising policy yields screen priority unconditionally to emergency stop alerts or critical path corrections; guest shoppers are restricted to single-product guidance, whereas authenticated members enjoy multi-product selection and TSP routing; ad mission automatically resumes post-guidance without manual operator intervention.", bold_prefix="Advertising & Guidance Policy:")

# Save master document
print(f"\n[Step 4] Saving updated Master Document to: {MASTER_DOCX}")
doc.save(MASTER_DOCX)
print("Master Document saved successfully!")

# Mirror copy to screenShot dir
print(f"\n[Step 5] Mirroring updated document to: {MIRROR_DOCX}")
shutil.copyfile(MASTER_DOCX, MIRROR_DOCX)
print("Document mirrored successfully!")
print("=" * 70)
