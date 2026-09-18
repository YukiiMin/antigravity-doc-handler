# -*- coding: utf-8 -*-
import openpyxl

out_file = r"d:\Minh\For_myself\ZSCORT_GSU26_SAP05\tool\AI-Docx-Testing-Product\dataset1\output\Report5_Unit Test_ChoHung.xlsx"
wb = openpyxl.load_workbook(out_file, data_only=False)

print("=== VERIFICATION SUMMARY ===")
print(f"Total sheets: {len(wb.sheetnames)}")
print(f"Sheet names: {wb.sheetnames}")

# 1. Guideline
print("\n[1] Guideline:")
ws_g = wb["Guideline"]
print("  A2:", ws_g["A2"].value)

# 2. Cover
print("\n[2] Cover:")
ws_cov = wb["Cover"]
print("  Images count:", len(ws_cov._images))
print("  B4 (Project Name):", ws_cov["B4"].value)
print("  B5 (Project Code):", ws_cov["B5"].value)
print("  B6 (Doc Code):", ws_cov["B6"].value)
print("  F4 (Creator):", ws_cov["F4"].value)
print("  F5 (Date):", ws_cov["F5"].value)
print("  F6 (Version):", ws_cov["F6"].value)

# 3. Functions
print("\n[3] Functions:")
ws_fn = wb["Functions"]
print("  E4:", ws_fn["E4"].value)
print("  E5:", ws_fn["E5"].value)
print(f"  Row 11: No={ws_fn.cell(11,1).value}, Sheet={ws_fn.cell(11,6).value}, FuncCode={ws_fn.cell(11,5).value}")
print(f"  Row 30: No={ws_fn.cell(30,1).value}, Sheet={ws_fn.cell(30,6).value}, FuncCode={ws_fn.cell(30,5).value}")

# 4. Statistics
print("\n[4] Statistics:")
ws_st = wb["Statistics"]
print(f"  Row 12 (NAV-SVC): Code={ws_st.cell(12,2).value}, Pass={ws_st.cell(12,3).value}, Total={ws_st.cell(12,9).value}")
print(f"  Row 31 (FACE-AI): Code={ws_st.cell(31,2).value}, Pass={ws_st.cell(31,3).value}, Total={ws_st.cell(31,9).value}")
print(f"  Row 32 (Sub total): B32={ws_st.cell(32,2).value}, C32={ws_st.cell(32,3).value}, I32={ws_st.cell(32,9).value}")
print(f"  Row 34 (Coverage): {ws_st.cell(34,4).value}")
print("  Charts count:", len(ws_st._charts))
for idx, ch in enumerate(ws_st._charts):
    print(f"    Chart {idx}: title={ch.title}, anchor={ch.anchor}")
    for ser in ch.series:
        print(f"      series data={ser.val.numRef.f}")

# 5. Function Sheets
for sname in ["NAV-SVC", "IMP-PROD", "FACE-AI"]:
    ws = wb[sname]
    print(f"\n[5] Function Sheet [{sname}]:")
    print("  C2 (Code):", ws["C2"].value, ", L2 (Name):", ws["L2"].value)
    print("  A7 (Pass formula):", ws["A7"].value)
    print("  C7 (Fail formula):", ws["C7"].value)
    print("  O7 (Total formula):", ws["O7"].value)
    print(f"  F9 (TC1 Header): {ws['F9'].value}, font={ws['F9'].font.name}, fill={ws['F9'].fill.start_color.rgb}")
    for r in range(10, 25):
        if ws.cell(r, 6).value == "O":
            c = ws.cell(r, 6)
            print(f"  First O mark at ({r},6): val={c.value}, font={c.font.name} {c.font.size}pt bold={c.font.bold}, border={c.border.left.style}")
            break
