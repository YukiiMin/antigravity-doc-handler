# -*- coding: utf-8 -*-
"""inspect_example_tokens.py — Extract format tokens from Example sheet."""
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

TMPL = r"d:\Minh\For_myself\ZSCORT_GSU26_SAP05\tool\AI-Docx-Testing-Product\dataset1\Template\Report5_Unit Test-template.xlsx"

wb = openpyxl.load_workbook(TMPL, data_only=False)
ws_ex = wb["Example"]

print("=== EXAMPLE SHEET: Format Token Extraction ===\n")

# 1. First 'O' mark
for r in range(10, 43):
    found = False
    for c in range(6, 25):
        cell = ws_ex.cell(r, c)
        if cell.value == "O":
            f = cell.font
            a = cell.alignment
            b = cell.border
            print("[ MARK 'O' ] ref cell:", cell.coordinate)
            print("  font.name  =", f.name)
            print("  font.size  =", f.size)
            print("  font.bold  =", f.bold)
            try: print("  font.color =", f.color.rgb)
            except: print("  font.color = (default)")
            print("  align.horizontal =", a.horizontal)
            print("  align.vertical   =", a.vertical)
            print("  border.left.style =", b.left.style if b.left else None)
            found = True
            break
    if found:
        break

print()

# 2. Row 9 header (first TC col)
print("[ ROW 9 HEADER ] first TC column:")
for c in range(6, 10):
    cell = ws_ex.cell(9, c)
    f = cell.font
    a = cell.alignment
    fi = cell.fill
    print(f"  {cell.coordinate}: val={repr(cell.value)}")
    print(f"    fill.fgColor = {fi.start_color.rgb}")
    print(f"    font = {f.name} {f.size}pt bold={f.bold}")
    try: print(f"    font.color  = {f.color.rgb}")
    except: print(f"    font.color  = (default)")
    print(f"    align.textRotation = {a.textRotation}")
    print(f"    align.horizontal   = {a.horizontal}")
    print(f"    align.vertical     = {a.vertical}")

print()

# 3. Column A section cells
print("[ COL A SECTION ] Condition/Confirm/Result:")
for r in [10, 31, 39]:
    cell = ws_ex.cell(r, 1)
    f = cell.font
    a = cell.alignment
    fi = cell.fill
    print(f"  A{r}: val={repr(cell.value)}")
    print(f"    fill.fgColor = {fi.start_color.rgb}")
    print(f"    font = {f.name} {f.size}pt bold={f.bold}")
    try: print(f"    font.color  = {f.color.rgb}")
    except: print(f"    font.color  = (default)")
    print(f"    align.textRotation = {a.textRotation}")

print()

# 4. Result rows 39/40/41
print("[ RESULT ROWS 39/40/41 ] type/status/date:")
for r in [39, 40, 41]:
    for c in range(6, 10):
        cell = ws_ex.cell(r, c)
        if cell.value is not None:
            f = cell.font
            fi = cell.fill
            print(f"  {cell.coordinate}: val={repr(cell.value)}")
            print(f"    fill.fgColor = {fi.start_color.rgb if fi.fill_type else 'none'}")
            print(f"    font = {f.name} {f.size}pt bold={f.bold}")
            try: print(f"    font.color  = {f.color.rgb}")
            except: print(f"    font.color  = (default)")
