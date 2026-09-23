import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

DOCX_PATH = r"E:\Do_an_SU26\SuperMarketBot-BE\docs\Report4_Software Design Document.docx"

TABLES_DATA = [
    ("01", "ACCOUNT", "Stores user account credentials, password hashes, email, role, OTP codes, and authentication tokens. - Primary keys: AccountID - Foreign keys: None"),
    ("02", "MEMBER", "Stores supermarket member profile information, biometric face vectors, spending limits, and accumulated loyalty points. - Primary keys: MemberID - Foreign keys: AccountID"),
    ("03", "MEMBERSHIP", "Defines loyalty membership tiers and status for registered customer profiles. - Primary keys: MembershipID - Foreign keys: MemberID"),
    ("04", "HEALTH_TAG", "Catalog of health, allergen, and dietary restriction tags (e.g., Gluten-Free, Vegan, Nut-Allergy). - Primary keys: HealthTagID - Foreign keys: None"),
    ("05", "MEMBERHEALTH_PREFERENCE", "Associates members with specific dietary, health, or allergen tags to enable personalized product warnings. - Primary keys: MemberID, HealthTagID - Foreign keys: MemberID, HealthTagID"),
    ("06", "HEALTH_TAG_CONFLICT", "Defines incompatibility rules and warnings between conflicting dietary and health tags. - Primary keys: TagId1, TagId2 - Foreign keys: TagId1, TagId2"),
    ("07", "CATEGORY", "Master high-level product classifications in the supermarket catalog. - Primary keys: CategoryID - Foreign keys: None"),
    ("08", "SUBCATEGORY", "Mid-level product categorization belonging to a specific primary Category. - Primary keys: SubcategoryID - Foreign keys: CategoryID"),
    ("09", "PRODUCT_TYPE", "Granular product classification linked to specific functional product groups. - Primary keys: ProductTypeID - Foreign keys: SubcategoryID"),
    ("10", "PRODUCT", "Stores supermarket item master data including barcodes, unit prices, stock, expiry dates, and specs. - Primary keys: ProductID - Foreign keys: ProductTypeID"),
    ("11", "PRODUCT_HEALTHTAG", "Maps products to health tags to indicate dietary properties and potential allergens. - Primary keys: ProductID, HealthTagID - Foreign keys: ProductID, HealthTagID"),
    ("12", "MEAL_SUGGESTION", "Stores meal recipe suggestions and dietary combination ideas generated dynamically by Gemini AI. - Primary keys: MealSuggestionID - Foreign keys: None"),
    ("13", "CART", "Represents a member's active shopping cart session created for automated pathfinding and order planning. - Primary keys: CartID - Foreign keys: MemberID"),
    ("14", "CART_ITEM", "Stores individual products and requested quantities placed within an active shopping cart. - Primary keys: CartItemID - Foreign keys: CartID, ProductID"),
    ("15", "INVOICE_HISTORY", "Stores completed checkout transactions, totals, and historical invoice header records. - Primary keys: InvoiceHistoryID - Foreign keys: MemberID"),
    ("16", "INVOICE_HISTORY_ITEM", "Stores detailed product breakdown, unit prices, and quantities for each historical invoice transaction. - Primary keys: InvoiceHistoryItemID - Foreign keys: InvoiceHistoryID, ProductID"),
    ("17", "FLOOR", "Represents supermarket building floors or physical level layouts. - Primary keys: FloorID - Foreign keys: None"),
    ("18", "ZONE", "Represents specific sections/zones designated on a supermarket floor. - Primary keys: ZoneID - Foreign keys: FloorID"),
    ("19", "AISLE", "Represents individual shopping aisles located within a specific Zone. - Primary keys: AisleID - Foreign keys: ZoneID"),
    ("20", "SHELF", "Represents physical shelf units installed along aisles, mapped directly to navigation waypoints. - Primary keys: ShelfID - Foreign keys: AisleID, NodeID"),
    ("21", "SLOT", "Represents granular product placement positions/slots across shelf tiers. - Primary keys: SlotID - Foreign keys: ShelfID"),
    ("22", "PRODUCT_SLOT", "Maps specific Product inventory to physical Slot locations on store shelves. - Primary keys: ProductsSlotID - Foreign keys: SlotID, ProductID"),
    ("23", "MAP", "Stores spatial grid maps, SLAM navigation layouts, and floorplans for autonomous mobile robots. - Primary keys: MapID - Foreign keys: FloorID"),
    ("24", "NAVIGATION_NODE", "Represents discrete navigation coordinate points (waypoints) on the floor map. - Primary keys: NodeID - Foreign keys: MapID"),
    ("25", "NAVIGATION_EDGE", "Defines navigable pathways, distances, and bidirectional connections between navigation nodes. - Primary keys: EdgeID - Foreign keys: FromNodeID, ToNodeID"),
    ("26", "ROBOT", "Stores hardware identifiers, battery levels, operational modes, and IP addresses of AMR units. - Primary keys: RobotID - Foreign keys: None"),
    ("27", "ROBOT_LOG", "Logs operational telemetry, battery levels, and real-time coordinates of robots. - Primary keys: LogID - Foreign keys: RobotID"),
    ("28", "SHELF_SCAN", "Stores camera/AI vision scanning logs, stock density, and restock alerts from robot shelf patrols. - Primary keys: ScanID - Foreign keys: ShelfID, RobotID, NavigationNodeID"),
    ("29", "ROBOT_ROUTE", "Stores planned patrol, navigation, and advertisement routes defined across map waypoints. - Primary keys: RobotRouteID - Foreign keys: MapID, RobotID, ZoneID"),
    ("30", "ROUTE_ASSIGNMENT", "Assigns generated RobotRoutes to specific Robot units for mission execution. - Primary keys: RouteAssignmentID - Foreign keys: RobotID, RobotRouteID"),
    ("31", "ROUTE_NODE_MAPPING", "Maps sequential NavigationNodes and dwell times that form a complete RobotRoute. - Primary keys: RouteNodeMappingID - Foreign keys: RobotRouteID, NodeID"),
    ("32", "BRAND", "Stores master records of product manufacturers and advertising brand partners. - Primary keys: BrandID - Foreign keys: None"),
    ("33", "AD_PACKAGE", "Defines advertising campaign packages, pricing tiers, and duration limits. - Primary keys: PackageID - Foreign keys: None"),
    ("34", "AD_CAMPAIGN", "Stores promotional campaigns created by brand partners and managed by Admins. - Primary keys: AdCampaignID - Foreign keys: PackageID, BrandID"),
    ("35", "AD_CAMPAIGN_LOG", "Logs real-time ad display events, viewer interactions, and impression metrics. - Primary keys: LogID - Foreign keys: AdCampaignID, ProductID, RobotID, ShelfID, ZoneID, SlotID, MemberID"),
    ("36", "AD_RESOURCE", "Stores multimedia assets (banners, videos, voice audio) used in advertising campaigns. - Primary keys: ResourceID - Foreign keys: AdCampaignID"),
    ("37", "SPONSORED_PRODUCT", "Links specific products promoted within an active advertising campaign. - Primary keys: SponsoredID - Foreign keys: AdCampaignID, ProductID"),
    ("38", "AD_CAMPAIGN_ZONE", "Maps targeted advertising campaigns to specific store Zones. - Primary keys: AdCampaignID, ZoneID - Foreign keys: AdCampaignID, ZoneID"),
    ("39", "AD_CAMPAIGN_SHELF", "Maps targeted advertising campaigns to specific physical Shelves. - Primary keys: AdCampaignID, ShelfID - Foreign keys: AdCampaignID, ShelfID"),
    ("40", "AD_CAMPAIGN_ROUTE", "Maps targeted advertising campaigns to Robot patrol routes. - Primary keys: AdCampaignID, RobotRouteID - Foreign keys: AdCampaignID, RobotRouteID"),
]

def format_cell(cell, text, font_name="Tahoma", font_size_pt=9.0, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold

    # OpenXML cell vertical alignment: center
    tcPr = cell._element.get_or_add_tcPr()
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), 'center')
    tcPr.append(vAlign)

def patch_table_descriptions():
    print(f"Loading DOCX: {DOCX_PATH}")
    doc = docx.Document(DOCX_PATH)
    t = doc.tables[3]

    header_cells = [c.text.strip() for c in t.rows[0].cells]
    print(f"Table 3 Header: {header_cells}")

    # Remove existing data rows (keep header row 0)
    for _ in range(len(t.rows) - 1):
        tr = t.rows[1]._element
        t._element.remove(tr)

    # Ensure header row has tblHeader and cantSplit
    trPr0 = t.rows[0]._element.get_or_add_trPr()
    if not trPr0.xpath('w:tblHeader'):
        trPr0.append(OxmlElement('w:tblHeader'))
    if not trPr0.xpath('w:cantSplit'):
        trPr0.append(OxmlElement('w:cantSplit'))

    # Format header row cells
    format_cell(t.rows[0].cells[0], "No", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell(t.rows[0].cells[1], "Table", bold=True)
    format_cell(t.rows[0].cells[2], "Description", bold=True)

    # Add 40 new data rows
    for no, tbl_name, desc in TABLES_DATA:
        row = t.add_row()
        trPr = row._element.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))

        format_cell(row.cells[0], no, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(row.cells[1], tbl_name, bold=True)
        format_cell(row.cells[2], desc)

    print(f"Updated Table 3 to {len(t.rows)} rows (1 header + {len(TABLES_DATA)} data rows).")
    doc.save(DOCX_PATH)
    print(f"[OK] Document successfully saved to: {DOCX_PATH}")

if __name__ == "__main__":
    patch_table_descriptions()
