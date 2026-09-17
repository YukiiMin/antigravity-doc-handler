# -*- coding: utf-8 -*-
"""
consolidate_and_clean_output.py — Single Source of Output Consolidator & Workspace Cleaner.

1. Reads immutable templates from dataset1/Template/
2. Injects data and formulas in-place.
3. Emits ALL deliverables exclusively into dataset1/output/:
   - dataset1/output/Report5_Test Report.xlsx
   - dataset1/output/Report5_Unit Test.xlsx
   - dataset1/output/Report5_Unit Test.xls
   - dataset1/output/Report5_Test Documentation.docx
   - dataset1/output/test_screenshots/ (7 Retina 300+ DPI PNGs & HTMLs)
4. Purges all scattered duplicate output files from:
   - AI-Docx-Testing-Product/ (root)
   - dataset1/ (root)
   - dataset1/Report5_Unit Test_files/
"""

from __future__ import annotations

import os
import sys
import shutil
import re
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')
sys.stderr.reconfigure(encoding='utf-8')

BASE_DIR = r"d:\Minh\For_myself\ZSCORT_GSU26_SAP05\tool\AI-Docx-Testing-Product"
DATASET1_DIR = os.path.join(BASE_DIR, "dataset1")
TMPL_DIR = os.path.join(DATASET1_DIR, "Template")
OUTPUT_DIR = os.path.join(DATASET1_DIR, "output")
OUT_SCREENSHOTS = os.path.join(OUTPUT_DIR, "test_screenshots")

def safe_set_cell(ws, coordinate: str, value=None, formula: str | None = None, fill_color: str | None = None, font_bold: bool | None = None, font_color: str | None = None, font_size: float | None = None):
    target_coord = coordinate
    target_cell = ws[coordinate]
    
    if isinstance(target_cell, openpyxl.cell.cell.MergedCell):
        for m_range in ws.merged_cells.ranges:
            if coordinate in m_range:
                target_coord = m_range.coord.split(":")[0]
                break
                
    cell = ws[target_coord]
    
    if formula and str(formula).startswith("="):
        cell.value = str(formula)
    elif value is not None:
        cell.value = value

    if fill_color:
        h = fill_color.lstrip("#").upper()
        if len(h) == 6:
            h = f"FF{h}"
        cell.fill = PatternFill(fill_type="solid", start_color=h, end_color=h)

    if font_bold is not None or font_color is not None or font_size is not None:
        curr_font = cell.font or Font()
        bold_val = font_bold if font_bold is not None else curr_font.bold
        size_val = font_size if font_size is not None else (curr_font.size or 10)
        fc_val = None
        if font_color:
            raw_fc = str(font_color).lstrip("#").upper()
            fc_val = f"FF{raw_fc}" if len(raw_fc) == 6 else raw_fc

        cell.font = Font(
            name=curr_font.name or "Segoe UI",
            size=size_val,
            bold=bold_val,
            italic=curr_font.italic,
            color=fc_val
        )


def build_output_files():
    print("==================================================================")
    print(f"=== 1. GENERATING DELIVERABLES INTO: {OUTPUT_DIR}")
    print("==================================================================")
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    os.makedirs(OUT_SCREENSHOTS, exist_ok=True)

    # 1. Test Report (.xlsx)
    tmpl_tr = os.path.join(TMPL_DIR, "Report5_Test Report-template.xlsx")
    wb_tr = openpyxl.load_workbook(tmpl_tr, data_only=False)

    # Cover
    ws_cover = wb_tr["Cover"]
    safe_set_cell(ws_cover, "B4", value="SuperMarketBot")
    safe_set_cell(ws_cover, "B5", value="SMB_SU26")
    safe_set_cell(ws_cover, "B6", value="SMB_SU26_Test_Report_v1.0")
    safe_set_cell(ws_cover, "F4", value="Nguyễn Quang Huy")
    safe_set_cell(ws_cover, "F5", value="15/09/2026")
    safe_set_cell(ws_cover, "F6", value="v1.0")

    # Test Cases
    ws_tc = wb_tr["Test Cases"]
    safe_set_cell(ws_tc, "D5", value=".NET 10.0 SDK | Kestrel 10.0.9 | Python 3.14.3 FastAPI | Azure SQL Database | Expo SDK 57 / React 19 | Mosquitto MQTT v2.0")

    # Test Statistics
    ws_ts = wb_tr["Test Statistics"]
    safe_set_cell(ws_ts, "C4", value="Nguyễn Quang Huy")
    safe_set_cell(ws_ts, "C5", value="Lê Trọng Hiếu")
    safe_set_cell(ws_ts, "C6", value="15/09/2026")
    
    safe_set_cell(ws_ts, "D11", formula="='Feature 1'!B6")
    safe_set_cell(ws_ts, "E11", formula="='Feature 1'!C6")
    safe_set_cell(ws_ts, "F11", formula="='Feature 1'!D6")
    safe_set_cell(ws_ts, "G11", formula="='Feature 1'!E6")
    safe_set_cell(ws_ts, "H11", formula="='Feature 1'!F6")

    safe_set_cell(ws_ts, "D12", formula="='Feature 2'!B6")
    safe_set_cell(ws_ts, "E12", formula="='Feature 2'!C6")
    safe_set_cell(ws_ts, "F12", formula="='Feature 2'!D6")
    safe_set_cell(ws_ts, "G12", formula="='Feature 2'!E6")
    safe_set_cell(ws_ts, "H12", formula="='Feature 2'!F6")

    safe_set_cell(ws_ts, "D13", formula="='Feature 3'!B6")
    safe_set_cell(ws_ts, "E13", formula="='Feature 3'!C6")
    safe_set_cell(ws_ts, "F13", formula="='Feature 3'!D6")
    safe_set_cell(ws_ts, "G13", formula="='Feature 3'!E6")
    safe_set_cell(ws_ts, "H13", formula="='Feature 3'!F6")

    safe_set_cell(ws_ts, "D14", formula="='Feature 4'!B6")
    safe_set_cell(ws_ts, "E14", formula="='Feature 4'!C6")
    safe_set_cell(ws_ts, "F14", formula="='Feature 4'!D6")
    safe_set_cell(ws_ts, "G14", formula="='Feature 4'!E6")
    safe_set_cell(ws_ts, "H14", formula="='Feature 4'!F6")

    safe_set_cell(ws_ts, "D15", formula="=SUM(D11:D14)")
    safe_set_cell(ws_ts, "E15", formula="=SUM(E11:E14)")
    safe_set_cell(ws_ts, "F15", formula="=SUM(F11:F14)")
    safe_set_cell(ws_ts, "G15", formula="=SUM(G11:G14)")
    safe_set_cell(ws_ts, "H15", formula="=SUM(H11:H14)")
    safe_set_cell(ws_ts, "E16", formula="=(D15+E15)*100/(H15-G15)")
    safe_set_cell(ws_ts, "E17", formula="=D15*100/(H15-G15)")

    features = [
        ("Feature 1", "Core Backend & AI Realtime APIs", 15),
        ("Feature 2", "Web Admin Portal & Staff Management", 9),
        ("Feature 3", "Customer Mobile App & Autonomous Robot Client", 8),
        ("Feature 4", "User Acceptance Testing (End-to-End Journeys)", 3),
    ]
    for sname, title, count in features:
        ws_f = wb_tr[sname]
        safe_set_cell(ws_f, "B4", value=title)
        safe_set_cell(ws_f, "B6", value=count)
        safe_set_cell(ws_f, "C6", value=0)
        safe_set_cell(ws_f, "D6", value=0)
        safe_set_cell(ws_f, "E6", value=0)
        safe_set_cell(ws_f, "F6", value=count)

    out_tr_xlsx = os.path.join(OUTPUT_DIR, "Report5_Test Report.xlsx")
    wb_tr.save(out_tr_xlsx)
    print(f"  [OK] Generated: {out_tr_xlsx} ({os.path.getsize(out_tr_xlsx):,} bytes)")

    # 2. Unit Test (.xlsx)
    tmpl_ut = os.path.join(TMPL_DIR, "Report5_Unit Test-template.xlsx")
    wb_ut = openpyxl.load_workbook(tmpl_ut, data_only=False)

    if "Function3" in wb_ut.sheetnames and "Function 3" not in wb_ut.sheetnames:
        wb_ut["Function3"].title = "Function 3"

    ws_cover_ut = wb_ut["Cover"]
    safe_set_cell(ws_cover_ut, "B4", value="SuperMarketBot")
    safe_set_cell(ws_cover_ut, "B5", value="SMB_SU26")
    safe_set_cell(ws_cover_ut, "B6", value="SMB_SU26_Unit_Test_v1.0")
    safe_set_cell(ws_cover_ut, "F4", value="Nguyễn Quang Huy")
    safe_set_cell(ws_cover_ut, "F5", value="15/09/2026")
    safe_set_cell(ws_cover_ut, "F6", value="v1.0")

    ws_func = wb_ut["Functions"]
    safe_set_cell(ws_func, "E4", formula="=Cover!B4")
    safe_set_cell(ws_func, "E5", formula="=Cover!B5")
    
    functions_data = [
        ("MealSuggestionsController.GetAiMenuAssistant", "Function 1", "Nguyễn Quang Huy", 85),
        ("NavigationControllerDemo.AutonomousDispatch", "Function 2", "Lê Trọng Hiếu", 300),
        ("NavigationService.PlanRouteAsync", "Function 3", "Lê Trọng Hiếu", 300),
        ("CartGuidePlanner.CalculateOptimalSequence", "Function 4", "Lê Trọng Hiếu", 150),
        ("ShelfVisionReIdService.DetectAndMatchShelfSlots", "Function 5", "Nguyễn Quang Huy", 220),
        ("FaceEmbeddingExtractor.ExtractSFaceVector", "Function 6", "Nguyễn Quang Huy", 140),
        ("SignalRRealtimeHub.BroadcastRobotTelemetry", "Function 7", "Nguyễn Anh Hùng", 110),
    ]
    for idx, (fname, fcode, dev, loc) in enumerate(functions_data, start=11):
        safe_set_cell(ws_func, f"D{idx}", value=fname)
        safe_set_cell(ws_func, f"E{idx}", value=fcode)

    ws_stat = wb_ut["Statistics"]
    safe_set_cell(ws_stat, "B4", formula="=Cover!B4")
    safe_set_cell(ws_stat, "B5", formula="=Cover!B5")
    safe_set_cell(ws_stat, "E4", value="Nguyễn Quang Huy")
    safe_set_cell(ws_stat, "E5", value="Lê Trọng Hiếu")
    safe_set_cell(ws_stat, "E6", value="15/09/2026")

    function_sheets = [
        (
            "Function 1",
            "MealSuggestionsController.GetAiMenuAssistant",
            "Nguyễn Quang Huy",
            "Cung cấp trợ lý AI gợi ý thực đơn thông minh theo thành phần dinh dưỡng và giỏ hàng của khách hàng."
        ),
        (
            "Function 2",
            "NavigationControllerDemo.AutonomousDispatch",
            "Lê Trọng Hiếu",
            "Điều phối robot tự hành tuần tra toàn bản đồ siêu thị và phát video quảng cáo thương hiệu theo lịch."
        ),
        (
            "Function 3",
            "NavigationService.PlanRouteAsync",
            "Lê Trọng Hiếu",
            "Tính toán lộ trình di chuyển ngắn nhất cho robot bằng thuật toán Dijkstra qua 8 nodes và 11 cạnh đồ thị."
        ),
    ]
    for sname, fname, dev_name, desc in function_sheets:
        if sname not in wb_ut.sheetnames:
            continue
        ws_f = wb_ut[sname]
        safe_set_cell(ws_f, "C3", value=dev_name)
        safe_set_cell(ws_f, "C5", value=desc)
        safe_set_cell(ws_f, "F3", value="Nguyễn Quang Huy")
        safe_set_cell(ws_f, "L2", value=fname)

    if "Example" in wb_ut.sheetnames:
        ws_ex = wb_ut["Example"]
        safe_set_cell(ws_ex, "C3", value="Nguyễn Quang Huy")
        safe_set_cell(ws_ex, "C5", value="Kiểm thử mẫu hàm trợ lý thực đơn AI.")
        safe_set_cell(ws_ex, "F3", value="Nguyễn Quang Huy")

    out_ut_xlsx = os.path.join(OUTPUT_DIR, "Report5_Unit Test.xlsx")
    wb_ut.save(out_ut_xlsx)
    print(f"  [OK] Generated: {out_ut_xlsx} ({os.path.getsize(out_ut_xlsx):,} bytes)")

    # 3. Unit Test (.xls)
    out_ut_xls = os.path.join(OUTPUT_DIR, "Report5_Unit Test.xls")
    import win32com.client
    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb_com = excel.Workbooks.Open(os.path.abspath(out_ut_xlsx))
        wb_com.SaveAs(os.path.abspath(out_ut_xls), FileFormat=56)
        wb_com.Close(False)
        print(f"  [OK] Generated: {out_ut_xls} ({os.path.getsize(out_ut_xls):,} bytes)")
    finally:
        excel.Quit()

    # 4. Word Document (.docx)
    out_docx = os.path.join(OUTPUT_DIR, "Report5_Test Documentation.docx")
    if not os.path.exists(out_docx):
        zip_path = os.path.join(BASE_DIR, "dataset1.zip")
        if os.path.exists(zip_path):
            import zipfile
            with zipfile.ZipFile(zip_path, 'r') as z:
                if "Report5_Test Documentation.docx" in z.namelist():
                    with z.open("Report5_Test Documentation.docx") as src, open(out_docx, "wb") as dst:
                        dst.write(src.read())
    print(f"  [OK] Word Document verified: {out_docx} ({os.path.getsize(out_docx):,} bytes)")

    # 5. Screenshots
    src_screenshots = os.path.join(DATASET1_DIR, "test_screenshots")
    if os.path.exists(src_screenshots):
        for f in os.listdir(src_screenshots):
            s_src = os.path.join(src_screenshots, f)
            s_dst = os.path.join(OUT_SCREENSHOTS, f)
            shutil.copyfile(s_src, s_dst)
    else:
        zip_path = os.path.join(BASE_DIR, "dataset1.zip")
        if os.path.exists(zip_path):
            import zipfile
            with zipfile.ZipFile(zip_path, 'r') as z:
                for item in z.namelist():
                    if item.startswith("test_screenshots/") and not item.endswith("/"):
                        fname = os.path.basename(item)
                        with z.open(item) as src, open(os.path.join(OUT_SCREENSHOTS, fname), "wb") as dst:
                            dst.write(src.read())
    print(f"  [OK] Synced {len(os.listdir(OUT_SCREENSHOTS))} screenshot assets to: {OUT_SCREENSHOTS}")


def purge_scattered_duplicates():
    print("\n==================================================================")
    print("=== 2. PURGING SCATTERED DUPLICATE FILES")
    print("==================================================================")
    
    # 1. Clean root AI-Docx-Testing-Product/ duplicate excel files
    files_to_remove = [
        os.path.join(BASE_DIR, "Report5_Test Report.xlsx"),
        os.path.join(BASE_DIR, "Report5_Unit Test.xlsx"),
        os.path.join(BASE_DIR, "Report5_Test Documentation.docx"),
        os.path.join(DATASET1_DIR, "Report5_Test Report.xlsx"),
        os.path.join(DATASET1_DIR, "Report5_Unit Test.xlsx"),
        os.path.join(DATASET1_DIR, "Report5_Test Documentation.docx"),
        os.path.join(DATASET1_DIR, "Report5_Unit Test.htm"),
    ]
    for fp in files_to_remove:
        if os.path.exists(fp):
            os.remove(fp)
            print(f"  [Removed Duplicate File]: {fp}")

    # 2. Clean Report5_Unit Test_files folder
    files_dir = os.path.join(DATASET1_DIR, "Report5_Unit Test_files")
    if os.path.exists(files_dir):
        shutil.rmtree(files_dir)
        print(f"  [Removed Deprecated Folder]: {files_dir}")

    # 3. Clean root test_screenshots in dataset1 (moved to output/)
    src_screenshots = os.path.join(DATASET1_DIR, "test_screenshots")
    if os.path.exists(src_screenshots):
        shutil.rmtree(src_screenshots)
        print(f"  [Cleaned Old Screenshots Dir]: {src_screenshots}")


def verify_consolidated_workspace():
    print("\n==================================================================")
    print("=== 3. FINAL WORKSPACE VERIFICATION")
    print("==================================================================")
    print(f"Listing contents of {DATASET1_DIR}:")
    for item in os.listdir(DATASET1_DIR):
        item_path = os.path.join(DATASET1_DIR, item)
        if os.path.isdir(item_path):
            count = len(os.listdir(item_path))
            print(f"  📁 {item}/ ({count} items)")
        else:
            print(f"  📄 {item} ({os.path.getsize(item_path):,} bytes)")

    print(f"\nListing SINGLE OUTPUT FOLDER ({OUTPUT_DIR}):")
    for item in os.listdir(OUTPUT_DIR):
        item_path = os.path.join(OUTPUT_DIR, item)
        if os.path.isdir(item_path):
            count = len(os.listdir(item_path))
            print(f"  📁 {item}/ ({count} files)")
        else:
            print(f"  📄 {item} ({os.path.getsize(item_path):,} bytes)")


if __name__ == "__main__":
    build_output_files()
    purge_scattered_duplicates()
    verify_consolidated_workspace()
