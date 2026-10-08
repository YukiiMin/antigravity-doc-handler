"""
doctools.operations.xlsx.write_catalog — Extended Write Catalog operations (FR-35, D-14, TC-55).
Provides typed, validated Pydantic write operations:
- set_format: atomic formatting updates (number format, font, fill, border, alignment)
- set_validation: Data Validation injection
- set_conditional_format: Conditional Formatting rule application
- copy_sheet: high-fidelity sheet cloning via SheetCloner parity engine
"""

from __future__ import annotations
import builtins
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, ConfigDict, Field

from doctools.contract.envelope import Diagnostics, FileRef, ResultEnvelope, Stats
from doctools.contract.issues import Engine, Issue, Severity
from doctools.core.xlsx.mutate.sheet_cloner import clone_sheet_with_parity
from doctools.infra.file_store import FileStore
from doctools.registry import ToolRegistry

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class SetFormatInput(BaseModel):
    model_config = ConfigDict(extra="ignore")
    file_path: str
    sheet: str
    range: str = Field(..., alias="cell_range")
    number_format: Optional[str] = None
    bold: Optional[bool] = None
    font_size: Optional[int] = None
    font_color: Optional[str] = None
    fill_color: Optional[str] = None
    border_style: Optional[str] = None
    border_color: Optional[str] = "000000"
    alignment: Optional[str] = None
    wrap_text: Optional[bool] = None


class SetValidationInput(BaseModel):
    model_config = ConfigDict(extra="ignore")
    file_path: str
    sheet: str
    range: str = Field(..., alias="cell_range")
    type: str = Field(default="list", description="Validation type: list, whole, decimal, date")
    formula1: str = Field(..., description="Validation formula or list (e.g. '\"O,X,N/A\"')")
    formula2: Optional[str] = None
    allow_blank: bool = True


class SetConditionalFormatInput(BaseModel):
    model_config = ConfigDict(extra="ignore")
    file_path: str
    sheet: str
    range: str = Field(..., alias="cell_range")
    operator: str = "equal"
    formula: List[str] = Field(default_factory=lambda: ['"FAIL"'])
    fill_color: Optional[str] = "FFEE1111"


class CopySheetInput(BaseModel):
    model_config = ConfigDict(extra="ignore")
    file_path: str
    source_sheet: str
    target_title: str
    rewire_self_refs: bool = True


def xlsx_set_format(
    file_path: Union[str, Path],
    sheet: str,
    range: Optional[str] = None,
    cell_range: Optional[str] = None,
    number_format: Optional[str] = None,
    bold: Optional[bool] = None,
    font_size: Optional[int] = None,
    font_color: Optional[str] = None,
    fill_color: Optional[str] = None,
    border_style: Optional[str] = None,
    border_color: Optional[str] = "000000",
    alignment: Optional[str] = None,
    wrap_text: Optional[bool] = None,
    file_store: Optional[FileStore] = None,
) -> ResultEnvelope:
    """Applies atomic formatting tokens to a range of cells."""
    diag = Diagnostics(engine=Engine.XLSX)
    target_rng = cell_range or range or "A1"
    p = Path(file_path).resolve()
    if not p.is_file():
        diag.add_issue(Issue(code="E-XLSX-NOT-FOUND", severity=Severity.ERROR, message=f"File not found: {p}"))
        return ResultEnvelope(success=False, diagnostics=diag)

    wb = openpyxl.load_workbook(p, data_only=False)
    if sheet not in wb.sheetnames:
        diag.add_issue(Issue(code="E-XLSX-SHEET-NOT-FOUND", severity=Severity.ERROR, message=f"Sheet not found: {sheet}"))
        return ResultEnvelope(success=False, diagnostics=diag)

    ws = wb[sheet]
    min_c, min_r, max_c, max_r = range_boundaries(target_rng)

    border_obj = None
    if border_style:
        side = Side(style=border_style, color=border_color or "000000")
        border_obj = Border(top=side, bottom=side, left=side, right=side)

    fill_obj = None
    if fill_color:
        fill_obj = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

    for r in builtins.range(min_r, max_r + 1):
        for c in builtins.range(min_c, max_c + 1):
            cell = ws.cell(row=r, column=c)
            if number_format is not None:
                cell.number_format = number_format
            if bold is not None or font_size is not None or font_color is not None:
                cur_font = cell.font or Font()
                cell.font = Font(
                    name=cur_font.name,
                    size=font_size or cur_font.size,
                    bold=bold if bold is not None else cur_font.bold,
                    color=font_color or cur_font.color,
                )
            if fill_obj:
                cell.fill = fill_obj
            if border_obj:
                cell.border = border_obj
            if alignment is not None or wrap_text is not None:
                cell.alignment = Alignment(horizontal=alignment, wrap_text=wrap_text)

    wb.save(p)
    out_ref = file_store.store_file(p, mime=_XLSX_MIME, engine="xlsx") if file_store else None
    return ResultEnvelope(
        success=True,
        diagnostics=diag,
        file_ref=out_ref,
        stats=Stats(extra={"cells_updated": (max_r - min_r + 1) * (max_c - min_c + 1)}),
    )


def xlsx_set_validation(
    file_path: Union[str, Path],
    sheet: str,
    range: Optional[str] = None,
    cell_range: Optional[str] = None,
    type: str = "list",
    formula1: str = '"O,X,N/A"',
    formula2: Optional[str] = None,
    allow_blank: bool = True,
    file_store: Optional[FileStore] = None,
) -> ResultEnvelope:
    """Configures Data Validation on specified cell range."""
    diag = Diagnostics(engine=Engine.XLSX)
    target_rng = cell_range or range or "A1"
    p = Path(file_path).resolve()
    if not p.is_file():
        diag.add_issue(Issue(code="E-XLSX-NOT-FOUND", severity=Severity.ERROR, message=f"File not found: {p}"))
        return ResultEnvelope(success=False, diagnostics=diag)

    wb = openpyxl.load_workbook(p, data_only=False)
    if sheet not in wb.sheetnames:
        diag.add_issue(Issue(code="E-XLSX-SHEET-NOT-FOUND", severity=Severity.ERROR, message=f"Sheet not found: {sheet}"))
        return ResultEnvelope(success=False, diagnostics=diag)

    ws = wb[sheet]
    dv = DataValidation(type=type, formula1=formula1, formula2=formula2, allow_blank=allow_blank)
    ws.add_data_validation(dv)
    dv.add(target_rng)

    wb.save(p)
    out_ref = file_store.store_file(p, mime=_XLSX_MIME, engine="xlsx") if file_store else None
    return ResultEnvelope(
        success=True,
        diagnostics=diag,
        file_ref=out_ref,
        stats=Stats(extra={"validation_added": target_rng, "type": type}),
    )


def xlsx_set_conditional_format(
    file_path: Union[str, Path],
    sheet: str,
    range: Optional[str] = None,
    cell_range: Optional[str] = None,
    operator: str = "equal",
    formula: Optional[List[str]] = None,
    fill_color: str = "FFEE1111",
    file_store: Optional[FileStore] = None,
) -> ResultEnvelope:
    """Adds a Conditional Formatting rule to a cell range."""
    diag = Diagnostics(engine=Engine.XLSX)
    target_rng = cell_range or range or "A1"
    p = Path(file_path).resolve()
    if not p.is_file():
        diag.add_issue(Issue(code="E-XLSX-NOT-FOUND", severity=Severity.ERROR, message=f"File not found: {p}"))
        return ResultEnvelope(success=False, diagnostics=diag)

    wb = openpyxl.load_workbook(p, data_only=False)
    if sheet not in wb.sheetnames:
        diag.add_issue(Issue(code="E-XLSX-SHEET-NOT-FOUND", severity=Severity.ERROR, message=f"Sheet not found: {sheet}"))
        return ResultEnvelope(success=False, diagnostics=diag)

    ws = wb[sheet]
    fill_obj = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    rule = CellIsRule(operator=operator, formula=formula or ['"FAIL"'], fill=fill_obj)
    ws.conditional_formatting.add(target_rng, rule)

    wb.save(p)
    out_ref = file_store.store_file(p, mime=_XLSX_MIME, engine="xlsx") if file_store else None
    return ResultEnvelope(
        success=True,
        diagnostics=diag,
        file_ref=out_ref,
        stats=Stats(extra={"cf_added": target_rng, "operator": operator}),
    )


def xlsx_copy_sheet(
    file_path: Union[str, Path],
    source_sheet: str,
    target_title: str,
    rewire_self_refs: bool = True,
    file_store: Optional[FileStore] = None,
) -> ResultEnvelope:
    """Clones a worksheet with 100% object parity via SheetCloner."""
    diag = Diagnostics(engine=Engine.XLSX)
    p = Path(file_path).resolve()
    if not p.is_file():
        diag.add_issue(Issue(code="E-XLSX-NOT-FOUND", severity=Severity.ERROR, message=f"File not found: {p}"))
        return ResultEnvelope(success=False, diagnostics=diag)

    wb = openpyxl.load_workbook(p, data_only=False)
    if source_sheet not in wb.sheetnames:
        diag.add_issue(Issue(code="E-XLSX-SHEET-NOT-FOUND", severity=Severity.ERROR, message=f"Sheet not found: {source_sheet}"))
        return ResultEnvelope(success=False, diagnostics=diag)

    clone_sheet_with_parity(wb, source_sheet, target_title, rewire_self_refs=rewire_self_refs)
    wb.save(p)

    out_ref = file_store.store_file(p, mime=_XLSX_MIME, engine="xlsx") if file_store else None
    return ResultEnvelope(
        success=True,
        diagnostics=diag,
        file_ref=out_ref,
        stats=Stats(extra={"source_sheet": source_sheet, "cloned_sheet": target_title}),
    )


def register_xlsx_write_catalog_tools(registry: ToolRegistry) -> None:
    """Registers extended write catalog tools into ToolRegistry."""
    @registry.register(name="xlsx.set_format", description="Áp dụng định dạng ô (số, font, viền, fill).", input_model=SetFormatInput)
    def handle_format(file_path: str, sheet: str, range: str, **kwargs: Any) -> ResultEnvelope:
        return xlsx_set_format(file_path, sheet, range=range, file_store=registry.file_store, **kwargs)

    @registry.register(name="xlsx.set_validation", description="Gán Data Validation cho dải ô.", input_model=SetValidationInput)
    def handle_val(file_path: str, sheet: str, range: str, **kwargs: Any) -> ResultEnvelope:
        return xlsx_set_validation(file_path, sheet, range=range, file_store=registry.file_store, **kwargs)

    @registry.register(name="xlsx.set_conditional_format", description="Gán quy tắc Conditional Formatting.", input_model=SetConditionalFormatInput)
    def handle_cf(file_path: str, sheet: str, range: str, **kwargs: Any) -> ResultEnvelope:
        return xlsx_set_conditional_format(file_path, sheet, range=range, file_store=registry.file_store, **kwargs)

    @registry.register(name="xlsx.copy_sheet", description="Nhân bản sheet bảo toàn 100% đối tượng và DV.", input_model=CopySheetInput)
    def handle_copy(file_path: str, source_sheet: str, target_title: str, rewire_self_refs: bool = True) -> ResultEnvelope:
        return xlsx_copy_sheet(file_path, source_sheet, target_title, rewire_self_refs, file_store=registry.file_store)
