"""
doctools.operations.xlsx.inspect_ops — Thao tác MCP kiểm tra chi tiết cấu trúc bảng tính (xlsx.inspect),
báo cáo ma trận bao phủ (xlsx.coverage_report), phân tích công thức (xlsx.analyze_formulas),
và mô tả định dạng (xlsx.describe_formats). Tuân thủ FR-30, FR-31, FR-33, FR-36, D-26, D-28.
"""

from __future__ import annotations
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import zipfile
import openpyxl
from pydantic import BaseModel, ConfigDict, Field

from doctools.contract.envelope import Diagnostics, FileRef, ResultEnvelope, Stats
from doctools.contract.issues import Engine, Issue, Severity
from doctools.core.xlsx.inspect.coverage_analyzer import CoverageAnalyzer
from doctools.core.xlsx.inspect.format_profiler import FormatProfiler
from doctools.core.xlsx.inspect.formula_profiler import FormulaProfiler
from doctools.core.xlsx.inspect.sheet_inspector import TieredSheetInspector
from doctools.infra.file_store import FileStore
from doctools.registry import ToolRegistry


class InspectInput(BaseModel):
    model_config = ConfigDict(extra="ignore")
    file_ref: Union[str, Dict[str, Any]]
    level: str = Field(default="summary", description="summary | structure | objects | cells")
    sheet: Optional[str] = Field(default=None, description="Tên worksheet mục tiêu")
    anchor: Optional[str] = Field(default=None, description="Anchor hoặc dải ô")
    page: int = Field(default=1, ge=1, description="Trang kết quả phân trang")
    max_items: int = Field(default=100, ge=1, le=500, description="Số lượng mục tối đa mỗi trang")


class CoverageReportInput(BaseModel):
    model_config = ConfigDict(extra="ignore")
    file_ref: Optional[Union[str, Dict[str, Any]]] = Field(default=None)


class AnalyzeFormulasInput(BaseModel):
    model_config = ConfigDict(extra="ignore")
    file_ref: Union[str, Dict[str, Any]]
    sheet: Optional[str] = Field(default=None)


class DescribeFormatsInput(BaseModel):
    model_config = ConfigDict(extra="ignore")
    file_ref: Union[str, Dict[str, Any]]
    sheet: Optional[str] = Field(default=None)


def _resolve_file(input_val: Any, file_store: Optional[FileStore] = None) -> Path:
    if isinstance(input_val, dict):
        input_val = FileRef(**input_val)
    if isinstance(input_val, FileRef):
        return file_store.resolve(input_val) if file_store else Path(input_val.uri.replace("file://", "")).resolve()
    s = str(input_val)
    if file_store and s.startswith("resource://"):
        return file_store.resolve(s)
    return Path(s[7:] if s.startswith("file://") else s).resolve()


def _get_package_parts(file_path: Path) -> List[str]:
    try:
        with zipfile.ZipFile(file_path, "r") as zf:
            return zf.namelist()
    except Exception:
        return []


def xlsx_inspect(
    file_ref_or_path: Union[str, FileRef, Dict[str, Any], Path],
    level: str = "summary",
    sheet: Optional[str] = None,
    anchor: Optional[str] = None,
    page: int = 1,
    max_items: int = 100,
    file_store: Optional[FileStore] = None,
) -> ResultEnvelope:
    diag = Diagnostics(engine=Engine.XLSX)
    try:
        real_path = _resolve_file(file_ref_or_path, file_store)
        wb = openpyxl.load_workbook(real_path, data_only=False)
        package_parts = _get_package_parts(real_path)
    except Exception as exc:
        diag.add_issue(Issue(code="E-XLSX-INSPECT-LOAD-FAILED", severity=Severity.ERROR, engine=Engine.XLSX, message=str(exc)))
        return ResultEnvelope(success=False, diagnostics=diag)

    inspector = TieredSheetInspector(wb, package_parts)
    norm = level.lower().strip()
    if norm == "structure":
        data = inspector.inspect_structure()
    elif norm == "objects":
        data = inspector.inspect_objects(page=page, max_items=max_items)
    elif norm == "cells":
        data = inspector.inspect_cells(sheet_name=sheet, anchor=anchor, page=page, max_items=max_items)
    else:
        norm = "summary"
        data = inspector.inspect_summary()

    extra = dict(data)
    if norm == "summary":
        extra["drawingml_count"] = sum(1 for p in package_parts if p.startswith("xl/drawings/"))
        extra["images_count"] = sum(1 for p in package_parts if p.startswith("xl/media/"))

    return ResultEnvelope(success=True, diagnostics=diag, guarantees_applied=["STRUCTURE_INSPECTED"], stats=Stats(elements_processed=len(wb.worksheets), extra=extra))


def xlsx_coverage_report(
    file_ref_or_path: Optional[Union[str, FileRef, Dict[str, Any], Path]] = None,
    file_store: Optional[FileStore] = None,
) -> ResultEnvelope:
    diag = Diagnostics(engine=Engine.XLSX)
    baseline = CoverageAnalyzer.get_baseline_matrix()
    report_extra: Dict[str, Any] = {"baseline_matrix": baseline, "total_feature_families": len(baseline)}

    if file_ref_or_path:
        try:
            real_path = _resolve_file(file_ref_or_path, file_store)
            wb = openpyxl.load_workbook(real_path, data_only=False)
            report_extra["workbook_audit"] = CoverageAnalyzer.analyze_workbook(wb, _get_package_parts(real_path))
        except Exception as exc:
            diag.add_issue(Issue(code="W-XLSX-COVERAGE-AUDIT-FAILED", severity=Severity.WARNING, engine=Engine.XLSX, message=str(exc)))

    return ResultEnvelope(success=True, diagnostics=diag, guarantees_applied=["COVERAGE_MATRIX_REPORTED"], stats=Stats(elements_processed=len(baseline), extra=report_extra))


def xlsx_analyze_formulas(
    file_ref_or_path: Union[str, FileRef, Dict[str, Any], Path],
    sheet: Optional[str] = None,
    file_store: Optional[FileStore] = None,
) -> ResultEnvelope:
    diag = Diagnostics(engine=Engine.XLSX)
    try:
        real_path = _resolve_file(file_ref_or_path, file_store)
        wb = openpyxl.load_workbook(real_path, data_only=False)
    except Exception as exc:
        diag.add_issue(Issue(code="E-XLSX-FORMULA-LOAD-FAILED", severity=Severity.ERROR, engine=Engine.XLSX, message=str(exc)))
        return ResultEnvelope(success=False, diagnostics=diag)

    report = FormulaProfiler(wb).profile(target_sheet=sheet)
    for items, fn in [
        (report.get("pattern_breaks", []), lambda it: it["message"]),
        (report.get("empty_cell_references", []), lambda it: it["message"]),
        (report.get("circular_references", []), lambda it: f"Phát hiện chu trình: {it['cycle']}"),
    ]:
        for it in items:
            diag.add_issue(Issue(code=it["code"], severity=Severity.WARNING, engine=Engine.XLSX, message=fn(it)))

    return ResultEnvelope(success=True, diagnostics=diag, guarantees_applied=["FORMULAS_ANALYZED"], stats=Stats(elements_processed=report.get("total_formulas", 0), extra=report))


def xlsx_describe_formats(
    file_ref_or_path: Union[str, FileRef, Dict[str, Any], Path],
    sheet: Optional[str] = None,
    file_store: Optional[FileStore] = None,
) -> ResultEnvelope:
    diag = Diagnostics(engine=Engine.XLSX)
    try:
        real_path = _resolve_file(file_ref_or_path, file_store)
        wb = openpyxl.load_workbook(real_path, data_only=False)
    except Exception as exc:
        diag.add_issue(Issue(code="E-XLSX-FORMAT-LOAD-FAILED", severity=Severity.ERROR, engine=Engine.XLSX, message=str(exc)))
        return ResultEnvelope(success=False, diagnostics=diag)

    report = FormatProfiler(wb).profile(target_sheet=sheet)
    for bg in report.get("border_gaps_detected", []):
        diag.add_issue(Issue(code=bg["code"], severity=Severity.WARNING, engine=Engine.XLSX, message=bg["message"]))

    return ResultEnvelope(success=True, diagnostics=diag, guarantees_applied=["FORMATS_DESCRIBED"], stats=Stats(elements_processed=report.get("unique_fonts_count", 0), extra=report))


def register_xlsx_inspect_tools(registry: ToolRegistry) -> None:
    @registry.register(name="xlsx.inspect", description="Kiểm tra chi tiết workbook 4 cấp độ.", input_model=InspectInput)
    def handle_inspect(file_ref: Any, level: str = "summary", sheet: Optional[str] = None, anchor: Optional[str] = None, page: int = 1, max_items: int = 100) -> ResultEnvelope:
        return xlsx_inspect(file_ref, level, sheet, anchor, page, max_items, registry.file_store)

    @registry.register(name="xlsx.coverage_report", description="Báo cáo ma trận bao phủ năng lực.", input_model=CoverageReportInput)
    def handle_cov(file_ref: Optional[Any] = None) -> ResultEnvelope:
        return xlsx_coverage_report(file_ref, registry.file_store)

    @registry.register(name="xlsx.analyze_formulas", description="Phân tích công thức R1C1 và rủi ro.", input_model=AnalyzeFormulasInput)
    def handle_formulas(file_ref: Any, sheet: Optional[str] = None) -> ResultEnvelope:
        return xlsx_analyze_formulas(file_ref, sheet, registry.file_store)

    @registry.register(name="xlsx.describe_formats", description="Mô tả 12 lớp định dạng số, Font, viền.", input_model=DescribeFormatsInput)
    def handle_formats(file_ref: Any, sheet: Optional[str] = None) -> ResultEnvelope:
        return xlsx_describe_formats(file_ref, sheet, registry.file_store)
