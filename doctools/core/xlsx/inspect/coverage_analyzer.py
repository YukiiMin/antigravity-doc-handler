"""
doctools.core.xlsx.inspect.coverage_analyzer
Coverage Matrix Analyzer following Annex I of Foundation Spec v1.2 (FR-36, D-28).
Defines the 5 standard coverage states and audits feature coverage.
"""

from __future__ import annotations
from enum import Enum
from typing import Any, Dict, List, Optional
import openpyxl


class CoverageState(str, Enum):
    """The 5 standard capability coverage states defined in Annex I."""
    READ_WRITE = "READ_WRITE"
    READ_ONLY = "READ_ONLY"
    PRESERVE_ONLY = "PRESERVE_ONLY"
    DETECT_ONLY = "DETECT_ONLY"
    UNSUPPORTED = "UNSUPPORTED"


class FeatureCategory:
    """Specification of an Excel feature and its expected coverage."""

    def __init__(
        self,
        key: str,
        name: str,
        target_state: CoverageState,
        current_state: CoverageState,
        notes: str,
    ) -> None:
        self.key = key
        self.name = name
        self.target_state = target_state
        self.current_state = current_state
        self.notes = notes

    def to_dict(self) -> Dict[str, Any]:
        return {
            "key": self.key,
            "name": self.name,
            "target_state": self.target_state.value,
            "current_state": self.current_state.value,
            "notes": self.notes,
        }


# Complete baseline definition matching Annex I (24 feature families)
ANNEX_I_BASELINE: List[FeatureCategory] = [
    FeatureCategory("values_datatypes", "Values & Primitive Datatypes", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "FR-24"),
    FeatureCategory("formulas_basic", "Formulas (standard, shared, array)", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "Shift Manager AST"),
    FeatureCategory("formulas_dynamic_spill", "Dynamic Arrays & Spill Formulas", CoverageState.READ_ONLY, CoverageState.READ_ONLY, "Read-only inspection"),
    FeatureCategory("functions_catalog", "Formula Functions Catalog", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "TC-50 versioned registry"),
    FeatureCategory("cell_formats_styles", "Cell Styles, Formats & Themes", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "FR-33, FR-35"),
    FeatureCategory("merged_cells", "Merged Cells", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "Top-left write + border sync"),
    FeatureCategory("conditional_formatting", "Conditional Formatting", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "Standard CF supported; x14 preserve"),
    FeatureCategory("data_validation", "Data Validation", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "Dropdowns, numbers, formulas"),
    FeatureCategory("tables_list_objects", "ListObject Tables & AutoFilters", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "Full XML table shifting"),
    FeatureCategory("defined_names", "Defined Names & Print Areas", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "Global and local sheet scope"),
    FeatureCategory("hyperlinks_comments", "Hyperlinks & Legacy Comments", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "Standard cell annotations"),
    FeatureCategory("threaded_comments", "Threaded Comments", CoverageState.DETECT_ONLY, CoverageState.DETECT_ONLY, "Modern Office 365 threads"),
    FeatureCategory("images_logos", "DrawingML Images & Logos", CoverageState.PRESERVE_ONLY, CoverageState.PRESERVE_ONLY, "Preserve via openpyxl/Pillow"),
    FeatureCategory("charts_native", "DrawingML Native Charts", CoverageState.READ_WRITE, CoverageState.PRESERVE_ONLY, "P1 Chart generation"),
    FeatureCategory("shapes_form_controls", "DrawingML Shapes & Form Controls", CoverageState.DETECT_ONLY, CoverageState.DETECT_ONLY, "TC-11 preserved in zip"),
    FeatureCategory("sparklines_slicers", "Sparklines, Slicers & Timelines", CoverageState.DETECT_ONLY, CoverageState.DETECT_ONLY, "Risk of losing layout"),
    FeatureCategory("pivot_tables", "Pivot Tables & Pivot Caches", CoverageState.PRESERVE_ONLY, CoverageState.PRESERVE_ONLY, "Requires client refresh on open"),
    FeatureCategory("rich_text", "Rich Text within Cells", CoverageState.PRESERVE_ONLY, CoverageState.PRESERVE_ONLY, "Complex inline runs"),
    FeatureCategory("sheet_protection", "Sheet & Workbook Protection", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "Protected ranges & passwords"),
    FeatureCategory("page_setup", "Page Setup, Headers & Footers", CoverageState.READ_WRITE, CoverageState.READ_WRITE, "Print margins & orientations"),
    FeatureCategory("hidden_sheets", "Hidden & VeryHidden Sheets", CoverageState.READ_ONLY, CoverageState.READ_ONLY, "Inspect visibility states"),
    FeatureCategory("external_links", "External Workbook Links", CoverageState.UNSUPPORTED, CoverageState.UNSUPPORTED, "SEC-03 blocked links"),
    FeatureCategory("vba_macros", "VBA Macros (.xlsm)", CoverageState.UNSUPPORTED, CoverageState.UNSUPPORTED, "SEC-02 rejected executables"),
    FeatureCategory("legacy_xls", "Legacy BIFF (.xls)", CoverageState.READ_ONLY, CoverageState.READ_ONLY, "FR-27 LibreOffice conversion"),
]


class CoverageAnalyzer:
    """Analyzes and reports feature coverage matrices according to Annex I."""

    @staticmethod
    def get_baseline_matrix() -> List[Dict[str, Any]]:
        """Returns the complete Annex I baseline coverage matrix."""
        return [cat.to_dict() for cat in ANNEX_I_BASELINE]

    @staticmethod
    def analyze_workbook(
        wb: openpyxl.Workbook,
        package_parts: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """
        Scans a live workbook and correlates detected features
        with their active Coverage Matrix status.
        """
        detected: Dict[str, Dict[str, Any]] = {}
        parts = package_parts or []

        # 1. Values & Types
        detected["values_datatypes"] = {
            "status": CoverageState.READ_WRITE.value,
            "present": True,
            "count": sum(ws.max_row * ws.max_column for ws in wb.worksheets),
        }

        # 2. Formulas
        f_count = sum(
            1
            for ws in wb.worksheets
            for row in ws.iter_rows(values_only=True)
            for val in row
            if str(val).startswith("=")
        )
        detected["formulas_basic"] = {
            "status": CoverageState.READ_WRITE.value,
            "present": f_count > 0,
            "count": f_count,
        }

        # 3. Merged Cells
        m_count = sum(len(ws.merged_cells.ranges) for ws in wb.worksheets)
        detected["merged_cells"] = {
            "status": CoverageState.READ_WRITE.value,
            "present": m_count > 0,
            "count": m_count,
        }

        # 4. Tables
        tbl_count = sum(len(ws.tables) for ws in wb.worksheets)
        detected["tables_list_objects"] = {
            "status": CoverageState.READ_WRITE.value,
            "present": tbl_count > 0,
            "count": tbl_count,
        }

        # 5. Data Validation
        dv_count = sum(
            len(getattr(ws.data_validations, "dataValidation", []))
            if hasattr(ws, "data_validations")
            else 0
            for ws in wb.worksheets
        )
        detected["data_validation"] = {
            "status": CoverageState.READ_WRITE.value,
            "present": dv_count > 0,
            "count": dv_count,
        }

        # 6. Conditional Formatting
        cf_count = sum(len(getattr(ws, "conditional_formatting", [])) for ws in wb.worksheets)
        detected["conditional_formatting"] = {
            "status": CoverageState.READ_WRITE.value,
            "present": cf_count > 0,
            "count": cf_count,
        }

        # 7. Images & Drawings
        img_count = sum(len(getattr(ws, "_images", [])) for ws in wb.worksheets)
        detected["images_logos"] = {
            "status": CoverageState.PRESERVE_ONLY.value,
            "present": img_count > 0 or any("xl/media/" in p for p in parts),
            "count": img_count,
        }

        # 8. Charts
        has_charts = any("xl/charts/" in p for p in parts)
        detected["charts_native"] = {
            "status": CoverageState.PRESERVE_ONLY.value,
            "present": has_charts,
            "count": 1 if has_charts else 0,
        }

        # 9. Slicers / Pivots
        has_pivots = any("xl/pivotTables/" in p for p in parts)
        detected["pivot_tables"] = {
            "status": CoverageState.PRESERVE_ONLY.value,
            "present": has_pivots,
            "count": 1 if has_pivots else 0,
        }

        has_slicers = any("xl/slicers/" in p for p in parts)
        detected["sparklines_slicers"] = {
            "status": CoverageState.DETECT_ONLY.value,
            "present": has_slicers,
            "count": 1 if has_slicers else 0,
        }

        # 10. VBA Macros
        has_vba = any("xl/vbaProject.bin" in p for p in parts)
        detected["vba_macros"] = {
            "status": CoverageState.UNSUPPORTED.value,
            "present": has_vba,
            "count": 1 if has_vba else 0,
        }

        # Status distribution
        summary_by_state: Dict[str, int] = {st.value: 0 for st in CoverageState}
        for info in detected.values():
            if info["present"]:
                summary_by_state[info["status"]] += 1

        return {
            "detected_features": detected,
            "active_states_distribution": summary_by_state,
            "has_unsupported_features": summary_by_state[CoverageState.UNSUPPORTED.value] > 0,
            "has_detect_only_features": summary_by_state[CoverageState.DETECT_ONLY.value] > 0,
        }
