"""
doctools.core.xlsx.preflight.scanner
Preflight Scanner for Excel workbooks (FR-01, FR-02).
Establishes OpenXML package inventory and selects Fidelity Tier (T1..T4).
"""

from __future__ import annotations
import io
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union
import zipfile

import openpyxl

from doctools.contract import Engine, FixableBy, Issue, Location, Severity
from doctools.contract.xlsx.preflight import (
    FidelityTier,
    PackageInventory,
    SheetInventory,
)


class PreflightScanResult:
    """Holds scanned package inventory, detected issues, and applied guarantees."""

    def __init__(
        self,
        inventory: PackageInventory,
        issues: List[Issue],
        guarantees_applied: List[str],
    ) -> None:
        self.inventory = inventory
        self.issues = issues
        self.guarantees_applied = guarantees_applied


class PreflightScanner:
    """
    Scans Excel OpenXML package parts, detects fidelity risks,
    and classifies the workbook into Fidelity Tier T1..T4.
    """

    def scan(self, workbook_input: Union[Path, bytes, str]) -> PreflightScanResult:
        """Executes full preflight audit on the workbook."""
        issues: List[Issue] = []
        guarantees: List[str] = ["preflight_inventory_established"]

        # 1. Resolve raw bytes
        raw_bytes: bytes
        filename: str = "workbook.xlsx"
        if isinstance(workbook_input, bytes):
            raw_bytes = workbook_input
        else:
            p = Path(workbook_input).resolve()
            filename = p.name
            if not p.is_file():
                issues.append(
                    Issue(
                        code="E-PKG-NOT-FOUND",
                        severity=Severity.ERROR,
                        engine=Engine.XLSX,
                        message=f"Workbook file does not exist: {p}",
                        location=Location(part="package"),
                    )
                )
                empty_inv = PackageInventory(file_name=filename, fidelity_tier="T1")
                return PreflightScanResult(empty_inv, issues, guarantees)
            raw_bytes = p.read_bytes()

        # 2. Inspect ZIP parts
        has_vba = False
        has_ext_links = False
        has_drawings = False
        has_images = False
        has_charts = False
        has_pivots = False
        has_slicers = False

        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes), "r") as zf:
                names = zf.namelist()

                if "xl/workbook.xml" not in names:
                    issues.append(
                        Issue(
                            code="E-PKG-CORRUPT",
                            severity=Severity.ERROR,
                            engine=Engine.XLSX,
                            message="Package is missing mandatory 'xl/workbook.xml' part.",
                            location=Location(part="xl/workbook.xml"),
                        )
                    )
                    empty_inv = PackageInventory(file_name=filename, fidelity_tier="T1")
                    return PreflightScanResult(empty_inv, issues, guarantees)

                for name in names:
                    if name.startswith("xl/vbaProject.bin"):
                        has_vba = True
                    elif name.startswith("xl/externalLinks/"):
                        has_ext_links = True
                    elif name.startswith("xl/drawings/"):
                        has_drawings = True
                    elif name.startswith("xl/media/"):
                        has_images = True
                    elif name.startswith("xl/charts/"):
                        has_charts = True
                    elif name.startswith("xl/pivotTables/"):
                        has_pivots = True
                    elif name.startswith("xl/slicers/"):
                        has_slicers = True

        except Exception as exc:
            issues.append(
                Issue(
                    code="E-PKG-UNREADABLE",
                    severity=Severity.ERROR,
                    engine=Engine.XLSX,
                    message=f"Cannot read OpenXML zip package: {exc}",
                    location=Location(part="package"),
                )
            )
            empty_inv = PackageInventory(file_name=filename, fidelity_tier="T1")
            return PreflightScanResult(empty_inv, issues, guarantees)

        # 3. Security & Feature Issues
        if has_vba:
            issues.append(
                Issue(
                    code="E-SEC-MACRO",
                    severity=Severity.ERROR,
                    engine=Engine.XLSX,
                    message="Macro-enabled workbook (xl/vbaProject.bin) detected. Execution forbidden by SEC-04.",
                    location=Location(part="xl/vbaProject.bin"),
                    fixable_by=FixableBy.HUMAN,
                )
            )

        if has_ext_links:
            issues.append(
                Issue(
                    code="W-PKG-EXTLINK",
                    severity=Severity.WARNING,
                    engine=Engine.XLSX,
                    message="External workbook link detected in xl/externalLinks. Engine will not follow external links.",
                    location=Location(part="xl/externalLinks"),
                    fixable_by=FixableBy.HUMAN,
                )
            )

        if has_pivots:
            issues.append(
                Issue(
                    code="W-PKG-PIVOT-REFRESH",
                    severity=Severity.WARNING,
                    engine=Engine.XLSX,
                    message="Pivot Tables detected. Cache will not be recalculated server-side; refresh on open required.",
                    location=Location(part="xl/pivotTables"),
                    fixable_by=FixableBy.HUMAN,
                )
            )

        if has_slicers:
            issues.append(
                Issue(
                    code="W-PKG-UNSUPPORTED-SLICER",
                    severity=Severity.WARNING,
                    engine=Engine.XLSX,
                    message="Slicers or Timeline controls detected. Writing back may compromise control layout.",
                    location=Location(part="xl/slicers"),
                    fixable_by=FixableBy.HUMAN,
                )
            )

        # 4. Open with openpyxl for deep sheet inspection
        sheet_inventories: List[SheetInventory] = []
        total_formulas = 0
        total_merges = 0
        defined_names_count = 0

        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=False, keep_vba=True)
            defined_names_count = len(wb.defined_names)

            for ws in wb.worksheets:
                state = ws.sheet_state or "visible"
                if state != "visible":
                    issues.append(
                        Issue(
                            code="W-PKG-HIDDEN",
                            severity=Severity.WARNING,
                            engine=Engine.XLSX,
                            message=f"Worksheet '{ws.title}' has hidden state '{state}'.",
                            location=Location(sheet=ws.title),
                            fixable_by=FixableBy.AI,
                        )
                    )

                # Count formulas
                f_count = 0
                for row in ws.iter_rows(values_only=True):
                    for val in row:
                        if isinstance(val, str) and val.startswith("="):
                            f_count += 1

                m_count = len(ws.merged_cells.ranges)
                total_formulas += f_count
                total_merges += m_count

                has_filter = bool(ws.auto_filter and ws.auto_filter.ref)
                cf_count = len(ws.conditional_formatting)
                dv_count = len(ws.data_validations.dataValidation) if hasattr(ws.data_validations, "dataValidation") else 0

                s_inv = SheetInventory(
                    name=ws.title,
                    visible_state=state,  # type: ignore[arg-type]
                    max_row=ws.max_row or 0,
                    max_column=ws.max_column or 0,
                    formulas_count=f_count,
                    merges_count=m_count,
                    has_autofilter=has_filter,
                    conditional_formats_count=cf_count,
                    data_validations_count=dv_count,
                )
                sheet_inventories.append(s_inv)
            wb.close()

        except Exception as exc:
            issues.append(
                Issue(
                    code="E-PKG-OPENPYXL-PARSE-FAILED",
                    severity=Severity.ERROR,
                    engine=Engine.XLSX,
                    message=f"openpyxl failed to parse sheets: {exc}",
                    location=Location(part="xl/workbook.xml"),
                )
            )

        # 5. Determine Fidelity Tier
        tier: FidelityTier
        if has_vba:
            tier = "T4"
        elif has_slicers:
            tier = "T3"
        elif has_drawings or has_charts or has_images:
            tier = "T2"
        else:
            tier = "T1"

        inventory = PackageInventory(
            file_name=filename,
            sheets=sheet_inventories,
            total_formulas=total_formulas,
            total_merges=total_merges,
            has_drawingml=has_drawings,
            has_images=has_images,
            has_charts=has_charts,
            has_pivot_tables=has_pivots,
            has_vba_macros=has_vba,
            has_external_links=has_ext_links,
            has_slicers=has_slicers,
            defined_names_count=defined_names_count,
            fidelity_tier=tier,
        )

        return PreflightScanResult(inventory, issues, guarantees)


# Global singleton
preflight_scanner = PreflightScanner()
