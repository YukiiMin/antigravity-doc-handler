"""
doctools.core.xlsx.template.manifest_parser — Trình phân tích và kiểm định Manifest của XLSX Template.
Chuyển đổi manifest dạng Dict / YAML / JSON thành model XlsxTemplateManifest.
Thực hiện kiểm tra tính nhất quán giữa Manifest và tài liệu Excel (.xlsx).
"""

from __future__ import annotations
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Tuple, Union
import openpyxl
from openpyxl.utils.cell import range_boundaries
import yaml

from doctools.contract.issues import Engine, Issue, Severity
from doctools.contract.xlsx.manifest import AnchorConfig, XlsxTemplateManifest


class ManifestParseError(Exception):
    """Lỗi khi phân tích cú pháp dữ liệu manifest."""
    pass


def parse_manifest(
    manifest_input: Union[Dict[str, Any], str, Path, XlsxTemplateManifest]
) -> XlsxTemplateManifest:
    """
    Phân tích đầu vào thành đối tượng XlsxTemplateManifest.
    Hỗ trợ đối tượng có sẵn, Dictionary, chuỗi JSON/YAML hoặc đường dẫn file .yaml/.json.
    """
    if isinstance(manifest_input, XlsxTemplateManifest):
        return manifest_input

    if isinstance(manifest_input, (str, Path)):
        p = Path(manifest_input)
        if p.exists() and p.is_file():
            raw_text = p.read_text(encoding="utf-8")
            data = yaml.safe_load(raw_text)
        else:
            try:
                data = yaml.safe_load(str(manifest_input))
            except Exception as e:
                raise ManifestParseError(f"Không thể giải mã YAML/JSON: {e}") from e
    elif isinstance(manifest_input, dict):
        data = manifest_input
    else:
        raise ManifestParseError(f"Đầu vào manifest không hợp lệ: {type(manifest_input)}")

    if not isinstance(data, dict):
        raise ManifestParseError("Manifest phải là một từ điển (dictionary/mapping).")

    try:
        return XlsxTemplateManifest.model_validate(data)
    except Exception as exc:
        raise ManifestParseError(f"Dữ liệu không khớp schema XlsxTemplateManifest: {exc}") from exc


def parse_locked_zone(zone_str: str) -> Tuple[str, Tuple[int, int, int, int]]:
    """
    Phân tích cú pháp chuỗi locked_zone dạng 'Sheet!A1:B10' hoặc '*!A1:T8'.
    Trả về (sheet_pattern, (min_col, min_row, max_col, max_row)).
    """
    if "!" not in zone_str:
        raise ValueError(f"Vùng khóa thiếu tên sheet hoặc '!': '{zone_str}'")
    sheet_part, cell_range = zone_str.split("!", 1)
    sheet_name = sheet_part.strip().strip("'")
    boundaries = range_boundaries(cell_range.strip())
    if any(b is None for b in boundaries):
        raise ValueError(f"Dải ô không hợp lệ: '{cell_range}'")
    return sheet_name, boundaries


def validate_manifest_against_template(
    manifest: XlsxTemplateManifest,
    template_path_or_wb: Union[Path, str, openpyxl.Workbook],
) -> List[Issue]:
    """
    Kiểm tra tính nhất quán giữa khai báo trong Manifest và workbook Excel thực tế.
    Trả về danh sách Issue nếu phát hiện bất thường (không làm crash hệ thống).
    """
    issues: List[Issue] = []
    wb: Optional[openpyxl.Workbook] = None
    should_close = False

    if isinstance(template_path_or_wb, openpyxl.Workbook):
        wb = template_path_or_wb
    else:
        p = Path(template_path_or_wb)
        if not p.exists():
            issues.append(Issue(
                code="E-TPL-FILE-NOT-FOUND",
                severity=Severity.ERROR,
                engine=Engine.XLSX,
                message=f"Tệp template không tồn tại: {p}",
            ))
            return issues
        try:
            wb = openpyxl.load_workbook(p, data_only=False, read_only=False)
            should_close = True
        except Exception as e:
            issues.append(Issue(
                code="E-TPL-LOAD-FAILED",
                severity=Severity.ERROR,
                engine=Engine.XLSX,
                message=f"Không thể mở tệp template Excel: {e}",
            ))
            return issues

    try:
        sheet_names = wb.sheetnames

        # 1. Kiểm tra reference_sheet
        if manifest.reference_sheet not in sheet_names:
            issues.append(Issue(
                code="E-TPL-REF-SHEET-NOT-FOUND",
                severity=Severity.ERROR,
                engine=Engine.XLSX,
                message=(
                    f"reference_sheet '{manifest.reference_sheet}' không tồn tại trong template. "
                    f"Các sheet có sẵn: {sheet_names}"
                ),
            ))

        # 2. Kiểm tra sibling_sheets
        for sib in manifest.sibling_sheets:
            if sib not in sheet_names:
                issues.append(Issue(
                    code="E-TPL-SIBLING-SHEET-NOT-FOUND",
                    severity=Severity.ERROR,
                    engine=Engine.XLSX,
                    message=f"sibling_sheet '{sib}' khai báo trong manifest nhưng không có trong workbook.",
                ))

        # 3. Kiểm tra cú pháp locked_zones
        for zone in manifest.locked_zones:
            try:
                parse_locked_zone(zone)
            except Exception as e:
                issues.append(Issue(
                    code="E-TPL-LOCKED-ZONE-SYNTAX",
                    severity=Severity.ERROR,
                    engine=Engine.XLSX,
                    message=f"Cú pháp locked_zone không hợp lệ '{zone}': {e}",
                ))

        # 4. Kiểm tra semantic anchors trên reference_sheet (nếu sheet tồn tại)
        if manifest.reference_sheet in sheet_names:
            ref_ws = wb[manifest.reference_sheet]
            for anchor_name, anchor_cfg in manifest.anchors.items():
                if isinstance(anchor_cfg, str):
                    cfg = AnchorConfig(keyword=anchor_cfg)
                elif isinstance(anchor_cfg, dict):
                    cfg = AnchorConfig(**anchor_cfg)
                elif isinstance(anchor_cfg, AnchorConfig):
                    cfg = anchor_cfg
                else:
                    continue

                kw = cfg.keyword or (cfg.keywords[0] if cfg.keywords else None)
                if not kw:
                    continue

                matches: List[str] = []
                max_row = min(cfg.scan_rows, ref_ws.max_row or 1)
                max_col = ref_ws.max_column or 1

                for row_idx in range(1, max_row + 1):
                    for col_idx in range(1, max_col + 1):
                        cell_val = ref_ws.cell(row=row_idx, column=col_idx).value
                        if cell_val is None:
                            continue
                        str_val = str(cell_val).strip()
                        matched = False
                        if cfg.match == "exact" and str_val == kw:
                            matched = True
                        elif cfg.match == "contains" and kw in str_val:
                            matched = True
                        elif cfg.match == "regex" and re.search(kw, str_val):
                            matched = True

                        if matched:
                            coord = f"{ref_ws.cell(row=row_idx, column=col_idx).coordinate}"
                            matches.append(coord)

                if len(matches) == 0:
                    issues.append(Issue(
                        code="E-TPL-ANCHOR-MISSING",
                        severity=Severity.ERROR,
                        engine=Engine.XLSX,
                        message=(
                            f"Anchor '{anchor_name}' với từ khóa '{kw}' không tìm thấy trên sheet "
                            f"'{manifest.reference_sheet}' trong {max_row} dòng đầu."
                        ),
                        evidence={"anchor": anchor_name, "keyword": kw, "scan_rows": max_row},
                    ))
                elif len(matches) > 1 and cfg.scope == "first_table":
                    issues.append(Issue(
                        code="E-TPL-ANCHOR-AMBIGUOUS",
                        severity=Severity.ERROR,
                        engine=Engine.XLSX,
                        message=(
                            f"Anchor '{anchor_name}' khớp nhiều ô ({matches}) trong phạm vi first_table. "
                            f"Cần chỉ định từ khóa hoặc phạm vi hẹp hơn."
                        ),
                        evidence={"anchor": anchor_name, "keyword": kw, "matches": matches},
                    ))

        # 5. Kiểm tra prototype_rows
        for sheet_key, p_row in manifest.prototype_rows.items():
            target_ws_name = manifest.reference_sheet if sheet_key == "default" else sheet_key
            if target_ws_name in sheet_names:
                target_ws = wb[target_ws_name]
                if p_row <= 0 or (target_ws.max_row and p_row > target_ws.max_row + 10):
                    issues.append(Issue(
                        code="E-TPL-PROTOTYPE-ROW-INVALID",
                        severity=Severity.WARNING,
                        engine=Engine.XLSX,
                        message=f"prototype_row {p_row} trên sheet '{target_ws_name}' vượt ngoài biên dữ liệu.",
                    ))

    finally:
        if should_close and wb is not None:
            wb.close()

    return issues
