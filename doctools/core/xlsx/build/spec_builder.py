"""
doctools.core.xlsx.build.spec_builder — Động cơ xây dựng bảng tính từ XlsxSpec (Spec Builder).
Tuân thủ Path B, FR-08, FR-15 của Foundation Plan v1.1:
- Dựng toàn diện workbook từ JSON đặc tả thuần XlsxSpec.
- Sinh biểu đồ DrawingML native (BarChart, LineChart, PieChart, AreaChart) không qua trung gian.
- Định dạng design tokens: font, fill, border, alignment, number_format.
- Ghi an toàn ô gộp và đồng bộ viền 4 phía (ERR_XLSX_004).
"""

from __future__ import annotations
from typing import Any, Dict, List, Optional, Tuple
import openpyxl
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.worksheet.worksheet import Worksheet

from doctools.contract.issues import Engine, Issue, Severity
from doctools.contract.xlsx.spec import ChartSpec, SheetSpec, StyleSpec, TableSpec, XlsxSpec
from doctools.core.xlsx.mutate.style_cloner import sync_merged_borders


class XlsxSpecBuilder:
    """Bộ tạo dựng workbook từ XlsxSpec."""

    def __init__(self) -> None:
        pass

    def build(self, spec: XlsxSpec) -> Tuple[openpyxl.Workbook, List[Issue]]:
        """Dựng workbook hoàn chỉnh từ đặc tả XlsxSpec."""
        issues: List[Issue] = []
        wb = openpyxl.Workbook()

        if not spec.sheets:
            issues.append(
                Issue(
                    code="E-XLSX-SPEC-NO-SHEETS",
                    severity=Severity.ERROR,
                    engine=Engine.XLSX,
                    message="XlsxSpec không chứa sheet nào.",
                )
            )
            return wb, issues

        # Xóa sheet mặc định ban đầu nếu có sheet được khai báo
        wb.remove(wb.active)

        for s_idx, sheet_spec in enumerate(spec.sheets):
            ws = wb.create_sheet(title=sheet_spec.name)

            # Cấu hình độ rộng cột
            for col_letter, width in sheet_spec.column_widths.items():
                ws.column_dimensions[col_letter].width = width

            # Cố định dòng/cột freeze_panes
            if sheet_spec.freeze_panes:
                ws.freeze_panes = sheet_spec.freeze_panes

            # Dựng các khối bảng (TableSpec)
            for tbl in sheet_spec.tables:
                self._build_table(ws, tbl)

            # Ghi các ô riêng lẻ (CellSpec)
            for cell_spec in sheet_spec.cells:
                r, c = coordinate_to_tuple(cell_spec.coordinate)
                cell = ws.cell(row=r, column=c)
                cell.value = cell_spec.value
                if cell_spec.style:
                    self._apply_style(cell, cell_spec.style)

            # Xử lý vùng gộp (Merged Ranges)
            for rng_str in sheet_spec.merged_ranges:
                try:
                    ws.merge_cells(rng_str)
                    sync_merged_borders(ws, rng_str)
                except Exception as exc:
                    issues.append(
                        Issue(
                            code="W-XLSX-MERGE-FAILED",
                            severity=Severity.WARNING,
                            engine=Engine.XLSX,
                            message=f"Không thể gộp dải ô '{rng_str}': {exc}",
                        )
                    )

            # Sinh biểu đồ DrawingML native (ChartSpec)
            for chart_spec in sheet_spec.charts:
                try:
                    chart_obj = self._create_chart(ws, chart_spec)
                    if chart_obj:
                        ws.add_chart(chart_obj, chart_spec.anchor)
                except Exception as exc:
                    issues.append(
                        Issue(
                            code="W-XLSX-CHART-FAILED",
                            severity=Severity.WARNING,
                            engine=Engine.XLSX,
                            message=f"Lỗi khi dựng biểu đồ '{chart_spec.title or chart_spec.chart_type}': {exc}",
                        )
                    )

        # Cấu hình tính toán tự động khi mở
        if hasattr(wb, "calculation"):
            wb.calculation.calcMode = "auto"
            wb.calculation.fullCalcOnLoad = True

        return wb, issues

    def _build_table(self, ws: Worksheet, tbl: TableSpec) -> None:
        """Dựng khối bảng dữ liệu trên worksheet."""
        start_r, start_c = coordinate_to_tuple(tbl.start_coordinate)
        cur_r = start_r

        # Ghi Headers nếu có
        if tbl.headers:
            for c_idx, h_text in enumerate(tbl.headers, start=start_c):
                cell = ws.cell(row=cur_r, column=c_idx, value=h_text)
                if tbl.header_style:
                    self._apply_style(cell, tbl.header_style)
            cur_r += 1

        # Ghi các dòng dữ liệu
        for row_data in tbl.rows:
            for c_offset, val in enumerate(row_data, start=0):
                c_idx = start_c + c_offset
                cell = ws.cell(row=cur_r, column=c_idx)
                if (c_offset + 1) in tbl.id_columns:
                    cell.number_format = "@"
                    cell.value = str(val) if val is not None else ""
                else:
                    cell.value = val

                if tbl.row_style:
                    self._apply_style(cell, tbl.row_style)
            cur_r += 1

    def _create_chart(self, ws: Worksheet, spec: ChartSpec) -> Optional[Any]:
        """Tạo đối tượng chart openpyxl native."""
        chart: Any
        if spec.chart_type == "bar":
            chart = BarChart()
            chart.type = "bar"
        elif spec.chart_type == "col":
            chart = BarChart()
            chart.type = "col"
        elif spec.chart_type == "line":
            chart = LineChart()
        elif spec.chart_type == "pie":
            chart = PieChart()
        elif spec.chart_type == "area":
            chart = AreaChart()
        elif spec.chart_type == "doughnut":
            chart = DoughnutChart()
        else:
            return None

        if spec.title:
            chart.title = spec.title

        # Tham chiếu dải dữ liệu
        data_ref = Reference(
            ws,
            min_col=spec.data_min_col,
            min_row=spec.data_min_row,
            max_col=spec.data_max_col,
            max_row=spec.data_max_row,
        )
        chart.add_data(data_ref, titles_from_data=spec.titles_from_data)

        # Tham chiếu nhãn trục / danh mục nếu có
        if spec.categories_min_col and spec.categories_min_row:
            cat_ref = Reference(
                ws,
                min_col=spec.categories_min_col,
                min_row=spec.categories_min_row,
                max_col=spec.categories_max_col or spec.categories_min_col,
                max_row=spec.categories_max_row or spec.categories_min_row,
            )
            chart.set_categories(cat_ref)

        chart.width = spec.width_cm
        chart.height = spec.height_cm
        return chart

    def _apply_style(self, cell: Any, style: StyleSpec) -> None:
        """Áp dụng StyleSpec vào ô."""
        if style.font_name or style.font_size or style.bold or style.italic or style.font_color:
            cell.font = Font(
                name=style.font_name or "Arial",
                size=style.font_size or 11,
                bold=style.bold or False,
                italic=style.italic or False,
                color=style.font_color,
            )

        if style.fill_color:
            cell.fill = PatternFill(
                start_color=style.fill_color,
                end_color=style.fill_color,
                fill_type="solid",
            )

        if style.border_style:
            side = Side(style=style.border_style, color=style.border_color or "000000")
            cell.border = Border(left=side, right=side, top=side, bottom=side)

        if style.alignment_h or style.alignment_v:
            cell.alignment = Alignment(
                horizontal=style.alignment_h,
                vertical=style.alignment_v,
            )

        if style.number_format:
            cell.number_format = style.number_format
