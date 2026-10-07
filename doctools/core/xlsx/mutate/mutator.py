"""
doctools.core.xlsx.mutate.mutator — Động cơ đột biến bảng tính Excel (XlsxMutator).
Tuân thủ E1..E14, D-02, D-04, FR-05, FR-06, FR-09 của Foundation Plan v1.1:
- Cưỡng chế vùng khóa locked_zones (E-XLSX-LOCK-001).
- Nhân bản dòng mẫu Prototype Row (consume prototype hoặc chèn mới).
- Ép kiểu định dạng chuỗi '@' cho cột id_columns để giữ số 0 đầu (E12).
- Tích hợp ShiftManager dịch công thức và dải ô tổng cộng tự động (E14).
- Ghi an toàn vào merged cells (chỉ ghi Top-Left) và đồng bộ viền (ERR_XLSX_004).
"""

from __future__ import annotations
from typing import Any, Dict, List, Optional, Set, Tuple, Union
import openpyxl
from openpyxl.utils.cell import column_index_from_string, coordinate_to_tuple
from openpyxl.worksheet.worksheet import Worksheet

from doctools.contract.issues import Engine, Issue, Severity
from doctools.contract.xlsx.mutation import CellUpdate, MutationSpec, TableExpansion
from doctools.core.xlsx.shift.shift_manager import ShiftManager
from doctools.core.xlsx.template.manifest_parser import parse_locked_zone
from .style_cloner import clone_row_style, sync_merged_borders


class XlsxMutator:
    """Bộ thực thi đột biến dữ liệu bảng tính theo MutationSpec."""

    def __init__(self, workbook: openpyxl.Workbook, locked_zones: Optional[List[str]] = None) -> None:
        self.workbook = workbook
        self.locked_zones = locked_zones or []

    def _is_cell_locked(self, sheet_name: str, row: int, col: int, extra_zones: List[str]) -> bool:
        """Kiểm tra xem tọa độ (row, col) có rơi vào vùng cấm ghi locked_zones hay không."""
        all_zones = self.locked_zones + extra_zones
        for zone_str in all_zones:
            try:
                z_sheet, (min_c, min_r, max_c, max_r) = parse_locked_zone(zone_str)
                if z_sheet == "*" or z_sheet == sheet_name:
                    if min_r <= row <= max_r and min_c <= col <= max_c:
                        return True
            except Exception:
                continue
        return False

    def _resolve_col_idx(self, col_key: Union[str, int]) -> int:
        """Chuyển đổi tên cột (vd: 'A' hoặc 1) thành chỉ số cột 1-indexed."""
        if isinstance(col_key, int):
            return col_key
        if col_key.isdigit():
            return int(col_key)
        return column_index_from_string(col_key)

    def _write_cell_safe(
        self,
        ws: Worksheet,
        row: int,
        col: int,
        value: Any,
        is_id_column: bool = False,
    ) -> None:
        """Ghi dữ liệu an toàn vào ô, tránh ghi vào ô phụ của MergedCell (ERR_XLSX_004)."""
        # Kiểm tra xem ô có phải là ô con trong dải merge không
        for rng in ws.merged_cells.ranges:
            min_c, min_r, max_c, max_r = rng.bounds
            if min_r <= row <= max_r and min_c <= col <= max_c:
                if row != min_r or col != min_c:
                    # Là ô phụ trong merged cell: bỏ qua để tránh lỗi ghi read-only
                    return

        cell = ws.cell(row=row, column=col)
        if is_id_column:
            cell.number_format = "@"
            cell.value = str(value) if value is not None else ""
        else:
            cell.value = value

    def mutate(self, spec: MutationSpec) -> List[Issue]:
        """Thực thi MutationSpec trên Workbook."""
        issues: List[Issue] = []

        # 1. Kiểm tra locked_zones cho các cập nhật ô đơn lẻ (CellUpdate)
        for update in spec.cell_updates:
            row, col = coordinate_to_tuple(update.coordinate)
            if self._is_cell_locked(update.sheet, row, col, spec.locked_zones):
                issues.append(Issue(
                    code="E-XLSX-LOCK-001",
                    severity=Severity.ERROR,
                    engine=Engine.XLSX,
                    message=f"Thao tác ghi bị từ chối: ô {update.sheet}!{update.coordinate} nằm trong vùng khóa.",
                    evidence={"sheet": update.sheet, "coordinate": update.coordinate},
                ))
                return issues

        # 2. Xử lý từng TableExpansion
        shift_mgr = ShiftManager(self.workbook)

        for exp in spec.expansions:
            if exp.sheet not in self.workbook.sheetnames:
                issues.append(Issue(
                    code="E-XLSX-SHEET-NOT-FOUND",
                    severity=Severity.ERROR,
                    engine=Engine.XLSX,
                    message=f"Sheet '{exp.sheet}' không tồn tại trong workbook.",
                ))
                continue

            ws = self.workbook[exp.sheet]
            num_records = len(exp.rows_data)
            if num_records == 0:
                continue

            # Xác định proto_row và start_row
            proto_row = exp.prototype_row or exp.start_row or 2
            start_row = exp.start_row or proto_row

            # Kiểm tra xem vùng bắt đầu ghi có nằm trong locked_zones không
            if self._is_cell_locked(exp.sheet, start_row, 1, spec.locked_zones):
                issues.append(Issue(
                    code="E-XLSX-LOCK-001",
                    severity=Severity.ERROR,
                    engine=Engine.XLSX,
                    message=f"Thao tác mở rộng bảng bị từ chối: dòng {start_row} trên sheet '{exp.sheet}' nằm trong vùng khóa.",
                    evidence={"sheet": exp.sheet, "row": start_row},
                ))
                return issues

            # Lập tập hợp các cột id_columns
            id_cols_set: Set[int] = set()
            for id_c in exp.id_columns:
                id_cols_set.add(self._resolve_col_idx(id_c))

            # Xác định ánh xạ cột
            col_map = exp.columns_mapping

            if exp.consume_prototype:
                # Ghi đè dòng mẫu bằng bản ghi đầu tiên
                first_row_data = exp.rows_data[0]
                self._populate_row(ws, proto_row, first_row_data, col_map, id_cols_set)

                # Chèn các dòng tiếp theo (nếu có từ 2 bản ghi trở lên)
                if num_records > 1:
                    insert_count = num_records - 1
                    insert_at = proto_row + 1
                    ws.insert_rows(insert_at, amount=insert_count)

                    # Dịch chuyển toàn bộ công thức và dải ô qua ShiftManager
                    shift_issues = shift_mgr.shift_sheet_rows(
                        sheet_name=exp.sheet,
                        insert_at_row=insert_at,
                        num_rows=insert_count,
                        range_policy=exp.range_policy,
                    )
                    issues.extend(shift_issues)

                    # Sao chép style từ proto_row và ghi dữ liệu các dòng tiếp theo
                    for idx in range(1, num_records):
                        current_r = proto_row + idx
                        clone_row_style(ws, proto_row, current_r, ws.max_column)
                        row_data = exp.rows_data[idx]
                        self._populate_row(ws, current_r, row_data, col_map, id_cols_set)
            else:
                # Không consume prototype: chèn mới num_records dòng tại start_row
                ws.insert_rows(start_row, amount=num_records)
                shift_issues = shift_mgr.shift_sheet_rows(
                    sheet_name=exp.sheet,
                    insert_at_row=start_row,
                    num_rows=num_records,
                    range_policy=exp.range_policy,
                )
                issues.extend(shift_issues)

                for idx in range(num_records):
                    current_r = start_row + idx
                    clone_row_style(ws, proto_row, current_r, ws.max_column)
                    row_data = exp.rows_data[idx]
                    self._populate_row(ws, current_r, row_data, col_map, id_cols_set)

            # Đồng bộ viền cho các ô gộp trên sheet
            for rng in list(ws.merged_cells.ranges):
                sync_merged_borders(ws, rng)

        # 3. Thực thi các cập nhật ô đơn lẻ (CellUpdate)
        for update in spec.cell_updates:
            if update.sheet in self.workbook.sheetnames:
                ws = self.workbook[update.sheet]
                r, c = coordinate_to_tuple(update.coordinate)
                self._write_cell_safe(ws, r, c, update.value)

        return issues

    def _populate_row(
        self,
        ws: Worksheet,
        row_idx: int,
        row_data: Union[Dict[str, Any], List[Any]],
        col_map: Optional[Dict[str, Union[str, int]]],
        id_cols: Set[int],
    ) -> None:
        """Điền dữ liệu một bản ghi vào dòng row_idx."""
        if isinstance(row_data, dict):
            for key, val in row_data.items():
                if col_map and key in col_map:
                    target_c = self._resolve_col_idx(col_map[key])
                elif key.isdigit():
                    target_c = int(key)
                elif len(key) <= 3 and key.isalpha():
                    target_c = column_index_from_string(key)
                else:
                    continue
                is_id = (target_c in id_cols) or (key in id_cols)
                self._write_cell_safe(ws, row_idx, target_c, val, is_id_column=is_id)
        elif isinstance(row_data, (list, tuple)):
            for c_offset, val in enumerate(row_data, start=1):
                is_id = (c_offset in id_cols)
                self._write_cell_safe(ws, row_idx, c_offset, val, is_id_column=is_id)
