"""
doctools.core.xlsx.mutate.style_cloner — Bộ sao chép định dạng và đồng bộ viền ô (Style Cloner).
Tuân thủ E1..E14, ERR_XLSX_004, ERR_XLSX_006 của Foundation Plan v1.1:
- Sao chép 100% token kiểu dáng từ dòng mẫu (Prototype Row): font, fill, border, alignment, number_format.
- Đồng bộ viền cho các ô gộp (sync_merged_borders) chống rách viền khi render trong Excel.
"""

from __future__ import annotations
from copy import copy
from typing import Optional, Union
import openpyxl
from openpyxl.cell import Cell
from openpyxl.styles import Border, Side
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet


def clone_cell_style(source_cell: Cell, target_cell: Cell) -> None:
    """Sao chép toàn bộ thuộc tính định dạng từ ô nguồn sang ô đích."""
    if source_cell.has_style:
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.number_format:
            target_cell.number_format = str(source_cell.number_format)
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)


def clone_row_style(
    ws: Worksheet,
    source_row_idx: int,
    target_row_idx: int,
    max_col: Optional[int] = None,
) -> None:
    """
    Sao chép định dạng của toàn bộ dòng nguồn sang dòng đích.
    Bao gồm định dạng từng ô và chiều cao dòng (row height).
    """
    limit_col = max_col or ws.max_column or 1
    for c_idx in range(1, limit_col + 1):
        src_cell = ws.cell(row=source_row_idx, column=c_idx)
        tgt_cell = ws.cell(row=target_row_idx, column=c_idx)
        clone_cell_style(src_cell, tgt_cell)

    src_height = ws.row_dimensions[source_row_idx].height
    if src_height is not None:
        ws.row_dimensions[target_row_idx].height = src_height


def sync_merged_borders(ws: Worksheet, cell_range: Union[CellRange, str]) -> None:
    """
    Đồng bộ đường viền 4 phía cho toàn bộ các ô trong vùng gộp (MergedCellRange).
    Tránh lỗi viền bị đứt đoạn hoặc biến mất khi Excel render vùng gộp (ERR_XLSX_004).
    """
    rng = CellRange(cell_range) if isinstance(cell_range, str) else cell_range
    min_col, min_row, max_col, max_row = rng.bounds

    top_left = ws.cell(row=min_row, column=min_col)
    tl_border = top_left.border or Border()

    top_side = tl_border.top or Side()
    left_side = tl_border.left or Side()

    # Lấy viền dưới và viền phải từ ô biên nếu có
    bottom_right = ws.cell(row=max_row, column=max_col)
    br_border = bottom_right.border or Border()
    bottom_side = br_border.bottom or tl_border.bottom or Side()
    right_side = br_border.right or tl_border.right or Side()

    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=r, column=c)
            current_border = cell.border or Border()

            new_top = top_side if r == min_row else current_border.top
            new_bottom = bottom_side if r == max_row else current_border.bottom
            new_left = left_side if c == min_col else current_border.left
            new_right = right_side if c == max_col else current_border.right

            cell.border = Border(
                top=new_top,
                bottom=new_bottom,
                left=new_left,
                right=new_right,
                diagonal=current_border.diagonal,
                diagonal_direction=current_border.diagonal_direction,
            )
