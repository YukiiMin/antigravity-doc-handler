import os
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r'e:\Do_an_SU26\SuperMarketBot-Android-Robot\src\Report3_Software Requirement Specification.docx'

print("=== VERIFYING PATCHED SRS DOCUMENT ===")
doc = docx.Document(DOC_PATH)

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

# Check Section 3.1
print("\n--- 1. Verification of Section 3.1 Screen Flows ---")
for i, p in enumerate(doc.paragraphs[:100]):
    txt = p.text.strip()
    if any(k in txt for k in ["3.1.1", "3.1.2", "3.1.3"]):
        print(f"[{i:3d}] ({p.style.name}): {txt}")
        # check next paragraph for image
        next_p = doc.paragraphs[i+1]
        has_img = any('graphic' in r._element.xml for r in next_p.runs)
        print(f"      Image in next paragraph: {has_img}")

# Check Table 3
print("\n--- 2. Verification of Table 3 (Robot Screen Descriptions) ---")
t3 = doc.tables[3]
print(f"Table 3 rows: {len(t3.rows)}")
last_row = [c.text.strip() for c in t3.rows[-1].cells]
print(f"Last row in Table 3: {last_row}")

# Check Section 3.2
print("\n--- 3. Verification of Section 3.2 Subsections and Images ---")
in_sec_32 = False
sec32_headings = []
sec32_images = 0
todos_found = []

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("3.2 App for Robot"):
        in_sec_32 = True
    elif txt.startswith("3.3 Android Customer App"):
        in_sec_32 = False
        
    if in_sec_32:
        if p.style.name == "Heading 3":
            sec32_headings.append((i, txt))
        img_count = sum(1 for r in p.runs if 'graphic' in r._element.xml)
        sec32_images += img_count
        if "[TODO" in txt:
            todos_found.append((i, txt))

print(f"Section 3.2 Subsections found ({len(sec32_headings)}):")
for idx, h in sec32_headings:
    print(f"   [{idx:3d}] {h}")

print(f"Section 3.2 total embedded images: {sec32_images}")
print(f"Section 3.2 TODO flags found: {len(todos_found)}")
if todos_found:
    for tf in todos_found:
        print(f"   [WARN] {tf}")

# Check Section 3.2.8 Business Flow text keywords
print("\n--- 4. Verification of 3.2.8 Autonomous Advertising & Interactive Guidance Content ---")
sec328_text = []
in_328 = False
for p in doc.paragraphs:
    txt = p.text.strip()
    if "3.2.8" in txt:
        in_328 = True
    elif in_328 and p.style.name in ["Heading 2", "Heading 3"]:
        in_328 = False
    if in_328:
        sec328_text.append(txt)

full_328_str = "\n".join(sec328_text)

keywords = [
    "AdScore", "VIP", "Pro", "Standard",
    "Theo Kệ", "Theo Khu vực", "Tuyến tuần tra tự do",
    "Playlist", "Khách vãng lai", "Khách hàng thành viên",
    "AdMultiProductSelectScreen", "CartGuideMapScreen",
    "AdInterruptionService", "ARRIVED", "khôi phục"
]

for kw in keywords:
    present = kw.lower() in full_328_str.lower()
    print(f"   Keyword '{kw}': {'FOUND' if present else 'MISSING!'}")

# Check Section 3.5 Heading
print("\n--- 5. Verification of Section 3.5 Heading Level ---")
for i, p in enumerate(doc.paragraphs):
    if "3.5 Staff Mobile Experience" in p.text:
        print(f"[{i}] Style: {p.style.name} | Text: {p.text}")
        break

print("\n=== ALL VERIFICATION CHECKS COMPLETED ===")
