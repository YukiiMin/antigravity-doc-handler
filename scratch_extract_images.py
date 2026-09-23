import docx
import os
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc_path = r'e:\Do_an_SU26\SuperMarketBot-Android-Robot\src\Report3_Software Requirement Specification.docx'
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")

current_h1 = ""
current_h2 = ""
current_h3 = ""

images_info = []

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if p.style.name == "Heading 1":
        current_h1 = txt
        current_h2 = ""
        current_h3 = ""
    elif p.style.name == "Heading 2":
        current_h2 = txt
        current_h3 = ""
    elif p.style.name in ["Heading 3", "Heading 4"]:
        if txt.startswith("3.") or txt.startswith("2.") or txt.startswith("4."):
            current_h3 = txt
            
    for r in p.runs:
        if 'graphic' in r._element.xml:
            # find blip r:embed
            import xml.etree.ElementTree as ET
            root = ET.fromstring(r._element.xml)
            # find all blip elements
            blips = root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
            for b in blips:
                rId = b.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                # get target part
                target = doc.part.rels[rId].target_ref
                images_info.append({
                    'p_idx': i,
                    'h1': current_h1,
                    'h2': current_h2,
                    'h3': current_h3,
                    'rId': rId,
                    'target': target,
                    'text': txt[:100]
                })

print(f"Found {len(images_info)} images in runs.")
import json
with open('images_info.json', 'w', encoding='utf-8') as f:
    json.dump(images_info, f, ensure_ascii=False, indent=2)

print("Saved to images_info.json")
